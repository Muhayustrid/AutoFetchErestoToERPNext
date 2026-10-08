"""One-off: injeksi agregat penjualan eResto -> Journal Entry per outlet per hari.

Fase agregat docs/03 (Phase 1): Dr kas/clearing + Cr pendapatan, TANPA
gerakkan stok (eResto tetap stock truth). Idempoten via
cheque_no = ERESTO-<branch>-<date>; rerun tidak duplikat.

Pemetaan bayar -> akun (pola CoA diverifikasi di 3 company):
  cash -> Kas Kasir | ewallet -> Qris | debit -> Debit Card
  kredit -> Credit Card | transfer -> E-commerce | Cr -> Penjualan
Baris bernilai 0 tidak diposting. Order total=0 (void/OC) diskip.

Env: /tmp/eresto_mig.env (ERPNEXT_URL, ERPNEXT_API_KEY). Token tak diprint.
Usage:
  python3 inject_sales_aggregate.py [--src XLSX] [--dry-run]
Keluar: Inject_Agregat_YYYY-MM-DD_sd_YYYY-MM-DD.xlsx + cek stdout (OK/LOLOS).
"""

import argparse
import json
import os
import sys
import time
import urllib.parse

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENV_PATH = "/tmp/eresto_mig.env"
TEMP_HTTP_CODES = {429, 500, 502, 503, 504}

# product bayar eResto -> nama akun ERPNext (per company, via filter API)
PAYMAP = [
    ("cash", "Kas Kasir"),
    ("ewallet", "Qris"),
    ("debit", "Debit Card"),
    ("kredit", "Credit Card"),
    ("transfer", "E-commerce"),
]
INCOME_ACCT = "Penjualan"


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def load_env(path):
    d = {}
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            d[k.strip()] = v.strip().strip('"').strip("'")
    for k in ("ERPNEXT_URL", "ERPNEXT_API_KEY"):
        if not d.get(k):
            fail(f"env {k} kosong di {path}")
    return d


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def style_sheet(ws, money_cols=()):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill, c.font = fill, Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 45)
    for c in money_cols:
        for cell in ws[c][1:]:
            cell.number_format = "#,##0"


class ERP:
    def __init__(self, base, key):
        self.s = requests.Session()
        self.s.headers.update({"Authorization": "token " + key})
        self.base = base
        self.companies = {}   # lower(outlet) -> company doc ringkas
        self.acct_cache = {}  # company -> {account_name: account_fullname}

    def req(self, method, path, payload=None, params=None, timeout=90):
        last = None
        for a in range(1, 4):
            try:
                r = self.s.request(method, self.base + path, json=payload,
                                   params=params, timeout=timeout)
            except requests.RequestException as e:
                last = str(e)[:200]
                time.sleep(2 ** a)
                continue
            if r.status_code in TEMP_HTTP_CODES:
                last = f"HTTP {r.status_code}"
                time.sleep(2 ** a)
                continue
            try:
                return r.status_code, r.json()
            except ValueError:
                return r.status_code, {"_raw": r.text[:1000]}
        fail(f"{method} {path} gagal 3x: {last}")

    def get(self, path, params=None):
        return self.req("GET", path, params=params, timeout=60)

    def load_companies(self):
        st, d = self.get("/api/resource/Company", params={
            "fields": '["name","abbr","cost_center","default_income_account"]',
            "limit_page_length": 1000})
        if st != 200:
            fail(f"list Company HTTP {st}")
        for c in d.get("data", []):
            self.companies[c["name"].strip().lower()] = c

    def company_of(self, outlet):
        return self.companies.get(outlet.strip().lower())

    def accounts(self, company):
        if company not in self.acct_cache:
            st, d = self.get("/api/resource/Account", params={
                "filters": json.dumps([["company", "=", company],
                                       ["is_group", "=", 0]]),
                "fields": '["name","account_name"]',
                "limit_page_length": 500})
            if st != 200:
                fail(f"list Account {company} HTTP {st}")
            self.acct_cache[company] = {a["account_name"]: a["name"]
                                       for a in d.get("data", [])}
        return self.acct_cache[company]

    def existing_je(self, key):
        st, d = self.get("/api/resource/Journal Entry", params={
            "filters": json.dumps([["cheque_no", "=", key]]),
            "fields": '["name","docstatus","total_debit"]',
            "limit_page_length": 5})
        return d.get("data", []) if st == 200 else []

    def delete(self, doctype, name):
        st, _ = self.req("DELETE",
                         f"/api/resource/{doctype}/{urllib.parse.quote(name)}")
        return st

    def post_je(self, payload):
        st, d = self.req("POST", "/api/resource/Journal Entry", payload)
        if st != 200:
            return None, f"POST HTTP {st}: {json.dumps(d)[:300]}"
        return d.get("data", {}).get("name"), ""

    def submit(self, name):
        st, d = self.req(
            "PUT",
            "/api/resource/Journal Entry/" + urllib.parse.quote(name),
            {"docstatus": 1})
        if st != 200:
            return f"SUBMIT HTTP {st}: {json.dumps(d)[:300]}"
        return ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default="pos-fetch/Penjualan_2026-10-01_sd_2026-10-04.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    env = load_env(ENV_PATH)
    erp = ERP(env["ERPNEXT_URL"].rstrip("/"), env["ERPNEXT_API_KEY"])
    print("== 1. master ERP ==")
    erp.load_companies()
    print(f"  companies: {len(erp.companies)}")

    print("== 2. agregat eResto ==")
    raw = pd.read_excel(args.src, sheet_name="Raw_Orders")
    raw["tanggal"] = raw["tanggal"].astype(str).str[:10]
    zero = raw[raw["total"] == 0]
    df = raw[raw["total"] > 0].copy()
    print(f"  {len(raw)} order -> {len(df)} bayar, {len(zero)} void/nol diskip")
    dmin, dmax = df["tanggal"].min(), df["tanggal"].max()
    grp = df.groupby(["outlet_name", "branch", "tanggal"], as_index=False).agg(
        trx=("order_no", "count"), total=("total", "sum"),
        cash=("cash", "sum"), ewallet=("ewallet", "sum"),
        debit=("debit", "sum"), kredit=("kredit", "sum"),
        transfer=("transfer", "sum"))
    print(f"  outlet-hari: {len(grp)} ({dmin}..{dmax})")

    print("== 3. posting Journal Entry ==")
    posted, skipped, failed = [], [], []
    for r in grp.to_dict("records"):
        outlet, branch, tgl = r["outlet_name"], r["branch"], r["tanggal"]
        total = round(num(r["total"]))
        if total <= 0:
            skipped.append({**r, "alasan": "total 0"})
            continue
        co = erp.company_of(outlet)
        if not co:
            failed.append({"outlet": outlet, "tanggal": tgl,
                           "total": total, "alasan": "tanpa company"})
            continue
        company, key = co["name"], f"ERESTO-{branch}-{tgl}"
        accts = erp.accounts(company)
        cc = co.get("cost_center") or f"Main - {co.get('abbr', '')}"
        lines, miss = [], []
        for col, aname in PAYMAP:
            amt = round(num(r[col]))
            if amt <= 0:
                continue
            if aname not in accts:
                miss.append(aname)
            else:
                lines.append({"account": accts[aname],
                              "debit_in_account_currency": amt,
                              "credit_in_account_currency": 0,
                              "cost_center": cc})
        if INCOME_ACCT not in accts:
            miss.append(INCOME_ACCT)
        if miss:
            failed.append({"outlet": outlet, "tanggal": tgl, "total": total,
                           "alasan": f"akun hilang: {','.join(miss)}"})
            continue
        lines.append({"account": accts[INCOME_ACCT],
                      "debit_in_account_currency": 0,
                      "credit_in_account_currency": total,
                      "cost_center": cc})
        bal = sum(x["debit_in_account_currency"]
                  - x["credit_in_account_currency"] for x in lines)
        if abs(bal) > 0.01:
            failed.append({"outlet": outlet, "tanggal": tgl, "total": total,
                           "alasan": f"tidak balance: {bal}"})
            continue
        if args.dry_run:
            posted.append({"key": key, "company": company, "tanggal": tgl,
                           "trx": r["trx"], "total": total, "je": "DRY-RUN"})
            continue
        for old in erp.existing_je(key):
            if old["docstatus"] == 2:
                erp.delete("Journal Entry", old["name"])  # uji yg dibatalkan
            else:
                skipped.append({"outlet": outlet, "branch": branch,
                                "tanggal": tgl, "trx": r["trx"],
                                "total": total,
                                "alasan": f"sudah ada {old['name']}"})
                lines = None
                break
        if lines is None:
            continue
        name, err = erp.post_je({
            "doctype": "Journal Entry", "company": company,
            "posting_date": tgl, "voucher_type": "Journal Entry",
            "cheque_no": key, "cheque_date": tgl,
            "user_remark": (f"ERESTO aggregate {outlet} {tgl}: "
                            f"{r['trx']} trx Rp{total:,.0f}"),
            "accounts": lines})
        if not name:
            failed.append({"outlet": outlet, "tanggal": tgl, "total": total,
                           "alasan": err})
            continue
        err = erp.submit(name)
        if err:
            failed.append({"outlet": outlet, "tanggal": tgl, "total": total,
                           "alasan": f"{name}: {err}"})
        else:
            posted.append({"key": key, "company": company, "tanggal": tgl,
                           "trx": r["trx"], "total": total, "je": name})

    print(f"  posted: {len(posted)} | sudah-ada: {len(skipped)} | "
          f"gagal: {len(failed)}")
    for f in failed[:20]:
        print(f"    GAGAL {f['outlet']} {f['tanggal']} Rp{f['total']:,.0f}: "
              f"{f['alasan']}")

    print("== 4. verifikasi GL vs eResto ==")
    checks = []
    ok_post = len(failed) == 0
    checks.append((f"tanpa gagal posting ({len(failed)} gagal)", ok_post))
    dfp = pd.DataFrame(posted)
    if not dfp.empty:
        keys_posted = set(zip(dfp["company"].str.lower(), dfp["tanggal"]))
        src_posted = grp[grp.apply(
            lambda r: ((erp.company_of(r["outlet_name"]) or {}).get(
                "name", "").lower(), r["tanggal"]) in keys_posted, axis=1)]
        checks.append((f"JE posted ({len(dfp)}) = outlet-hari sumber "
                       f"({len(src_posted)})",
                       len(dfp) == len(src_posted)))
        if not src_posted.empty:
            diff = abs(round(src_posted["total"].sum())
                       - round(dfp["total"].sum()))
            checks.append((f"Rp posted = Rp sumber (selisih {diff:,.0f})",
                           diff < 1))
    st, d = erp.get("/api/resource/Journal Entry", params={
        "filters": json.dumps([["cheque_no", "like", "ERESTO-%"]]),
        "fields": '["name","docstatus"]', "limit_page_length": 1000})
    jes = d.get("data", []) if st == 200 else []
    jes = [j for j in jes if j["docstatus"] != 2]  # abaikan yg dibatalkan
    submitted = [j for j in jes if j["docstatus"] == 1]
    checks.append((f"semua JE ERESTO-* submitted "
                   f"({len(submitted)}/{len(jes)})",
                   len(jes) == len(submitted) and len(jes) > 0))
    all_ok = True
    for name, ok in checks:
        print(f"  [{'LOLOS' if ok else 'GAGAL'}] {name}")
        all_ok = all_ok and ok

    print("== 5. Excel ==")
    by_day = df.groupby("tanggal").agg(
        trx=("order_no", "count"), outlet=("outlet_name", "nunique"),
        omzet=("total", "sum"), cash=("cash", "sum"),
        ewallet=("ewallet", "sum")).reset_index()
    by_outlet = grp.copy()
    by_outlet["company"] = by_outlet["outlet_name"].apply(
        lambda o: (erp.company_of(o) or {}).get("name", "TANPA COMPANY"))
    out = f"Inject_Agregat_{dmin}_sd_{dmax}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        by_day.to_excel(w, sheet_name="Rekap_Harian", index=False)
        by_outlet.to_excel(w, sheet_name="Rekap_Outlet", index=False)
        pd.DataFrame(posted).to_excel(w, sheet_name="Raw_Posted", index=False)
        pd.DataFrame([{"cek": n, "hasil": "LOLOS" if ok else "GAGAL"}
                      for n, ok in checks]).to_excel(
            w, sheet_name="Validation", index=False)
        pd.DataFrame(failed).to_excel(w, sheet_name="Failed", index=False)
        style_sheet(w.sheets["Rekap_Harian"], ("D", "E", "F"))
        style_sheet(w.sheets["Rekap_Outlet"], ("E", "F", "G", "H", "I", "J"))
        style_sheet(w.sheets["Raw_Posted"], ("E",))
        style_sheet(w.sheets["Validation"])
        style_sheet(w.sheets["Failed"])
    print(f"  Excel: {out}")
    if not all_ok or failed:
        fail("ada cek GAGAL — lihat sheet Validation/Failed")
    print("INJEKSI OK")


if __name__ == "__main__":
    main()
