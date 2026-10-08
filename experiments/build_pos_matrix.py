"""Bangun matriks POS setup: tiap company -> tier takeaway eResto ->
Price List ERPNext + status kebutuhan POS Profile.
Sumber lokal saja (tanpa API):
- experiments/Company_vs_PriceList.xlsx (163 company lama + tier)
- experiments/Outlet_vs_PriceGroup.xlsx (tier per outlet eResto)
- experiments/Kandidat_Company_Baru.xlsx (tier Anna)
Output: experiments/Company_POS_Setup.xlsx.
Aturan nama Price List:
- Tier mulai "MENU ROPI" -> "PG-" + Title Case (ROPI tetap caps),
  cocok dgn yg sudah ada: PG-Menu ROPI 10K 15 Juli 2026.
- Tier lain -> "PG-" + nama eResto apa adanya (akronim aman)."""
import pandas as pd

BARU = {
    # company: (abbr, tier_takeaway_mentah)
    "RSUD Hj. ANNA LASMANAH BANJARNEGARA": ("RHALB", "MENU ROPI 10K 15 JULI 2026"),
    "RSUP Palembang": ("RSUPP", "RSUP Palembang"),
    "RSU San Medica (FRC)": ("RSSMF", "ROPI SAN MEDICA CENTER"),
    "Bangka Tengah (FRC)": ("BTF", "ROPI BANGKA TENGAH FRC"),
    "Plaza Haji alQosbah (FRC)": ("PHAQF", "ROPI 12K MARET 2026"),
    "RS Medina Garut (FRC)": ("RMGF", "ROPI RS MEDINA GARUT"),
    "Jambi Sarulangun (FRC)": ("JSF", "JAMBI SOROLANGUN FRC"),
    "Sungai Penuh (FRC)": ("SPF", "ROPI SUNGAI PENUH"),
    "Mitra Swalayan Gombong (FRC)": ("MSGF", "ROPI 10 K"),
}

PRICE_LIST_SUDAH_ADA = {"PG-Menu ROPI 10K 15 Juli 2026"}
POS_PROFILE_SUDAH_ADA = {"RS PKU MUHAMMADIYAH DELANGGU"}


def price_list_name(tier):
    if not tier:
        return ""
    if tier.startswith("MENU ROPI"):
        titled = tier.title().replace("Ropi", "ROPI")
        return "PG-" + titled
    return "PG-" + tier


def main():
    c = pd.read_excel("experiments/Company_vs_PriceList.xlsx")
    o = pd.read_excel("experiments/Outlet_vs_PriceGroup.xlsx",
                      sheet_name="Outlet_PriceGroup")
    # tier bersih per outlet eResto (buang bagian kurir).
    # key dinormalisasi (lower) karena beda case:
    # "RSUD Hj. ANNA..." (ERPNext) vs "RSUD HJ. ANNA..." (eResto).
    tier_of = {}
    div_of = {}
    for r in o.to_dict("records"):
        t = str(r.get("price_group_takeaway") or "").split(" + ")[0].strip()
        if t and t.lower() != "nan":
            key = str(r["outlet_eresto"]).strip().lower()
            tier_of[key] = t
            div_of.setdefault(t, r.get("division_id"))
    # fallback tier dari Kandidat (juga dinormalisasi)
    try:
        k = pd.read_excel("experiments/Kandidat_Company_Baru.xlsx",
                          sheet_name="Kandidat")
        for r in k.to_dict("records"):
            t = str(r.get("price_group_takeaway") or "").strip()
            if t and t.lower() != "nan":
                tier_of.setdefault(str(r["outlet"]).strip().lower(), t)
    except Exception:
        pass

    rows = []
    for r in c.to_dict("records"):
        co = str(r["company"])
        tier = str(r["tier_takeaway"]) if pd.notna(r["tier_takeaway"]) else ""
        tier = "" if tier.lower() == "nan" else tier
        sumber = "Company_vs_PriceList"
        if not tier:
            # fallback: cocokkan nama company ke outlet eResto (case-insensitive)
            t2 = tier_of.get(co.strip().lower(), "")
            if t2:
                tier, sumber = t2, "Outlet_vs_PriceGroup (case-insensitive)"
        rows.append({"company": co, "abbr": r["abbr"],
                     "tier_takeaway": tier,
                     "sumber_tier": sumber})
    have = {r["company"] for r in rows}
    for co, (abbr, tier) in BARU.items():
        if co not in have:
            rows.append({"company": co, "abbr": abbr, "tier_takeaway": tier,
                         "sumber_tier": "kandidat/eResto-langsung"})

    for r in rows:
        tier = r["tier_takeaway"]
        r["price_list"] = price_list_name(tier)
        r["price_list_status"] = ("SUDAH ADA" if r["price_list"] in
                                  PRICE_LIST_SUDAH_ADA
                                  else ("BARU" if r["price_list"] else
                                        "TANPA TIER"))
        r["pos_profile_status"] = ("SUDAH ADA (2 shift)" if r["company"] in
                                    POS_PROFILE_SUDAH_ADA else "BUTUH 2 shift")
        r["division_contoh"] = div_of.get(tier, "")

    df = pd.DataFrame(rows)
    kebutuhan = (df[df["price_list"].ne("")]
                 .groupby(["price_list", "tier_takeaway"], as_index=False)
                 .agg(jumlah_company=("company", "count"),
                      status=("price_list_status", "first"),
                      contoh_company=("company", "first"),
                      division_contoh=("division_contoh", "first"))
                 .sort_values("jumlah_company", ascending=False))

    with pd.ExcelWriter("experiments/Company_POS_Setup.xlsx",
                        engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Company_Tier", index=False)
        kebutuhan.to_excel(w, sheet_name="Kebutuhan_PriceList", index=False)
        for ws in w.sheets.values():
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Alignment, Font, PatternFill
            fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.fill, cell.font = fill, Font(bold=True)
                cell.alignment = Alignment(horizontal="center",
                                           vertical="center")
    print(f"company: {len(df)} | price list unik: {df['price_list'].nunique()}")
    print(kebutuhan.to_string())
    tanpa = df[df["price_list"] == ""]
    print(f"tanpa tier: {len(tanpa)} -> {tanpa['company'].tolist()}")


if __name__ == "__main__":
    main()
