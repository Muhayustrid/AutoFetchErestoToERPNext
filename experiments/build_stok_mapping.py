"""Bangun workbench mapping SKU inventory eResto -> Item ERPNext (stok SO).

Sumber SKU: experiments/stock_so_monthly.json (SO Monthly September,
semua outlet). Keputusan user 2026-10-04:
- RR0001..dst (barang jadi) -> PJO* matang (saldo stok); RP00x/CR001/GM001
  tetap item jual (bundle non-stok).
- BB0085 Krim Kopi -> PJ260016.
- Bubuk tanpa padanan -> item baru (sudah dibuat: BB260057-62).
- Sisa (Olesan Blueberry/Sarikaya, RR0004, RR0010, Dough Srikaya/
  Tiramisu/Vanilla) -> SKIP dulu.
Output: experiments/stok_sku_to_item.csv + Template_Mapping_Stok.xlsx.
Env: ERPNEXT_URL, ERPNEXT_API_KEY (utk baca nama/UOM item target).
"""
import csv
import json
import os
import urllib.parse

import requests

ENV_PATH = "/tmp/eresto_mig.env"
SRC = "experiments/stock_so_monthly.json"
OUT_CSV = "experiments/stok_sku_to_item.csv"
OUT_XLSX = "experiments/Template_Mapping_Stok.xlsx"


def load_env(path):
    d = {}
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            d[k.strip()] = v.strip().strip('"').strip("'")
    return d


ENV = load_env(ENV_PATH)
BASE = ENV["ERPNEXT_URL"].rstrip("/")
S = requests.Session()
S.headers.update({"Authorization": "token " + ENV["ERPNEXT_API_KEY"]})

# sku -> (item_code, status, faktor, catatan)
MAP = {
    # --- FIX: keputusan user ---
    "BB0046": ("PJ260012", "FIX", "1 Pack = 1 Pcs", "Dough Butter -> Dough Ropi Butter (pabrik)"),
    "BB0085": ("PJ260016", "FIX", "1:1", "Krim Kopi (pilih PJ260016, bukan PJ260008 700gr)"),
    "RR0001": ("PJO2603", "FIX", "1:1", "Ropi Butter matang; jual = RP001 (bundle)"),
    "RR0002": ("PJO2601", "FIX", "1:1", "Ropi Coklat matang; jual = RP002"),
    "RR0003": ("PJO2602", "FIX", "1:1", "Ropi Keju matang; jual = RP003"),
    "RR0008": ("PJO2605", "FIX", "1:1", "Croissant matang; jual = CR001"),
    "RR0009": ("PJO2604", "FIX", "1:1", "Gempi matang; jual = GM001"),
    # --- BARU: item dibuat 2026-10-04 (create_stok_items.py) ---
    "BB0308": ("BB260057", "BARU", "1:1", "Item baru Bubuk Hepi Drink Moccacino (30 g)"),
    "BB0309": ("BB260058", "BARU", "1:1", "Item baru Bubuk Hepi Drink Cappucino (30 g)"),
    "BB0310": ("BB260059", "BARU", "1:1", "Item baru Bubuk Hepi Drink Kopi Karamel (30 g)"),
    "BB0311": ("BB260060", "BARU", "1:1", "Item baru Bubuk Hepi Drink Chocolate"),
    "BB0322": ("BB260061", "BARU", "1:1", "Item baru Kopi Hitam @6gr (Isi 20 Sachet)"),
    "BB0145": ("BB260062", "BARU", "x1000", "Item baru Susu Bubuk Campuran Minuman (1 kg; kg->g)"),
    # --- COCOK: nama/pola ---
    "BB0047": ("PJ260013", "COCOK", "1 Pack = 1 Pcs", "Dough Ropi Coklat"),
    "BB0049": ("PJ260014", "COCOK", "1 Pack = 1 Pcs", "Dough Ropi Keju"),
    "BB0168": ("PJ260010", "COCOK", "1 Pack = 1 Pcs", "Dough Croissant"),
    "BB0261": ("PJ260011", "COCOK", "1 Pack = 1 Pcs", "Dough Gempi"),
    "BB0071": ("BB260039", "COCOK", "1:1", "Gula Stik"),
    "BB0080": ("PEL260074", "COCOK", "1:1", "Kardus Ropi 20 roti (25 pcs)"),
    "BB0102": ("BB260043", "COCOK", "x1000", "Olesan Choco Crunchy (1 kg; kg->g)"),
    "BB0103": ("BB260044", "COCOK", "x1000", "Olesan Dark Choco (1 kg; kg->g)"),
    "BB0110": ("BB260047", "COCOK", "x1000", "Olesan Tiramisu (1 kg; kg->g)"),
    "BB0112": ("PEL260076", "COCOK", "1:1", "Paper Bag Ropi"),
    "BB0117": ("PEL260077", "COCOK", "1:1", "Plastik Alas Gempi"),
    "BB0171": ("PEL260069", "COCOK", "1:1", "Box Gempi (100 pcs)"),
    "BB0257": ("PEL260071", "COCOK", "1:1", "Gelas Cup (9 oz) (50 pcs)"),
    "BB0298": ("PEL260070", "COCOK", "1:1", "Gelas (14 oz); beda label isi pack di eResto"),
    "BB0301": ("BB260034", "COCOK", "1 Pcs = 1 Sachet", "Bubuk Beras Kencur"),
    "BB0305": ("BB260038", "COCOK", "1:1", "Gula Cair Sachet"),
    "BB0312": ("PJ260009", "COCOK", "1:1", "Air Mineral Outlet (600 ml)"),
    "BB0329": ("BB260030", "COCOK", "1 Pack = 750 g", "Jasmine Green Tea"),
    "BB0340": ("BB260031", "COCOK", "1 Pack = 500 g", "Coffee Latte"),
    "BB0341": ("BB260032", "COCOK", "1 Pack = 120 g", "Coffee Classic"),
    "BB0342": ("BB260033", "COCOK", "1 Pack = 1000 g", "Dark Cocoa Drink"),
    "BB0345": ("BB260050", "COCOK", "1 Pack = 450 g", "Coffeemate"),
    "BB0346": ("BB260052", "COCOK", "1 Pack = 960 g", "Milo Activ"),
    "BB0347": ("BB260054", "COCOK", "1 Pack = 960 g", "Teh Tarik"),
    "BB0349": ("BB260051", "COCOK", "1 Pack = 640 g", "Lemonade"),
    # --- keputusan user 2026-10-09: Goodie Bag = PEL260072 (PG001 legacy di-merge) ---
    "BB0066": ("PEL260072", "COCOK", "1:1", "Goodie Bag = PEL260072 'Goodie Bag' (keputusan user 2026-10-09; PG001 di-merge)"),
    # --- SKIP: keputusan user ---
    "BB0052": (None, "SKIP", "", "Dough Srikaya: tidak ada padanan, stok 0"),
    "BB0053": (None, "SKIP", "", "Dough Tiramisu: tidak ada padanan, stok 0"),
    "BB0054": (None, "SKIP", "", "Dough Vanilla: tidak ada padanan, stok 0"),
    "BB0101": (None, "SKIP", "", "Olesan Blueberry: tidak ada di master (skip dulu)"),
    "BB0106": (None, "SKIP", "", "TIDAK AKTIF di eResto"),
    "BB0107": (None, "SKIP", "", "Olesan Sarikaya: tidak ada di master (skip dulu)"),
    "BB0108": (None, "SKIP", "", "TIDAK AKTIF di eResto"),
    "RR0004": (None, "SKIP", "", "Ropi Tiramisu: tidak ada item (skip dulu)"),
    "RR0010": (None, "SKIP", "", "Teh Enak Seduh: tidak ada item (skip dulu)"),
}


def fetch_item(code):
    r = S.get(f"{BASE}/api/resource/Item/{urllib.parse.quote(code)}",
              timeout=60)
    if r.status_code != 200:
        return None
    d = r.json().get("data", {})
    return {"name": d.get("item_name"), "uom": d.get("stock_uom"),
            "group": d.get("item_group"), "is_stock": d.get("is_stock_item")}


def main():
    data = json.load(open(SRC))
    agg = {}
    for r in data["rows"]:
        a = agg.setdefault(r["sku"], {"name": r["item_name"],
                                      "cat": r["category"] or "",
                                      "unit": r["unit"] or "",
                                      "apps": 0, "qty": 0.0})
        a["apps"] += 1
        a["qty"] += (r["real_qty"] or 0)
    missing = [s for s in agg if s not in MAP]
    if missing:
        print("PERINGATAN: SKU tanpa mapping:", missing)

    cache, rows = {}, []
    for sku in sorted(agg):
        a = agg[sku]
        code, status, faktor, catatan = MAP.get(
            sku, (None, "TANYA", "", "belum diputuskan"))
        info = None
        if code:
            if code not in cache:
                cache[code] = fetch_item(code)
            info = cache[code]
            if info is None:
                print(f"  PERINGATAN: item {code} tidak ada di ERPNext!")
        rows.append({
            "sku_eresto": sku, "nama_eresto": a["name"],
            "kategori_eresto": a["cat"], "unit_eresto": a["unit"],
            "outlet": a["apps"], "total_real_qty": round(a["qty"], 3),
            "item_code_erpnext": code or "",
            "item_name_erpnext": (info or {}).get("name", ""),
            "uom_erpnext": (info or {}).get("uom", ""),
            "faktor_konversi": faktor, "status": status, "catatan": catatan,
        })

    with open(OUT_CSV, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)

    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    df = pd.DataFrame(rows)

    def style(ws):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        fill = PatternFill("solid", fgColor="D9EAF7")
        for c in ws[1]:
            c.fill, c.font = fill, Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        for col in ws.columns:
            w_ = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(w_ + 2, 48)

    rekap = (df.groupby("status")
               .agg(sku=("sku_eresto", "count"),
                    outlets=("outlet", "sum"),
                    total_real=("total_real_qty", "sum"))
               .reset_index().sort_values("status"))
    keputusan = pd.DataFrame([
        {"topik": "RR* (barang jadi)", "keputusan": "Saldo stok -> PJO* matang (PJO2603/2601/2602/2605/2604). RP001/RP002/RP003/CR001/GM001 tetap item jual (bundle, non-stok)."},
        {"topik": "BB0085 Krim Kopi", "keputusan": "-> PJ260016 (bukan PJ260008 700gr)."},
        {"topik": "Bubuk tanpa padanan", "keputusan": "Dibuat item baru BB260057-62 (BB0308/09/10/11/BB0322/BB0145)."},
        {"topik": "Sisa tanpa padanan", "keputusan": "SKIP dulu: Olesan Blueberry/Sarikaya, RR0004, RR0010, Dough Srikaya/Tiramisu/Vanilla, item TIDAK AKTIF."},
        {"topik": "Konversi satuan", "keputusan": "Stok SO disimpan dalam stock_uom ERPNext: Pack->g pakai gramasi nama, kg->g x1000, Pack dough->Pcs 1:1."},
        {"topik": "Goodie Bag (dulu terbuka)", "keputusan": "BB0066 Goodie Bag -> PEL260072 'Goodie Bag' (keputusan user 2026-10-09). Item legacy PG001 'Goodie Bag 6 Roti' sudah DI-MERGE ke PEL260072 & dihapus; jangan pakai PG001."},
    ])
    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Mapping", index=False)
        rekap.to_excel(w, sheet_name="Rekap", index=False)
        keputusan.to_excel(w, sheet_name="Keputusan", index=False)
        df[df.status == "SKIP"][["sku_eresto", "nama_eresto", "catatan"]] \
            .to_excel(w, sheet_name="Sisa_Skip", index=False)
        for ws in w.sheets.values():
            style(ws)

    mapped = sum(1 for r in rows if r["item_code_erpnext"])
    skipped = sum(1 for r in rows if r["status"] == "SKIP")
    print(f"SKU total: {len(rows)} | ada target: {mapped} | skip: {skipped}")
    print(rekap.to_string(index=False))
    print(f"CSV : {OUT_CSV}")
    print(f"XLSX: {OUT_XLSX}")


if __name__ == "__main__":
    main()
