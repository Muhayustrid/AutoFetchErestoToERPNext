# SPIKE sekali pakai (lanjutan): tarik detail /v3/orders/{id} utk semua order
# di Excel v1 -> sheet per-baris-item + validasi silang Sigma(item) vs total.
# Bukan kode produksi. Kredensial WAJIB via env.
import argparse
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://ropi.eresto.id"
MAX_WORKERS = 10
MAX_RETRIES = 3
TIMEOUT = 90
TEMP_HTTP_CODES = {429, 500, 502, 503, 504}
SRC_XLSX = "POS_Orders_045_2026-10-03.xlsx"


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def get_with_retry(session, url, params):
    last = None
    for att in range(1, MAX_RETRIES + 1):
        try:
            r = session.get(url, params=params, timeout=TIMEOUT)
            if r.status_code in TEMP_HTTP_CODES:
                last = f"HTTP {r.status_code}"
                time.sleep(2 ** att)
                continue
            r.raise_for_status()
            return r.json()
        except (requests.RequestException, ValueError) as e:
            last = str(e)[:150]
            time.sleep(2 ** att)
    return {"_error": last}


def fetch_detail(session, token, user_id, order_id):
    url = f"{BASE_URL}/v3/orders/{order_id}"
    payload = get_with_retry(session, url, {
        "domain": "ropi", "token": token,
        "current_user_id": user_id, "_": int(time.time() * 1000),
    })
    if "_error" in payload:
        return order_id, None, payload["_error"]
    try:
        return order_id, payload["data"]["attributes"], ""
    except (KeyError, TypeError):
        return order_id, None, "payload tanpa data.attributes"


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill, c.font = fill, Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 45)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=SRC_XLSX)
    src = ap.parse_args().src
    token = os.getenv("ERESTO_TOKEN", "").strip()
    email = os.getenv("ERESTO_EMAIL", "").strip()
    user_id = os.getenv("ERESTO_USER_ID", "").strip()
    if not (token and email and user_id):
        fail("env ERESTO_TOKEN/EMAIL/USER_ID kurang")
    orders = pd.read_excel(src, sheet_name="Raw_Orders").to_dict("records")
    print(f"{len(orders)} order dari {src}")

    s = requests.Session()
    s.headers.update({
        "accept": "application/json, text/plain, */*",
        "authorization": f"Token token={token}, email={email}",
        "referer": "https://ropi.eresto.id/manager/v2/index.html",
        "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "cache-control": "no-cache", "pragma": "no-cache",
    })

    details, failed = {}, []
    t0 = time.time()
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        futs = {pool.submit(fetch_detail, s, token, user_id, o["order_id"]): o
                for o in orders}
        done = 0
        for f in as_completed(futs):
            oid, attrs, err = f.result()
            done += 1
            print(f"\r  {done}/{len(orders)}", end="", flush=True)
            if err:
                failed.append({"order_id": oid,
                               "order_no": futs[f]["order_no"], "error": err})
            else:
                details[oid] = attrs
    print(f"\nOK: {len(details)}, gagal: {len(failed)}, "
          f"{time.time()-t0:.1f}s")

    item_rows, recon = [], []
    for o in orders:
        oid, ono = o["order_id"], o["order_no"]
        attrs = details.get(oid)
        if attrs is None:
            recon.append({"order_no": ono, "header_total": o["total"],
                          "sum_item": "", "selisih": "",
                          "status": "DETAIL GAGAL"})
            continue
        items = attrs.get("order_items") or []
        subtotal = 0.0
        if not items:
            recon.append({"order_no": ono, "header_total": o["total"],
                          "sum_item": 0.0, "selisih": round(o["total"], 2),
                          "status": "TANPA ITEM"})
        for it in items:
            qty, price = num(it.get("quantity")), num(it.get("price"))
            disc = num(it.get("discount_item_amount"))
            line = round(qty * price - disc, 2)
            subtotal += line
            item_rows.append({
                "order_no": ono, "order_id": oid,
                "product_name": it.get("product_name", ""),
                "product_sku": it.get("product_sku", ""),
                "product_id": it.get("product_id", ""),
                "quantity": qty, "price": price,
                "discount_item_amount": disc,
                "tax_amount": num(it.get("tax_amount")),
                "line_amount": line,
                "void": bool(it.get("void")),
                "void_quantity": num(it.get("void_quantity")),
                "paid_quantity": num(it.get("paid_quantity")),
            })
        if items:
            diff = round(o["total"] - subtotal, 2)
            recon.append({
                "order_no": ono, "header_total": o["total"],
                "sum_item": round(subtotal, 2), "selisih": diff,
                "status": "COCOK" if abs(diff) < 0.01 else "SELISIH",
            })

    out = src.replace(".xlsx", "_v2_detail.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        pd.DataFrame(item_rows).to_excel(w, sheet_name="Order_Items", index=False)
        pd.DataFrame(recon).to_excel(w, sheet_name="Rekonsiliasi", index=False)
        pd.DataFrame(failed).to_excel(w, sheet_name="Failed", index=False)
        for ws in w.sheets.values():
            style_sheet(ws)
    print(f"Excel: {out}")

    n_cocok = sum(1 for r in recon if r["status"] == "COCOK")
    n_selisih = [r for r in recon if r["status"] == "SELISIH"]
    print(f"Rekonsiliasi: {n_cocok}/{len(recon)} COCOK, "
          f"{len(n_selisih)} SELISIH, "
          f"{sum(1 for r in recon if r['status']!='COCOK' and r['status']!='SELISIH')} tanpa-item/gagal")
    for r in n_selisih[:10]:
        print(f"  SELISIH {r['order_no']}: header={r['header_total']:g} "
              f"sum={r['sum_item']:g} diff={r['selisih']:g}")
    skus = {}
    for r in item_rows:
        skus.setdefault(r["product_sku"], set()).add(r["product_name"])
    ganda = {k: v for k, v in skus.items() if len(v) > 1}
    print(f"SKU unik: {len(skus)}, SKU dgn >1 nama: {len(ganda)}")
    for k, v in list(ganda.items())[:10]:
        print(f"  {k}: {sorted(v)}")


if __name__ == "__main__":
    main()
