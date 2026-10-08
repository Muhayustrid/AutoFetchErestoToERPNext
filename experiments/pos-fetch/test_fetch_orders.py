# SPIKE sekali pakai: test tarik /v3/orders eResto -> Excel verifikasi.
# Bukan kode produksi. Kode permanen pos_sync/ dibuat setelah hasil disetujui.
# Kredensial WAJIB via env: ERESTO_TOKEN, ERESTO_EMAIL, ERESTO_USER_ID.
import argparse
import os
import sys
import time

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://ropi.eresto.id"
DOMAIN = "ropi"
PER_PAGE = 200
MAX_RETRIES = 3
TIMEOUT = 90
TEMP_HTTP_CODES = {429, 500, 502, 503, 504}


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def load_creds():
    token = os.getenv("ERESTO_TOKEN", "").strip()
    email = os.getenv("ERESTO_EMAIL", "").strip()
    user_id = os.getenv("ERESTO_USER_ID", "").strip()
    missing = [n for n, v in
               (("ERESTO_TOKEN", token), ("ERESTO_EMAIL", email),
                ("ERESTO_USER_ID", user_id)) if not v]
    if missing:
        fail(f"env kurang: {', '.join(missing)}")
    return token, email, user_id


def create_session(token, email):
    s = requests.Session()
    s.headers.update({
        "accept": "application/json, text/plain, */*",
        "authorization": f"Token token={token}, email={email}",
        "referer": "https://ropi.eresto.id/manager/v2/index.html",
        "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "cache-control": "no-cache",
        "pragma": "no-cache",
    })
    return s


def request_json(session, url, params):
    last_err = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            r = session.get(url, params=params, timeout=TIMEOUT)
            if r.status_code in TEMP_HTTP_CODES:
                last_err = f"HTTP {r.status_code}"
                time.sleep(2 ** attempt)
                continue
            r.raise_for_status()
            return r.json()
        except (requests.RequestException, ValueError) as e:
            last_err = str(e)[:200]
            time.sleep(2 ** attempt)
    fail(f"GET {url} gagal {MAX_RETRIES}x: {last_err}")


def fetch_orders(session, token, user_id, date, outlet_id):
    url = f"{BASE_URL}/v3/orders"
    orders, meta_last, page = [], {}, 1
    while True:
        params = {
            "current_user_id": user_id,
            "date[geq]": f"{date} 00:00:00.000",
            "date[leq]": f"{date} 23:59:59.999",
            "domain": DOMAIN,
            "outlet_id": outlet_id,
            "page": page,
            "per_page": PER_PAGE,
            "sort[property]": "date",
            "sort[order]": "asc",
            "token": token,
            "_": int(time.time() * 1000),
        }
        payload = request_json(session, url, params)
        batch = payload.get("data", [])
        meta_last = payload.get("meta", {})
        print(f"  page {page}: {len(batch)} order")
        orders.extend(batch)
        if len(batch) < PER_PAGE or not meta_last.get("next_page"):
            break
        page += 1
        if page > 500:
            fail("loop paginasi > 500 page, hentikan")
    seen, uniq = set(), []
    for o in orders:
        oid = o.get("id")
        if oid and oid not in seen:
            seen.add(oid)
            uniq.append(o)
    return uniq, meta_last


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def flatten(order):
    a = order.get("attributes", {}) or {}
    p = a.get("payment", {}) or {}
    return {
        "order_id": order.get("id", ""),
        "order_no": a.get("order_no", ""),
        "receipt_number": p.get("receipt_number", ""),
        "date": str(a.get("date", "")),
        "local_time": a.get("local_time", ""),
        "outlet_id": a.get("outlet_id", ""),
        "outlet_name": a.get("outlet_name", ""),
        "branch_number": a.get("outlet_branch_number", ""),
        "cashier": a.get("cashier_name", ""),
        "category": a.get("category", ""),
        "source": a.get("source", ""),
        "queue_number": a.get("queue_number", ""),
        "transaction_group_id": a.get("transaction_group_id", ""),
        "item_quantity": num(a.get("item_quantity")),
        "void_quantity": num(a.get("void_quantity")),
        "total": num(p.get("total")),
        "cash_amount": num(p.get("cash_amount")),
        "ewallet_amount": num(p.get("ewallet_amount")),
        "credit_amount": num(p.get("credit_amount")),
        "debit_amount": num(p.get("debit_amount")),
        "transfer_amount": num(p.get("transfer_amount")),
    }


def validate(rows, meta, date, outlet_id):
    checks = []
    checks.append(("HTTP 200 + ada data/meta",
                   bool(rows) and isinstance(meta, dict)))
    checks.append(("semua outlet_id = filter",
                   all(r["outlet_id"] == outlet_id for r in rows)))
    checks.append(("semua date dalam rentang",
                   all(r["date"][:10] == date for r in rows)))
    checks.append(("payment.total numerik >= 0",
                   all(r["total"] >= 0 for r in rows)))
    ids = [r["order_id"] for r in rows]
    nos = [r["order_no"] for r in rows]
    checks.append(("tanpa duplikat id/order_no",
                   len(set(ids)) == len(ids)
                   and len(set(nos)) == len(nos)))
    total_count = meta.get("total_count")
    checks.append((f"jumlah baris ({len(rows)}) = meta.total_count ({total_count})",
                   total_count == len(rows)))
    return checks


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill, c.font = fill, Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 45)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", required=True, help="YYYY-MM-DD")
    ap.add_argument("--outlet-id", required=True)
    args = ap.parse_args()

    token, email, user_id = load_creds()
    session = create_session(token, email)
    print(f"Fetch {args.date} outlet {args.outlet_id} ...")
    orders, meta = fetch_orders(session, token, user_id,
                               args.date, args.outlet_id)
    rows = [flatten(o) for o in orders]
    print(f"Terkumpul {len(rows)} order unik")

    checks = validate(rows, meta, args.date, args.outlet_id)
    print("\nHasil validasi:")
    all_ok = True
    for name, ok in checks:
        print(f"  [{'OK' if ok else 'GAGAL'}] {name}")
        all_ok = all_ok and ok

    branch = (rows[0]["branch_number"] or "OUTLET") if rows else "EMPTY"
    safe_branch = "".join(c if c.isalnum() else "_" for c in branch)
    fname = f"POS_Orders_{safe_branch}_{args.date}.xlsx"
    with pd.ExcelWriter(fname, engine="openpyxl") as w:
        pd.DataFrame(rows).to_excel(w, sheet_name="Raw_Orders", index=False)
        pd.DataFrame([{"param": k, "value": v}
                      for k, v in {"date": args.date,
                                   "outlet_id": args.outlet_id,
                                   "per_page": PER_PAGE,
                                   **meta}.items()]
                     ).to_excel(w, sheet_name="Meta", index=False)
        pd.DataFrame([{"cek": n, "hasil": "LOLOS" if ok else "GAGAL"}
                      for n, ok in checks]
                     ).to_excel(w, sheet_name="Validation", index=False)
        for ws in w.sheets.values():
            style_sheet(ws)
    print(f"\nExcel: {fname}")
    print("Contoh untuk cek silang di UI sales_report/orders/orders:")
    for r in rows[:3]:
        print(f"  {r['order_no']} | {r['outlet_name']} | "
              f"{r['local_time']} | total={r['total']:g}")
    if not all_ok:
        fail("ada cek validasi GAGAL — jangan lanjut ke ERPNext")


if __name__ == "__main__":
    main()
