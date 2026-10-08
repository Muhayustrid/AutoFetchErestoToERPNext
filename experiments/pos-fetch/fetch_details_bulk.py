# SPIKE sekali pakai: detail /v3/orders/{id} utk SEMUA order 1-4 Okt 2026
# -> Excel per-baris-SKU + join mapping sku_to_item_code.csv.
# Cache resume di .cache_details/ (list JSON + 1 JSON per order_id);
# rerun hanya fetch yg belum ada. Kredensial WAJIB via env.
import argparse
import csv
import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://ropi.eresto.id"
DOMAIN = "ropi"
PER_PAGE = 200
MAX_WORKERS = 10
MAX_RETRIES = 3
TIMEOUT = 90
TEMP_HTTP_CODES = {429, 500, 502, 503, 504}
HERE = Path(__file__).resolve().parent
CACHE = HERE / ".cache_details"


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def creds():
    t = os.getenv("ERESTO_TOKEN", "").strip()
    e = os.getenv("ERESTO_EMAIL", "").strip()
    u = os.getenv("ERESTO_USER_ID", "").strip()
    if not (t and e and u):
        fail("env ERESTO_TOKEN/EMAIL/USER_ID kurang")
    return t, e, u


def session(token, email):
    s = requests.Session()
    s.headers.update({
        "accept": "application/json, text/plain, */*",
        "authorization": f"Token token={token}, email={email}",
        "referer": "https://ropi.eresto.id/manager/v2/index.html",
        "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "cache-control": "no-cache", "pragma": "no-cache",
    })
    return s


def getj(s, url, params):
    last = None
    for a in range(1, MAX_RETRIES + 1):
        try:
            r = s.get(url, params=params, timeout=TIMEOUT)
            if r.status_code in TEMP_HTTP_CODES:
                last = f"HTTP {r.status_code}"
                time.sleep(2 ** a)
                continue
            r.raise_for_status()
            return r.json(), ""
        except (requests.RequestException, ValueError) as e:
            last = str(e)[:150]
            time.sleep(2 ** a)
    return None, last


def fetch_list(s, token, user_id, dfrom, dto):
    orders, page = [], 1
    while True:
        j, err = getj(s, f"{BASE_URL}/v3/orders", {
            "current_user_id": user_id,
            "date[geq]": f"{dfrom} 00:00:00.000",
            "date[leq]": f"{dto} 23:59:59.999",
            "domain": DOMAIN, "outlet_id": "",
            "page": page, "per_page": PER_PAGE,
            "sort[property]": "date", "sort[order]": "asc",
            "token": token, "_": int(time.time() * 1000)})
        if err:
            fail(f"list page {page}: {err}")
        batch = j.get("data", [])
        meta = j.get("meta", {})
        orders.extend(batch)
        print(f"\r  list page {page}: +{len(batch)} (total {len(orders)})",
              end="", flush=True)
        if len(batch) < PER_PAGE or not meta.get("next_page"):
            break
        page += 1
    print()
    seen, uniq = set(), []
    for o in orders:
        if o.get("id") and o["id"] not in seen:
            seen.add(o["id"])
            uniq.append(o)
    return uniq


def flat_header(o):
    a = o.get("attributes", {}) or {}
    p = a.get("payment", {}) or {}
    return {
        "order_id": o.get("id", ""), "order_no": a.get("order_no", ""),
        "tanggal": str(a.get("date", ""))[:10],
        "outlet_name": a.get("outlet_name", ""),
        "branch": a.get("outlet_branch_number", ""),
        "kasir": a.get("cashier_name", ""), "kategori": a.get("category", ""),
        "total": num(p.get("total")), "cash": num(p.get("cash_amount")),
        "ewallet": num(p.get("ewallet_amount")),
        "transfer": num(p.get("transfer_amount")),
        "debit": num(p.get("debit_amount")),
        "kredit": num(p.get("credit_amount")),
    }


def fetch_detail(s, token, user_id, oid):
    p = CACHE / f"{oid}.json"
    if p.exists():
        try:
            return oid, json.loads(p.read_text()), ""
        except ValueError:
            p.unlink()
    j, err = getj(s, f"{BASE_URL}/v3/orders/{oid}", {
        "domain": DOMAIN, "token": token,
        "current_user_id": user_id, "_": int(time.time() * 1000)})
    if err:
        return oid, None, err
    root = j.get("data") if isinstance(j.get("data"), dict) else j
    attrs = (root or {}).get("attributes")
    if not isinstance(attrs, dict):
        return oid, None, "payload tanpa data.attributes"
    p.write_text(json.dumps(attrs))
    return oid, attrs, ""


def attrs_items(attrs):
    items = attrs.get("order_items") or []
    rows = []
    for it in items:
        pq = num(it.get("paid_quantity"))
        rows.append({
            "product_sku": it.get("product_sku", "") or "",
            "product_name": it.get("product_name", "") or "",
            "quantity": num(it.get("quantity")),
            "paid_quantity": pq,
            "price": num(it.get("price")),
            "line_paid": round(pq * num(it.get("price")), 2),
            "discount_item_amount": num(it.get("discount_item_amount")),
            "void": bool(it.get("void")),
            "void_quantity": num(it.get("void_quantity")),
            "void_note": it.get("void_note", "") or "",
        })
    return rows


def style_sheet(ws, money_cols=()):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill, c.font = fill, Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 45)
    for c in money_cols:
        for cell in ws[c][1:]:
            cell.number_format = "#,##0"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date-from", default="2026-10-01")
    ap.add_argument("--date-to", default="2026-10-04")
    ap.add_argument("--limit", type=int, default=0,
                    help="sample N order pertama (uji struktur)")
    ap.add_argument("--skip-list", action="store_true",
                    help="pakai cache list JSON, tanpa fetch list")
    ap.add_argument("--map", default="../sku_to_item_code.csv")
    args = ap.parse_args()

    token, email, user_id = creds()
    CACHE.mkdir(exist_ok=True)
    tag = f"{args.date_from}_sd_{args.date_to}"
    list_path = HERE / f".cache_details_list_{tag}.json"
    s = session(token, email)

    if args.skip_list and list_path.exists():
        orders = json.loads(list_path.read_text())
        print(f"list dari cache: {len(orders)} order")
    else:
        print("== 1. fetch list ==")
        orders = fetch_list(s, token, user_id, args.date_from, args.date_to)
        list_path.write_text(json.dumps(orders))
        print(f"  unik: {len(orders)} -> cache {list_path.name}")
    headers = [flat_header(o) for o in orders]
    if args.limit:
        headers = headers[:args.limit]
        print(f"  SAMPLE: {len(headers)} order pertama")

    print("== 2. fetch detail ==")
    details, failed, done = {}, [], 0
    t0 = time.time()
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        futs = {pool.submit(fetch_detail, s, token, user_id, h["order_id"]): h
                for h in headers}
        for f in as_completed(futs):
            oid, attrs, err = f.result()
            done += 1
            if done % 250 == 0 or done == len(headers):
                print(f"\r  {done}/{len(headers)} gagal={len(failed)} "
                      f"{time.time()-t0:.0f}s", end="", flush=True)
            if err:
                failed.append({"order_no": futs[f]["order_no"],
                               "order_id": oid, "error": err})
            else:
                details[oid] = attrs
    print(f"\n  OK: {len(details)}, gagal: {len(failed)}, "
          f"{time.time()-t0:.0f}s")

    mp = {}
    with open(HERE / args.map) as f:
        for r in csv.DictReader(f):
            mp[r["sku_eresto"]] = r

    print("== 3. rakit baris SKU ==")
    item_rows, recon = [], []
    for h in headers:
        ono = h["order_no"]
        attrs = details.get(h["order_id"])
        if attrs is None:
            recon.append({"order_no": ono, "outlet_name": h["outlet_name"],
                          "header_total": h["total"], "sum_paid": "",
                          "selisih": "", "status": "DETAIL GAGAL"})
            continue
        items = attrs_items(attrs)
        sub = round(sum(i["line_paid"] for i in items
                        if not i["void"] and i["paid_quantity"] > 0), 2)
        if not items:
            recon.append({"order_no": ono, "outlet_name": h["outlet_name"],
                          "header_total": h["total"], "sum_paid": 0,
                          "selisih": h["total"], "status": "TANPA ITEM"})
        else:
            diff = round(h["total"] - sub, 2)
            recon.append({
                "order_no": ono, "outlet_name": h["outlet_name"],
                "header_total": h["total"], "sum_paid": sub,
                "selisih": diff,
                "status": "COCOK" if abs(diff) < 1 else "SELISIH"})
        for it in items:
            m = mp.get(it["product_sku"], {})
            item_rows.append({
                "tanggal": h["tanggal"], "order_no": ono,
                "outlet_name": h["outlet_name"], "branch": h["branch"],
                "kasir": h["kasir"], "kategori": h["kategori"],
                "product_sku": it["product_sku"],
                "product_name": it["product_name"],
                "item_code_erpnext": m.get("item_code_erpnext", ""),
                "item_name_erpnext": m.get("item_name_erpnext", ""),
                "sumber_mapping": m.get("sumber", "UNMAPPED"),
                "quantity": it["quantity"],
                "paid_quantity": it["paid_quantity"], "price": it["price"],
                "line_paid": it["line_paid"],
                "void": it["void"], "void_note": it["void_note"],
            })

    df = pd.DataFrame(item_rows)
    pay = df[(~df["void"]) & (df["paid_quantity"] > 0)].copy()
    by_sku = pay.groupby(["product_sku", "item_code_erpnext",
                          "item_name_erpnext", "sumber_mapping"],
                         as_index=False).agg(
        nama_contoh=("product_name", "first"),
        paid_qty=("paid_quantity", "sum"), omzet=("line_paid", "sum"),
        n_order=("order_no", "nunique"),
        n_outlet=("outlet_name", "nunique")).sort_values("omzet",
                                                         ascending=False)
    unmapped = by_sku[by_sku["sumber_mapping"] == "UNMAPPED"]
    by_day = pd.DataFrame(headers).groupby("tanggal").agg(
        trx=("order_no", "count"), outlet=("outlet_name", "nunique"),
        omzet=("total", "sum")).reset_index()

    checks = [
        (f"detail gagal = 0 ({len(failed)})", len(failed) == 0),
        (f"order COCOK header-vs-item: "
         f"{sum(1 for r in recon if r['status'] == 'COCOK')}/{len(recon)}",
         True),
        (f"SKU termap {len(by_sku) - len(unmapped)}/{len(by_sku)}, "
         f"unmapped Rp{unmapped['omzet'].sum():,.0f}" if len(by_sku) else "tanpa SKU",
         True),
    ]
    print("validasi:")
    for n, _ in checks:
        print(f"  [INFO] {n}")
    print(f"  SELISIH: {sum(1 for r in recon if r['status'] == 'SELISIH')}, "
          f"TANPA ITEM: {sum(1 for r in recon if r['status'] == 'TANPA ITEM')}")

    out = f"Penjualan_{tag}_detail_SKU.xlsx"
    with pd.ExcelWriter(HERE / out, engine="openpyxl") as w:
        by_day.to_excel(w, sheet_name="Rekap_Harian", index=False)
        by_sku.to_excel(w, sheet_name="Rekap_SKU", index=False)
        df.to_excel(w, sheet_name="Order_Items", index=False)
        pd.DataFrame(recon).to_excel(w, sheet_name="Rekonsiliasi", index=False)
        pd.DataFrame([{"cek": n} for n, _ in checks]).to_excel(
            w, sheet_name="Validation", index=False)
        pd.DataFrame(failed).to_excel(w, sheet_name="Failed", index=False)
        style_sheet(w.sheets["Rekap_Harian"], ("D",))
        style_sheet(w.sheets["Rekap_SKU"], ("F", "G"))
        style_sheet(w.sheets["Order_Items"], ("M", "N", "O"))
        style_sheet(w.sheets["Rekonsiliasi"], ("C", "D", "E"))
        style_sheet(w.sheets["Validation"])
        style_sheet(w.sheets["Failed"])
    print(f"Excel: {out} ({len(df)} baris item, {len(by_sku)} SKU)")


if __name__ == "__main__":
    main()
