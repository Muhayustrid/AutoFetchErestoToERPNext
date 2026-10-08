# SPIKE sekali pakai: rekap penjualan SEMUA outlet 1-4 Okt 2026.
# List-only (tanpa detail per item, 27rb order). Kredensial via env.
import os
import sys
import time

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://ropi.eresto.id"
PER_PAGE = 200
MAX_RETRIES = 3
TIMEOUT = 90
TEMP_HTTP_CODES = {429, 500, 502, 503, 504}
DATE_FROM = "2026-10-01"
DATE_TO = "2026-10-04"


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def get_with_retry(session, url, params):
    last = None
    for att in range(1, MAX_RETRIES + 1):
        try:
            r = session.get(url, params=params, timeout=TIMEOUT)
            if r.status_code in TEMP_HTTP_CODES:
                last = f"HTTP {r.status_code}"
                time.sleep(2 ** att)
                continue
            r.raise_for_status()
            return r.json()
        except (requests.RequestException, ValueError) as e:
            last = str(e)[:150]
            time.sleep(2 ** att)
    fail(f"GET gagal {MAX_RETRIES}x: {last}")


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def style_sheet(ws, money_cols=()):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill, c.font = fill, Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="center")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 40)
    for c in money_cols:
        for cell in ws[c][1:]:
            cell.number_format = '#,##0'


def main():
    token = os.getenv("ERESTO_TOKEN", "").strip()
    email = os.getenv("ERESTO_EMAIL", "").strip()
    user_id = os.getenv("ERESTO_USER_ID", "").strip()
    if not (token and email and user_id):
        fail("env ERESTO_TOKEN/EMAIL/USER_ID kurang")

    s = requests.Session()
    s.headers.update({
        "accept": "application/json, text/plain, */*",
        "authorization": f"Token token={token}, email={email}",
        "referer": "https://ropi.eresto.id/manager/v2/index.html",
        "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "cache-control": "no-cache", "pragma": "no-cache",
    })

    url = f"{BASE_URL}/v3/orders"
    orders, page, total_count = [], 1, None
    t0 = time.time()
    while True:
        payload = get_with_retry(s, url, {
            "current_user_id": user_id,
            "date[geq]": f"{DATE_FROM} 00:00:00.000",
            "date[leq]": f"{DATE_TO} 23:59:59.999",
            "domain": "ropi", "outlet_id": "",
            "page": page, "per_page": PER_PAGE,
            "sort[property]": "date", "sort[order]": "asc",
            "token": token, "_": int(time.time() * 1000),
        })
        batch = payload.get("data", [])
        meta = payload.get("meta", {})
        total_count = meta.get("total_count", total_count)
        orders.extend(batch)
        print(f"\r  page {page}: +{len(batch)} (total {len(orders)}/{total_count})",
              end="", flush=True)
        if len(batch) < PER_PAGE or not meta.get("next_page"):
            break
        page += 1
    print(f"\nTerkumpul {len(orders)} order dalam {time.time()-t0:.0f}s")

    seen, rows = set(), []
    for o in orders:
        oid = o.get("id")
        if not oid or oid in seen:
            continue
        seen.add(oid)
        a = o.get("attributes", {}) or {}
        p = a.get("payment", {}) or {}
        rows.append({
            "tanggal": str(a.get("date", ""))[:10],
            "order_no": a.get("order_no", ""),
            "outlet_name": a.get("outlet_name", ""),
            "branch": a.get("outlet_branch_number", ""),
            "kasir": a.get("cashier_name", ""),
            "kategori": a.get("category", ""),
            "qty_item": num(a.get("item_quantity")),
            "void_qty": num(a.get("void_quantity")),
            "total": num(p.get("total")),
            "cash": num(p.get("cash_amount")),
            "ewallet": num(p.get("ewallet_amount")),
            "transfer": num(p.get("transfer_amount")),
            "debit": num(p.get("debit_amount")),
            "kredit": num(p.get("credit_amount")),
        })
    df = pd.DataFrame(rows)
    print(f"Unik: {len(df)} | outlet: {df['outlet_name'].nunique()} | "
          f"omzet: Rp{df['total'].sum():,.0f}")

    by_day = df.groupby("tanggal").agg(
        trx=("order_no", "count"), outlet=("outlet_name", "nunique"),
        qty=("qty_item", "sum"), omzet=("total", "sum"),
        cash=("cash", "sum"), ewallet=("ewallet", "sum")).reset_index()
    rekap_outlet = df.groupby(["outlet_name", "branch"]).agg(
        trx=("order_no", "count"), hari_aktif=("tanggal", "nunique"),
        qty=("qty_item", "sum"), omzet=("total", "sum"),
        cash=("cash", "sum"), ewallet=("ewallet", "sum")).reset_index()
    rekap_outlet["franchise"] = rekap_outlet["outlet_name"].str.contains(
        "FRC", case=False, na=False)
    rekap_outlet = rekap_outlet.sort_values("omzet", ascending=False)
    by_pay = pd.DataFrame([{
        "metode": m, "nominal": df[m].sum(),
        "porsi_%": round(100 * df[m].sum() / df["total"].sum(), 1)}
        for m in ["cash", "ewallet", "transfer", "debit", "kredit"]])

    out = (f"Penjualan_{DATE_FROM}_sd_{DATE_TO}.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        by_day.to_excel(w, sheet_name="Rekap_Harian", index=False)
        rekap_outlet.to_excel(w, sheet_name="Rekap_Outlet", index=False)
        by_pay.to_excel(w, sheet_name="Rekap_Pembayaran", index=False)
        df.to_excel(w, sheet_name="Raw_Orders", index=False)
        style_sheet(w.sheets["Rekap_Harian"], ("E", "F", "G"))
        style_sheet(w.sheets["Rekap_Outlet"], ("E", "F", "G"))
        style_sheet(w.sheets["Rekap_Pembayaran"], ("B",))
        style_sheet(w.sheets["Raw_Orders"], ("H", "I", "J", "K", "L", "M", "N"))
    print(f"Excel: {out}")
    print(by_day.to_string(index=False))
    print("\nTop 10 outlet:")
    print(rekap_outlet.head(10)[["outlet_name", "trx", "omzet"]].to_string(index=False))
    print(f"\nFranchise: {int(rekap_outlet['franchise'].sum())} outlet, "
          f"omzet Rp{rekap_outlet[rekap_outlet['franchise']]['omzet'].sum():,.0f}")


if __name__ == "__main__":
    main()
