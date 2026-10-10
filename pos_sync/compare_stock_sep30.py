#!/usr/bin/env python3
"""READ-ONLY: banding saldo stok ERPNext LAMA vs BARU per 2026-09-30 (23:00 & EOD).

Instruksi user 2026-10-08: "cek pada stock balance di erpnext lama dan baru
per 30 september jam 23:00, cek apakah ada yang sama dan yang beda; kalau ada
yang beda cari tau penyebabnya, laporin dulu ke saya kenapa beda."
Revisi scope 2026-10-09: "kecualikan company selain PT. JUARA ROTI INDONESIA"
-> hanya gudang company PT. JUARA ROTI INDONESIA (ROPI<->JURI) yang dibanding;
Stores - RPMD / company lain dikeluarkan.

Metode (ledger, bukan Bin):
  saldo(wh,item)@cutoff = qty_after_transaction baris SLE TERAKHIR (urut
  posting_date, posting_time, name) dengan waktu <= cutoff, is_cancelled=0.
  Cutoff: A = 23:00:00 (literal permintaan), B = EOD 23:59:59.
  Catatan: baris ADJ legacy LAMA kadang act=0 dgn qat melompat (koreksi lama)
  — qat tetap otoritatif; sum(actual_qty) TIDAK dipakai.

Fase:
  --dump    tarik ulang SLE kedua site + item master + RECO (cache).
  --compare hitung saldo + anotasi penyebab -> stock_sep30_result.json.
  --report  tulis pos_sync/Laporan_Stock_Sep30_LAMA_vs_BARU.xlsx.
Default: semua. Read-only di kedua site.
"""
import argparse
import json
import os
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prod_common import Client, load_env  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DUMP = os.path.join(HERE, "stock_sep30_dump")
RESULT = os.path.join(DUMP, "stock_sep30_result.json")
REPORT = os.path.join(HERE, "Laporan_Stock_Sep30_LAMA_vs_BARU.xlsx")
OLD_BASE = "https://rotiropi.j.frappe.cloud"
TOL = 1e-6
# scope: HANYA company ini (instruksi user 2026-10-09). Nama company sama di
# kedua site; company outlet lain (RPMD dst.) dikeluarkan dari pembandingan.
COMPANY = "PT. JUARA ROTI INDONESIA"

WH_PAIRS = [
    ("Gudang Bahan Baku - ROPI", "Gudang Bahan Baku - JURI"),
    ("Gudang Barang Jadi - ROPI", "Gudang Barang Jadi - JURI"),
    ("Gudang Jakarta - ROPI", "Gudang Jakarta - JURI"),
    ("Gudang Produksi - ROPI", "Gudang Produksi - JURI"),
    ("Work In Progress - ROPI", "Work In Progress - JURI"),
    ("Cold Storage Produksi - ROPI", "Cold Storage Produksi - JURI"),
    ("Gudang RnD - ROPI", "Gudang R&D - JURI"),
    ("Gudang Rejec - ROPI", "Gudang Reject - JURI"),
    ("Pabrik Klaten - ROPI", "Pabrik - JURI"),
    ("Goods In Transit - ROPI", "Goods In Transit - JURI"),
    ("Gudang Outlet - ROPI", "Gudang Outlet - JURI"),
    ("Gudang Peralatan - ROPI", "Gudang Peralatan - JURI"),
    ("Gudang Vendor - ROPI", "Gudang Vendor - JURI"),
    ("Gudang GA - ROPI", "Gudang GA - JURI"),
    # 3 gudang testing DIHAPUS 2026-10-09 (perintah user: "cuma testing aja")
    # di kedua site via pos_sync/hapus_warehouse_testing.py:
    #   Gudang Finished Goods Testing - ROPI, Gudang Testing Manufacture - ROPI,
    #   WIP Testing - ROPI  (-> tanpa pasangan lagi).
    ("Gudang Finished Goods Testing - ROPI", None),
    ("Gudang HRD - ROPI", None),
    ("Gudang Testing Manufacture - ROPI", None),
    ("Mini Outlet - ROPI", None),
    ("WIP Testing - ROPI", None),
]
PAIR_MAP = {a: b for a, b in WH_PAIRS}

# mapping bundle/penjualan -> matang (docs/05). Kunci = (LAMA code -> BARU code).
# Stok yang sama hidup di level berbeda; dipakai utk mengenali pasangan.
BUNDLE2MATANG = {
    "RR0001": "PJO2603",   # Ropi Butter -> Ropi Butter Matang
    "RR0002": "PJO2601",   # Ropi Coklat -> Ropi Coklat Matang
    "RR0003": "PJO2602",   # Ropi Keju   -> Ropi Keju Matang
    "RR0008": "PJO2605",   # Croissant   -> Croissant Matang
    "RR0009": "PJO2604",   # Gempi       -> Gempi Matang
    "RP001": "PJO2603", "RP002": "PJO2601", "RP003": "PJO2602",
}

# gudang BARU yang diinjeksi opening 30 Sep (window audit) — sisanya (GA, R&D,
# Peralatan, Vendor, dst.) diinjeksi dgn tanggal Okt / di luar window ini.
BIG6 = {"Gudang Barang Jadi - JURI", "Gudang Bahan Baku - JURI",
        "Gudang Jakarta - JURI", "Gudang Produksi - JURI",
        "Work In Progress - JURI", "Cold Storage Produksi - JURI"}

# koreksi disengaja 7-8 Okt (terdokumentasi di execute_fix_wip_and_wos.py /
# catchup_oct7.py): item -> catatan
DELIBERATE = {}


def _add_delib(wh, items, note):
    for it in items:
        DELIBERATE[(wh, it)] = note


_add_delib("Gudang Bahan Baku - JURI",
           ["BB260011", "BB260012", "BB260018", "BB260021", "BB260022"],
           "koreksi opname (RECO-JURI-261008001)")
_add_delib("Work In Progress - JURI",
           ["BB260002", "BB260011", "BB260015", "BB260016", "BB260019",
            "BB260022", "BB260029"],
           "kompensasi over-consumption MFG (RECO-JURI-261008002)")
_add_delib("Cold Storage Produksi - JURI",
           ["PJ260010", "PJ260011", "PJ260012", "PJ260013", "PJ260014",
            "PJ260016"],
           "FG serah terima 7 Okt (RECO-JURI-261007006)")
_add_delib("Gudang Barang Jadi - JURI", ["BB26008"],
           "opening susulan gula cair Repack (RECO-JURI-261007007)")
_add_delib("Gudang Produksi - JURI", ["BB260055"],
           "opening susulan telur (RECO-JURI-261007001)")
_add_delib("Gudang Produksi - JURI", ["BHP260022"],
           "opening susulan kertas alas (RECO-JURI-261007003)")

# jelas terdokumentasi tapi bukan kategori di atas
KNOWN_NOTES = {
    ("Gudang Jakarta - JURI", "PEL260075"):
        "konversi injeksi 1 Pack=10 Pcs (150 Pack -> 1500); LAMA mencatat "
        "150 di unit Pcs - verifikasi satuan yang benar",
    ("Gudang Barang Jadi - JURI", "BB260060"):
        "kode bentrok: LAMA BB260060=Olesan Pastry (250 g); BARU "
        "BB260060=Bubuk Hepi Drink Chocolate (item baru hasil opname GBJ)",
    ("Gudang Barang Jadi - JURI", "BB260061"):
        "item baru di BARU (Kopi Hitam) hasil opname GBJ; LAMA tidak punya",
    ("Gudang Barang Jadi - JURI", "BB260062"):
        "item baru di BARU (Susu Bubuk) hasil opname GBJ; LAMA tidak punya",
    ("Gudang Barang Jadi - JURI", "PEL260087"):
        "RECO-JURI-261007004 menetapkan 201.000 (klaim 'saldo 30 Sep EOD'); "
        "buku LAMA EOD=238.100 — selisih 37.100 perlu verifikasi",
    ("Gudang Jakarta - JURI", "PEL260087"):
        "RECO-JURI-261007005 menetapkan 34.000 (klaim 'saldo 30 Sep EOD'); "
        "buku LAMA EOD=27.850 — selisih 6.150 perlu verifikasi",
    ("Gudang Jakarta - JURI", "PEL260019"):
        "injeksi Jakarta memetakan 'BB0077 Kaos Pendek' (opname real=20) ke "
        "PEL260019; buku LAMA PEL260019=0 — cek mapping item",
}


def log(m):
    print(m, flush=True)


def tosec(t):
    s = str(t or "0").split(".")
    p = s[0].split(":")
    while len(p) < 3:
        p.append("0")
    try:
        return int(float(p[0])) * 3600 + int(float(p[1])) * 60 + float(p[2])
    except Exception:
        return -1


def secstr(sec):
    if sec is None or sec < 0:
        return "-"
    return "%02d:%02d:%02d" % (int(sec // 3600), int((sec % 3600) // 60),
                               int(sec % 60))


# ---------------------------------------------------------------- fetch/dump
def fetch(dt, fields, filters, order, c):
    out, start = [], 0
    while True:
        st, d = c.get("/api/resource/" + dt, {
            "fields": json.dumps(fields), "filters": json.dumps(filters),
            "order_by": order, "limit_page_length": 500,
            "limit_start": start})
        rows = d.get("data", []) if st == 200 else []
        if st != 200:
            raise SystemExit("GET %s %s gagal HTTP %s %s" %
                             (dt, c.base, st, str(d)[:200]))
        out += rows
        start += len(rows)
        if len(rows) < 500:
            break
    return out


SLE_F = ["name", "company", "warehouse", "item_code", "posting_date",
         "posting_time", "actual_qty", "qty_after_transaction",
         "voucher_type", "voucher_no", "is_cancelled"]


def do_dump():
    os.makedirs(DUMP, exist_ok=True)
    mig = load_env("/tmp/eresto_mig.env")
    prod = load_env("/tmp/eresto_prod.env")
    old = Client(OLD_BASE, mig["ERPNEXT_OLD_CLOUD_API_KEY"], "LAMA")
    new = Client(prod["ERPNEXT_URL"], prod["ERPNEXT_API_KEY"], "BARU")
    for site, c in (("lama", old), ("baru", new)):
        rows = fetch("Stock Ledger Entry", SLE_F,
                     [["posting_date", "<=", "2026-09-30"],
                      ["is_cancelled", "=", 0],
                      ["company", "=", COMPANY]],
                     "posting_date asc, posting_time asc, name asc", c)
        json.dump(rows, open(os.path.join(DUMP, "sle_%s_sep30.json" % site), "w"),
                  default=str)
        log("  dump SLE %s: %d baris <= 2026-09-30 (company=%s)"
            % (site, len(rows), COMPANY))
        items = fetch("Item", ["name", "item_name", "stock_uom",
                               "is_stock_item", "disabled"], [], "name asc", c)
        json.dump(items, open(os.path.join(DUMP, "items_%s.json" % site), "w"),
                  default=str)
        log("  dump items %s: %d" % (site, len(items)))
    recos = fetch("Stock Reconciliation",
                  ["name", "company", "set_warehouse", "posting_date",
                   "posting_time", "docstatus", "purpose", "creation"],
                  [], "creation asc", new)
    json.dump(recos, open(os.path.join(DUMP, "recos_baru.json"), "w"),
              default=str)
    log("  dump RECO BARU: %d" % len(recos))


def jload(name, default=None):
    p = os.path.join(DUMP, name)
    if not os.path.exists(p):
        return default
    return json.load(open(p))


# ---------------------------------------------------------------- balances
def build_index(rows):
    """{(wh,item): [(date, sec, name, qat, voucher, act)] urut naik}.

    Tanpa filter tanggal: baris terakhir bisa jauh sebelum 30 Sep (barang
    tanpa mutasi) — qat-nya adalah saldo berjalan sampai cutoff.
    """
    idx = defaultdict(list)
    for r in rows:
        d = str(r["posting_date"])[:10]
        idx[(r["warehouse"], r["item_code"])].append(
            (d, tosec(r["posting_time"]), r["name"],
             float(r["qty_after_transaction"] or 0), r["voucher_no"],
             float(r["actual_qty"] or 0)))
    for k in idx:
        idx[k].sort()
    return idx


def last_at(idx, key, cut_sec, cut_date="2026-09-30"):
    rows = idx.get(key)
    if not rows:
        return None
    best = None
    for e in rows:
        if (e[0], e[1]) <= (cut_date, cut_sec):
            best = e
    return best


def qat_map(idx, key, cut_sec):
    """scan seluruh rantai qat LAMA <= cut: {qat: (sec, voucher)} (terakhir menang)."""
    out = {}
    for e in idx.get(key, []):
        if e[0] == "2026-09-30" and e[1] <= cut_sec:
            out[round(e[3], 4)] = (e[1], e[4])
        elif e[0] < "2026-09-30":
            out[round(e[3], 4)] = (e[1], e[4])
    return out


# ---------------------------------------------------------------- compare
def annotate(rows, idx_l, name2lama, name2baru, items_l, items_b):
    for r in rows:
        if abs(r["delta"]) < TOL:
            r["cause"] = "SAMA"
            continue
        wa, wb, item = r["wh_lama"], r["wh_baru"], r["item"]
        r["note"] = ""
        if wb is None:
            r["cause"] = "GUDANG_TANPA_PASANGAN"
            r["note"] = "gudang LAMA ini tidak ada / belum dibuat di BARU"
            continue
        nm = (items_b.get(item) or items_l.get(item) or {}).get("item_name")
        r["item_name"] = nm
        # UOM beda
        if r["uom_lama"] and r["uom_baru"] and r["uom_lama"] != r["uom_baru"]:
            r["note"] += "UOM master beda: LAMA=%s BARU=%s. " % (
                r["uom_lama"], r["uom_baru"])
        b = r["baru"]
        # 1) koreksi disengaja 7-8 Okt (dokumen CATCHUP/fix)
        if (wb, item) in DELIBERATE:
            r["cause"] = "KOREKSI_DISENGAJA"
            r["note"] += DELIBERATE[(wb, item)]
            continue
        # 2) remap kode (nama sama tapi kode pindah; item BARU tdk ada di LAMA)
        if item not in items_l and name2lama.get(nm):
            r["cause"] = "REMAP_KODE"
            r["note"] += ("kode LAMA utk nama ini: %s -> dipetakan ke %s di BARU"
                          % (name2lama[nm][0], item))
            continue
        # 3) item betul-betul baru (tidak ada namanya di LAMA)
        if item not in items_l:
            r["cause"] = "ITEM_BARU"
            r["note"] += "item ada di BARU, tidak ada di master LAMA "
            r["note"] += "(item baru / level rantai berbeda, docs/05)"
            continue
        # 4) khusus yg sudah diketahui jelas
        if (wb, item) in KNOWN_NOTES:
            r["cause"] = "CEK_KHUSUS"
            r["note"] += KNOWN_NOTES[(wb, item)]
            continue
        # 5) BARU = posisi buku LAMA pd jam beda (opname pagi/siang vs malam)
        e_tb = last_at(idx_l, (wa, item), r["baru_last_sec"] or 0)
        if e_tb is not None and abs(e_tb[3] - b) < TOL:
            after = [e for e in idx_l.get((wa, item), [])
                     if (e[0], e[1]) > ("2026-09-30", r["baru_last_sec"] or 0)
                     and e[0] == "2026-09-30"]
            mv = ", ".join("%s %+.0f" % (e[4], e[5]) for e in after[:4])
            r["cause"] = "WAKTU_OPNAME"
            r["note"] += ("BARU = buku LAMA per jam opname %s; LAMA masih "
                          "bergerak sesudahnya: %s"
                          % (secstr(r["baru_last_sec"]), mv or "-"))
            continue
        # 6) BARU kosong = tidak dihitung di opname / belum diinjeksi
        if abs(b) < TOL:
            if wb in BIG6:
                r["cause"] = "TIDAK_DIHITUNG_OPNAME"
                r["note"] += ("BARU tidak menyimpan saldo item ini (tidak ada "
                              "di daftar opening 30 Sep); buku LAMA=%.2f "
                              "(%s %s)" % (r["lama"],
                                           secstr(r["lama_last_sec"]),
                                           r["lama_last_voucher"] or ""))
            else:
                r["cause"] = "BELUM_DIINJEKSI"
                if wb in ("Gudang GA - JURI", "Gudang Peralatan - JURI",
                          "Gudang Vendor - JURI"):
                    r["note"] += ("opening gudang ini di BARU diposting 2 Okt "
                                  "(di luar window audit); buku LAMA=%.2f"
                                  % r["lama"])
                else:
                    r["note"] += ("belum ada opening di BARU utk gudang ini; "
                                  "buku LAMA=%.2f" % r["lama"])
            continue
        # 7) BARU = posisi buku LAMA pada jam lain (opname dari sumber lain)
        qm = qat_map(idx_l, (wa, item), 23 * 3600 + 3599)
        hit = qm.get(round(b, 4))
        if hit:
            r["cause"] = "POSISI_LAMA_JAM_LAIN"
            r["note"] += ("BARU = posisi buku LAMA jam %s (%s) — bukan EOD"
                          % (secstr(hit[0]), hit[1]))
            continue
        # 8) sisa: buku LAMA vs opname fisik eResto (drift sebelum migrasi)
        r["cause"] = "BUKU_LAMA_VS_FISIK"
        r["note"] += ("BARU = hasil opname fisik eResto; buku LAMA berbeda "
                      "%.0f — drift sebelum migrasi, perlu keputusan angka"
                      % r["delta"])


def run_compare():
    lama = jload("sle_lama_sep30.json")
    baru = jload("sle_baru_sep30.json")
    if lama is None or baru is None:
        raise SystemExit("dump belum ada — jalankan --dump dulu")
    idx_l = build_index(lama)
    idx_b = build_index(baru)
    items_l = {i["name"]: i for i in (jload("items_lama.json") or [])}
    items_b = {i["name"]: i for i in (jload("items_baru.json") or [])}
    name2lama, name2baru = defaultdict(list), defaultdict(list)
    for c, i in items_l.items():
        name2lama[i["item_name"]].append(c)
    for c, i in items_b.items():
        name2baru[i["item_name"]].append(c)

    rows = []
    seen_b = set()
    for (wa, wb) in WH_PAIRS:
        keys_l = {k[1] for k in idx_l if k[0] == wa}
        keys_b = {k[1] for k in idx_b if k[0] == wb} if wb else set()
        for item in sorted(keys_l | keys_b):
            ea = last_at(idx_l, (wa, item), 23 * 3600 + 3599)
            eb = last_at(idx_b, (wb, item), 23 * 3600 + 3599) if wb else None
            a = ea[3] if ea else 0.0
            b = eb[3] if eb else 0.0
            if abs(a) < TOL and abs(b) < TOL:
                continue
            e23a = last_at(idx_l, (wa, item), 23 * 3600)
            e23b = last_at(idx_b, (wb, item), 23 * 3600) if wb else None
            a23 = e23a[3] if e23a else 0.0
            b23 = e23b[3] if e23b else 0.0
            if wb:
                seen_b.add((wb, item))
            rows.append({
                "wh_lama": wa, "wh_baru": wb, "item": item,
                "item_name": (items_b.get(item) or items_l.get(item) or {}
                              ).get("item_name"),
                "lama": a, "baru": b, "delta": b - a,
                "lama_2300": a23, "baru_2300": b23, "delta_2300": b23 - a23,
                "uom_lama": (items_l.get(item) or {}).get("stock_uom"),
                "uom_baru": (items_b.get(item) or {}).get("stock_uom"),
                "lama_last_sec": ea[1] if ea else None,
                "lama_last_voucher": ea[4] if ea else None,
                "baru_last_sec": eb[1] if eb else None,
                "baru_last_voucher": eb[4] if eb else None,
                "cause": "", "note": "",
            })
    annotate(rows, idx_l, name2lama, name2baru, items_l, items_b)

    # BARU keys yang tidak masuk pasangan gudang
    only_b = []
    paired_b = {p[1] for p in WH_PAIRS if p[1]}
    for k, v in sorted(idx_b.items()):
        if k[0] in paired_b and (k[0], k[1]) not in seen_b:
            e = last_at(idx_b, k, 23 * 3600 + 3599)
            if e and abs(e[3]) >= TOL:
                only_b.append({"wh": k[0], "item": k[1], "qty": e[3],
                               "voucher": e[4]})
    only_l_wh = defaultdict(lambda: [0, 0.0])
    for k in idx_l:
        if k[0] in PAIR_MAP and PAIR_MAP[k[0]] is None:
            e = last_at(idx_l, k, 23 * 3600 + 3599)
            if e and abs(e[3]) >= TOL:
                only_l_wh[k[0]][0] += 1
                only_l_wh[k[0]][1] += e[3]

    res = {"rows": rows, "only_baru": only_b,
           "lama_wh_tanpa_pasangan": {k: v for k, v in only_l_wh.items()}}
    json.dump(res, open(RESULT, "w"), default=str, indent=1)

    n = defaultdict(int)
    for r in rows:
        n[r["cause"]] += 1
    log("")
    log("=== RINGKASAN KATEGORI (gudang berpasangan) ===")
    order = ["SAMA", "WAKTU_OPNAME", "BUKU_LAMA_VS_FISIK", "KOREKSI_DISENGAJA",
             "POSISI_LAMA_JAM_LAIN", "REMAP_KODE", "ITEM_BARU", "CEK_KHUSUS",
             "TIDAK_DIHITUNG_OPNAME", "BELUM_DIINJEKSI",
             "GUDANG_TANPA_PASANGAN"]
    for c in order:
        if n[c] or c == "SAMA":
            log("  %-22s %d" % (c, n[c]))
    log("  HANYA_BARU             %d" % len(only_b))
    log("  gudang LAMA tanpa pasangan: %d gudang"
        % len(res["lama_wh_tanpa_pasangan"]))
    log("")
    log("=== BUKU_LAMA_VS_FISIK (perlu keputusan angka) ===")
    for r in rows:
        if r["cause"] == "BUKU_LAMA_VS_FISIK":
            log("  %s|%s|%s LAMA=%.2f BARU=%.2f delta=%+.2f"
                % (r["wh_baru"], r["item"], r.get("item_name"),
                   r["lama"], r["baru"], r["delta"]))
    log("")
    log("=== CEK_KHUSUS ===")
    for r in rows:
        if r["cause"] in ("CEK_KHUSUS", "PERLU_DIPERIKSA"):
            log("  %s|%s|%s LAMA=%.2f BARU=%.2f delta=%+.2f | %s"
                % (r["wh_baru"], r["item"], r.get("item_name"),
                   r["lama"], r["baru"], r["delta"], r["note"]))
    return res


# ---------------------------------------------------------------- report
def run_report(res=None):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    if res is None:
        res = json.load(open(RESULT))
    rows = res["rows"]
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="D9EAF7")
    bold = Font(bold=True)
    widths = [30, 30, 11, 34, 14, 14, 13, 14, 14, 9, 9, 11, 26, 12, 24, 46]
    cols = ["Wh LAMA", "Wh BARU", "Item", "Nama Item", "Qty LAMA(EOD)",
            "Qty BARU(EOD)", "Delta EOD", "Qty LAMA(23:00)", "Qty BARU(23:00)",
            "UOM L", "UOM B", "LAMA jam", "LAMA voucher",
            "BARU jam", "BARU voucher", "Catatan"]

    def fmt(r):
        return [r["wh_lama"], r["wh_baru"] or "-", r["item"],
                r.get("item_name") or "", r["lama"], r["baru"], r["delta"],
                r["lama_2300"], r["baru_2300"],
                r["uom_lama"] or "", r["uom_baru"] or "",
                secstr(r["lama_last_sec"]), r["lama_last_voucher"] or "",
                secstr(r["baru_last_sec"]), r["baru_last_voucher"] or "",
                r.get("note") or ""]

    def sheet(name, sel):
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = bold
            cell.alignment = Alignment(horizontal="center")
        for r in sel:
            ws.append(fmt(r))
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    from collections import Counter
    cnt = Counter(r["cause"] for r in rows)
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["Laporan", "Banding Stock Balance LAMA vs BARU per 2026-09-30"])
    ws.append(["Metode", "qty_after_transaction baris SLE terakhir <= cutoff"])
    ws.append(["Cutoff", "23:00:00 (literal) & EOD 23:59:59"])
    ws.append(["", ""])
    ws.append(["Kategori", "Jumlah"])
    for c in ["SAMA", "WAKTU_OPNAME", "BUKU_LAMA_VS_FISIK",
              "KOREKSI_DISENGAJA", "POSISI_LAMA_JAM_LAIN", "REMAP_KODE",
              "ITEM_BARU", "CEK_KHUSUS", "TIDAK_DIHITUNG_OPNAME",
              "BELUM_DIINJEKSI", "GUDANG_TANPA_PASANGAN"]:
        ws.append([c, cnt[c]])
    n23 = sum(1 for r in rows if abs(r["delta_2300"]) < TOL)
    nE = sum(1 for r in rows if abs(r["delta"]) < TOL)
    ws.append(["Cocok di potret 23:00 literal", n23])
    ws.append(["Cocok di potret EOD 23:59", nE])
    ws.append(["HANYA_BARU", len(res["only_baru"])])
    ws.append(["Gudang LAMA tanpa pasangan (jumlah gudang)",
               len(res["lama_wh_tanpa_pasangan"])])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 12

    sheet("Sama", [r for r in rows if r["cause"] == "SAMA"])
    sheet("Sama_EOD_Beda_2300",
          [r for r in rows if abs(r["delta_2300"]) >= TOL
           and abs(r["delta"]) < TOL])
    sheet("Beda_Buku_vs_Fisik",
          [r for r in rows if r["cause"] == "BUKU_LAMA_VS_FISIK"])
    sheet("Beda_Waktu_Opname", [r for r in rows if r["cause"] == "WAKTU_OPNAME"])
    sheet("Beda_Koreksi_Disengaja",
          [r for r in rows if r["cause"] == "KOREKSI_DISENGAJA"])
    sheet("Beda_Remap_Kode", [r for r in rows if r["cause"] == "REMAP_KODE"])
    sheet("Beda_Item_Baru", [r for r in rows if r["cause"] == "ITEM_BARU"])
    sheet("Beda_Cek_Khusus",
          [r for r in rows if r["cause"] in ("CEK_KHUSUS", "PERLU_DIPERIKSA")])
    sheet("Beda_Tidak_Dihitung",
          [r for r in rows if r["cause"] == "TIDAK_DIHITUNG_OPNAME"])
    sheet("Beda_Belum_Diinjeksi",
          [r for r in rows if r["cause"] == "BELUM_DIINJEKSI"])
    sheet("Beda_Posisi_Jam_Lain",
          [r for r in rows if r["cause"] == "POSISI_LAMA_JAM_LAIN"])
    sheet("Beda_Gudang_Tanpa_Pasangan",
          [r for r in rows if r["cause"] == "GUDANG_TANPA_PASANGAN"])

    ws2 = wb.create_sheet("Hanya_Baru")
    ws2.append(["Wh", "Item", "Nama", "Qty", "Voucher"])
    for c in range(1, 6):
        ws2.cell(row=1, column=c).fill = head_fill
        ws2.cell(row=1, column=c).font = bold
    for r in res["only_baru"]:
        ws2.append([r["wh"], r["item"], "", r["qty"], r["voucher"]])
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Gudang_Lama_Tanpa_Pasangan")
    ws3.append(["Gudang LAMA", "Jml item", "Total qty"])
    for c in range(1, 4):
        ws3.cell(row=1, column=c).fill = head_fill
        ws3.cell(row=1, column=c).font = bold
    for k, v in sorted(res["lama_wh_tanpa_pasangan"].items()):
        ws3.append([k, v[0], v[1]])
    ws3.freeze_panes = "A2"

    wb.save(REPORT)
    log("")
    log("Laporan ditulis: %s" % REPORT)
    log("OK")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dump", action="store_true", help="tarik ulang data dasar")
    ap.add_argument("--compare", action="store_true", help="hitung + JSON")
    ap.add_argument("--report", action="store_true", help="tulis Excel")
    ap.add_argument("--probe", action="store_true", help="uji koneksi saja")
    a = ap.parse_args()
    any_phase = a.dump or a.compare or a.report or a.probe
    if a.probe:
        mig = load_env("/tmp/eresto_mig.env")
        prod = load_env("/tmp/eresto_prod.env")
        old = Client(OLD_BASE, mig["ERPNEXT_OLD_CLOUD_API_KEY"], "LAMA")
        new = Client(prod["ERPNEXT_URL"], prod["ERPNEXT_API_KEY"], "BARU")
        st, d = old.get("/api/method/frappe.client.get_count",
                        {"doctype": "Stock Ledger Entry"})
        log("LAMA ping: HTTP %s %s" % (st, d.get("message")))
        st, d = new.get("/api/method/frappe.client.get_count",
                        {"doctype": "Stock Ledger Entry"})
        log("BARU ping: HTTP %s %s" % (st, d.get("message")))
        return
    if not any_phase or a.dump:
        log("[dump] tarik data dasar kedua site ...")
        do_dump()
    res = None
    if not any_phase or a.compare:
        log("[compare] hitung saldo 30 Sep + anotasi penyebab ...")
        res = run_compare()
    if not any_phase or a.report:
        log("[report] tulis Excel ...")
        run_report(res)


if __name__ == "__main__":
    main()
