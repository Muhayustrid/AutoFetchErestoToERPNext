#!/usr/bin/env python3
"""PO eResto supplier internal JURI (PT. JUARA ROTI INDONESIA) 9-12 Okt 2026 -> ERPNext BARU.

Konteks (lihat KNOWLEDGE.md §5 "Sumber input pembelian" + PROJECT_STATE):
- Sumber: eResto `/v2/purchases` jendela `date[]` DAN `created_at[]` 9-12 Okt 2026,
  filter `internal_supplier = true` (lane supplier gudang JURI:
  'Gudang Barang Jadi - PH Klaten', 'Gudang Jakarta', 'Gudang Outlet - PH Klaten',
  'Gudang Peralatan Klaten').
- Pemetaan internal gudang (sama dgn `migrate_po_gudang_jakarta_9okt.py`):
    * Company       = divisi outlet pemesan (`purchase.division.name`)
    * Supplier      = 'PT. JUARA ROTI INDONESIA'
    * set_warehouse = 'Stores - <ABBR>' (gudang penerimaan outlet)
    * set_from_warehouse = gudang asal JURI (WAREHOUSE_MAP)
    * Buying Price List = 'Harga Outlet'
    * Beban baris  = '4210.000 - HPP Pembelian - <ABBR>', Cost Center 'Main - <ABBR>'
- ID dokumen = nomor PO eResto (`rename_doc`, form-encoded). docstatus mengikuti
  status eResto: `processing`/`completed` -> submit (1); `pending` -> draft (0).
  `rejected`/`canceled` -> TIDAK dimigrasi (keputusan user 2026-10-12: skip).
- Idempoten: dokumen yang sudah ada (by name ATAU title 'ERESTO <PO>') di-skip;
  rerun aman. Self-heal rename bila dokumen tertinggal dengan nama fallback.
- SKU tak termap TIDAK ditebak: PO masuk sheet `Skipped` dengan alasan.

Pemetaan SKU baru (BUKTI, bukan tebakan) — dari padanan 1:1 eResto PO332 series
di site LAMA (`PO332.220926.001/.002` = counter of eResto PO332.220926.*),
baris identik nama/qty/harga:
    BB0006 Apron             -> PEL260006 Apron Ropi      (3 Pcs @66.000)
    BB0042 Crispy Ball       -> BB260048 Crispy Ball (500 g) (2 Pack @42.000)
    BB0076 Kaos Panjang      -> PEL260068                 (6 Pcs @70.000)
    BB0078 Kaos Polo         -> PEL260019 Kaos Polo Ropi  (1 Pcs @82.000)
    BB0160 Topi              -> PEL260052 Topi            (4 Pcs @40.000)
    BB0185 Botol Spray       -> PEL260010                 (1 Pcs @4.000)
    BB0270 Kuas Silikon      -> PEL260023                 (1 Pcs @12.000)
    BB0328 Jilbab Outlet     -> SG260012 Jilbab           (3 Pcs @32.000)
    BB0353 Gula Cair 250g    -> BB26009 'Gula Cair (250 g)' (repack bottle,
                                stock_uom Botol = unit eResto 'Botol'; KNOWLEDGE §5)
  CATATAN: BARU BB260058/BB260059 = 'Bubuk Hepi Drink ...' (BUKAN Gula Cair) —
  jangan dipakai. BB0043 'Daun Teh Enak Kata Orang' belum ada padanannya dan
  hanya dipakai PO FRC Pasaman (di luar scope) — dibiarkan unmapped.

FRC di luar scope (company belum ada di BARU): 'Jogokariyan (FRC)',
'Pasaman (FRC)', 'Ugm (FRC)' -> PO masuk sheet `Skipped` (butuh onboarding).

CLI:
  python pos_sync/migrate_po_juri_9_12_okt.py [--dump] [--dry-run] [--verify]
      [--only PO042.091026.001 ...] [--date-from 2026-10-09] [--date-to 2026-10-12]

Env: /tmp/eresto_mig.env (ERESTO_*) + /tmp/eresto_prod.env (ERPNEXT_URL,
ERPNEXT_API_KEY). Token tidak pernah ditulis ke output.
"""
import argparse
import json
import os
import sys
import time
import urllib.parse
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prod_common import Client, JURI, olesan_pastry_item  # noqa: E402

# ============================================================
# ENV & KONSTANTA
# ============================================================

def load_env_file(filepath):
    env = {}
    if os.path.exists(filepath):
        with open(filepath) as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env[k.strip()] = v.strip().strip('"').strip("'")
    return env

ENV_MIG = load_env_file("/tmp/eresto_mig.env")
ENV_PROD = load_env_file("/tmp/eresto_prod.env")

ERESTO_BASE_URL = "https://ropi.eresto.id"
ERESTO_DOMAIN = "ropi"
ERESTO_TOKEN = os.environ.get("ERESTO_TOKEN") or ENV_MIG.get("ERESTO_TOKEN", "")
ERESTO_EMAIL = os.environ.get("ERESTO_EMAIL") or ENV_MIG.get("ERESTO_EMAIL", "")
ERESTO_USER_ID = os.environ.get("ERESTO_USER_ID") or ENV_MIG.get("ERESTO_USER_ID", "")

ERPNEXT_URL = os.environ.get("ERPNEXT_URL") or ENV_PROD.get(
    "ERPNEXT_URL", "https://ropierpnext.frappe.cloud")
ERPNEXT_KEY = os.environ.get("ERPNEXT_API_KEY") or ENV_PROD.get("ERPNEXT_API_KEY", "")

DATE_FROM_DEFAULT = "2026-10-09"
DATE_TO_DEFAULT = "2026-10-12"
PRICE_LIST = "Harga Outlet"
DUMP_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "po_9_12_dump")
DUMP_FILE = os.path.join(DUMP_DIR, "eresto_internal_9_12.json")
REPORT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "Laporan_Migrasi_PO_JURI_9_12_Okt2026.xlsx")

SKIP_STATUS = {"rejected", "canceled"}
SUBMIT_STATUS = {"processing", "completed"}

# Supplier eResto (nama gudang) -> source warehouse JURI.
WAREHOUSE_MAP = {
    "Gudang Bahan Baku - PH Klaten": "Gudang Bahan Baku - JURI",
    "Gudang Barang Jadi - PH Klaten": "Gudang Barang Jadi - JURI",
    "Gudang Office - PH Klaten": "Gudang GA - JURI",
    "Gudang Outlet - PH Klaten": "Gudang Outlet - JURI",
    "Gudang Peralatan Klaten": "Gudang Peralatan - JURI",
    "Gudang Produksi Klaten": "Gudang Produksi - JURI",
    "Gudang Jakarta": "Gudang Jakarta - JURI",
    "Gudang Vendor": "Gudang Vendor - JURI",
    "Head Office": "Gudang GA - JURI",
    "Gudang Maintenance": "Gudang Maintenance - JURI",
}

# Mapping SKU eResto -> Item Code ERPNext BARU (gabungan proven map + 9 SKU baru
# dengan bukti; lihat docstring & sheet Catatan).
ITEM_MAP = {
    "BB0006": "PEL260006",  # Apron                 [BARU 12 Okt, bukti PO332.220926.002 LAMA]
    "BB0042": "BB260048",   # Crispy Ball 500 gram  [BARU 12 Okt, bukti PO332.220926.001 LAMA]
    "BB0046": "PJ260012",   # Dough Ropi Butter
    "BB0047": "PJ260013",   # Dough Ropi Coklat
    "BB0049": "PJ260014",   # Dough Ropi Keju
    "BB0066": "PEL260072",  # Goodie Bag
    "BB0076": "PEL260068",  # Kaos Panjang          [BARU 12 Okt, bukti PO332.220926.002 LAMA]
    "BB0078": "PEL260019",  # Kaos Polo             [BARU 12 Okt, bukti PO332.220926.002 LAMA]
    "BB0081": "BB260040",   # Keju Chedar
    "BB0085": "PJ260016",   # Krim Kopi
    "BB0100": "BB260042",   # Minyak Oles (1 kg)
    "BB0102": "BB260043",   # Olesan Choco Crunchy (1 kg)
    "BB0103": "BB260044",   # Olesan Dark Choco (1 kg)
    "BB0105": "BB260045",   # Olesan Margarine
    "BB0110": "BB260047",   # Olesan Tiramisu (1 kg)
    "BB0112": "PEL260076",  # Paper Bag Ropi
    "BB0117": "PEL260077",  # Plastik Alas Gempi
    "BB0131": "PEL260038",  # Sarung Tangan Oven  [BARU 13 Okt, mapping user]
    "BB0136": "PEL260041",  # Sendok Es Batu      [BARU 13 Okt, mapping user]
    "BB0118": "PEL260081",  # Plastik Kresek Kecil (17)
    "BB0119": "PEL260080",  # Plastik Kresek Besar (24)
    "BB0120": "BB260019",   # Plastik Krim
    "BB0122": "PEL260083",  # Plastik Seal Ropi
    "BB0263": "ATK260041",  # Pulpen Outlet       [BARU 13 Okt, mapping user]
    # --- Mapping SKU keperluan outlet 12 Okt (bukti nama item master BARU /
    #     mapping user 13 Okt) ---
    "BB0068": "PEL260021",  # Kuas Oles Loyang (1,5 inch)
    "BB0071": "BB260039",   # Gula Stik           [bukti stok_sku_to_item.csv COCOK]
    "BB0074": "PEL260018",  # Kanebo
    "BB0080": "PEL260074",  # Kardus Ropi 20 roti [bukti stok_sku_to_item.csv COCOK]
    "BB0082": "BHP260023",  # Lakban Besar -> Lakban (2 inch) Outlet [mapping user]
    "BB0087": "BHP260020",  # Lakban Kecil -> Selotip (0,5 inch) Outlet [user]
    "BB0090": "PEL260022",  # Kuas Oles Loyang (2 inch)
    "BB0096": "BB260041",   # Meises Coklat (1 kg)
    "BB0141": "PEL260043",  # Solet Oles Gempi
    "BB0143": "PEL260085",  # Stirrer
    "BB0220": "PEL260037",  # Sapu
    "BB0226": "ATK260044",  # Buku Tulis
    "BB0228": "PEL260020",  # Keset
    "BB0229": "PEL260025",  # Lap Microfiber
    "BB0230": "ATK260038",  # Sticker Label Tom & Jerry -> Stiker Label [user]
    "BB0231": "ATK260043",  # Spidol
    "BB0234": "ATK260027",  # Penggaris Merah -> Penggaris
    "BB0260": "PEL260009",  # Bill Holder
    "BB0264": "ATK260046",  # Isi Staples
    "BB0267": "BHP260033",  # Tinta Isi Ulang Outlet -> Tinta Refill Spidol Outlet
    "BB0268": "PEL260014",  # Dispenser Sabun Cair
    "BB0272": "PEL260045",  # Spons Cuci Outlet
    "BB0273": "BHP260035",  # Sabun Lantai Outlet
    "BB0274": "BHP260034",  # Sabun Cuci Outlet
    "BB0283": "PEL260030",  # Papan Menu Outlet
    "BB0305": "BB260038",   # Gula Cair Sachet    [bukti stok_sku_to_item.csv COCOK]
    "BB0307": "PEL260028",  # Mika Display Menu A3
    "BB0351": "PEL260095",  # Kemoceng
    "BB0132": "BHP260027",  # Sarung Tangan Plastik
    "BB0134": "PEL260084",  # Sedotan
    "BB0160": "PEL260052",  # Topi                   [BARU 12 Okt, bukti PO332.220926.002 LAMA]
    "BB0168": "PJ260010",   # Dough Croissant
    "BB0171": "PEL260069",  # Box Gempi (100 pcs)
    "BB0174": "BB260046",   # TIDAK AKTIF Olesan Pastry -> Olesan Pastry (lihat aturan tanggal)
    "BB0185": "PEL260010",  # Botol Spray            [BARU 12 Okt, bukti PO332.220926.002 LAMA]
    "BB0191": "PEL260075",  # Kertas Thermal
    "BB0257": "PEL260071",  # Gelas Cup 9 Oz (isi 50 pcs)
    "BB0261": "PJ260011",   # Dough Gempi
    "BB0270": "PEL260023",  # Kuas Silikon           [BARU 12 Okt, bukti PO332.220926.002 LAMA]
    "BB0296": "PEL260078",  # Plastik Kresek 1 Cup
    "BB0297": "PEL260079",  # Plastik Kresek 2 Cup
    "BB0298": "PEL260070",  # Gelas (14 oz) (50 pcs)
    "BB0312": "PJ260009",   # Air Mineral Outlet 600 ml
    "BB0328": "SG260012",   # Jilbab Outlet          [BARU 12 Okt, name match LAMA/BARU]
    "BB0329": "BB260030",   # Jasmine Green Tea (750 g)
    "BB0340": "BB260031",   # Coffee Latte (500 g)
    "BB0341": "BB260032",   # Coffee Classic (120 g)
    "BB0342": "BB260033",   # Dark Cocoa Drink (1 kg)
    "BB0352": "BB260046",   # Olesan Pastry 250g (lihat aturan tanggal)
    "BB0353": "BB26009",    # Gula Cair 250g -> 'Gula Cair (250 g)' (Botol)
}

# (item_code, uom_eresto) -> (uom_erpnext, cf, catatan). Kosong = fallback otomatis.
UOM_SPECIAL = {
    ("PEL260071", "Pack"): (
        "slop", 50.0,
        "Gelas Cup 9 oz: 1 Pack eResto = 1 slop = 50 Pcs "
        "(preseden PO120.051026.001 / rate 55.000 = Item Price 1.100/Pcs x 50)",
    ),
}

# ============================================================
# KLIEN eRESTO (read-only)
# ============================================================

def safe_float(v):
    try:
        if v is None or v == "":
            return 0.0
        if isinstance(v, str):
            return float(v.strip().replace(",", "."))
        return float(v)
    except Exception:
        return 0.0


def safe_str(v):
    return "" if v is None else str(v).strip()


class ErestoClient:
    def __init__(self, base_url, token, email, user_id):
        self.base_url = base_url.rstrip("/")
        self.token = token
        self.email = email
        self.user_id = user_id
        self.session = requests.Session()
        self.session.headers.update({
            "accept": "application/json, text/plain, */*",
            "authorization": f"Token token={self.token}, email={self.email}",
            "referer": f"{self.base_url}/manager/v2/index.html",
        })

    def _fetch_paged(self, window):
        out = {}
        page = 1
        while True:
            params = dict(window)
            params.update({"current_user_id": self.user_id, "domain": ERESTO_DOMAIN,
                           "token": self.token, "per_page": 200, "page": page})
            res = self.session.get(f"{self.base_url}/v2/purchases", params=params,
                                   timeout=90)
            if res.status_code != 200:
                raise SystemExit(f"[eResto] GET /v2/purchases HTTP "
                                 f"{res.status_code}: {res.text[:200]}")
            rows = res.json().get("purchases", [])
            for p in rows:
                out[p["id"]] = p
            if len(rows) < 200:
                break
            page += 1
        return out

    def fetch_purchases(self, date_from, date_to):
        """Union window date[] dan created_at[] (PO backdate / input beda hari)."""
        results = self._fetch_paged({"date[geq]": date_from, "date[leq]": date_to})
        results.update(self._fetch_paged({
            "created_at[geq]": f"{date_from}T00:00:00",
            "created_at[leq]": f"{date_to}T23:59:59",
        }))
        return list(results.values())

    def fetch_purchase_detail(self, purchase_id):
        params = {"domain": ERESTO_DOMAIN, "token": self.token}
        res = self.session.get(f"{self.base_url}/v2/purchases/{purchase_id}/detail",
                               params=params, timeout=90)
        return res.json().get("purchase_detail", {}) if res.status_code == 200 else {}

# ============================================================
# PENDUKUNG ERPNext
# ============================================================

def resolve_uom(code, unit, master):
    """(uom, cf, catatan). Tidak pernah menebak diam-diam: fallback selalu dicatat."""
    su = master["stock_uom"]
    uoms = master["uoms"]
    unit_norm = {"Kg": "kg", "KG": "kg"}.get(unit, unit)
    if unit_norm == su:
        return unit_norm, uoms.get(su, 1.0), ""
    if unit_norm in uoms:
        return unit_norm, uoms[unit_norm], ""
    special = UOM_SPECIAL.get((code, unit_norm))
    if special:
        return special[0], special[1], special[2]
    return su, uoms.get(su, 1.0), \
        f"UOM '{unit}' tidak ada di master -> fallback {su} cf=1"


def get_item_master(client, code):
    st, body = client.get(f"/api/resource/Item/{urllib.parse.quote(code)}")
    if st != 200:
        return None
    d = body.get("data", body)
    return {
        "name": d.get("name"),
        "item_name": d.get("item_name"),
        "stock_uom": d.get("stock_uom"),
        "disabled": bool(d.get("disabled")),
        "uoms": {u["uom"]: float(u["conversion_factor"]) for u in d.get("uoms", [])},
    }


def get_items_master(client, codes):
    out = {}
    with ThreadPoolExecutor(max_workers=5) as ex:
        futs = {ex.submit(get_item_master, client, c): c for c in sorted(set(codes))}
        for fu in as_completed(futs):
            out[futs[fu]] = fu.result()
    return out


def get_all(client, doctype, fields, page_size=500):
    out, offset = [], 0
    while True:
        st, d = client.get(f"/api/resource/{urllib.parse.quote(doctype)}",
                           params={"fields": json.dumps(fields),
                                   "limit_page_length": page_size,
                                   "limit_start": offset})
        if st != 200:
            raise SystemExit(f"[ERPNext] GET {doctype} HTTP {st}: {str(d)[:200]}")
        rows = d.get("data", [])
        out.extend(rows)
        if len(rows) < page_size:
            break
        offset += page_size
    return out


def get_existing_po_index(client):
    """{name: row} + alias title 'ERESTO <PO>' -> row."""
    index = {}
    for row in get_all(client, "Purchase Order",
                       ["name", "title", "docstatus", "grand_total", "status"]):
        index[row["name"]] = row
        title = row.get("title") or ""
        if title.startswith("ERESTO "):
            index[title.replace("ERESTO ", "").strip()] = row
    return index


def get_po(client, name):
    st, body = client.get(f"/api/resource/Purchase Order/{urllib.parse.quote(name)}")
    if st != 200:
        return None
    return body.get("data", body)

# ============================================================
# TARIK / DUMP
# ============================================================

def fetch_targets(date_from, date_to):
    eresto = ErestoClient(ERESTO_BASE_URL, ERESTO_TOKEN, ERESTO_EMAIL, ERESTO_USER_ID)
    raw = eresto.fetch_purchases(date_from, date_to)
    internal = [p for p in raw if p.get("internal_supplier")]
    details = []
    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = {ex.submit(eresto.fetch_purchase_detail, p["id"]): p for p in internal}
        for fu in as_completed(futs):
            d = fu.result()
            if d:
                details.append(d)
    details.sort(key=lambda x: safe_str(x.get("purchase_number")))
    return raw, details


def run_dump(args):
    print("=" * 84)
    print(f"  DUMP PO internal JURI eResto {args.date_from}..{args.date_to}")
    print("=" * 84)
    raw, details = fetch_targets(args.date_from, args.date_to)
    os.makedirs(DUMP_DIR, exist_ok=True)
    with open(DUMP_FILE, "w") as f:
        json.dump(details, f, ensure_ascii=False, indent=1)
    print(f"  {len(raw)} PO di jendela | {len(details)} internal_supplier")
    print(f"  dump -> {DUMP_FILE}")


def load_details(refresh=False, date_from=DATE_FROM_DEFAULT, date_to=DATE_TO_DEFAULT):
    if not refresh and os.path.exists(DUMP_FILE):
        with open(DUMP_FILE) as f:
            return json.load(f)
    _, details = fetch_targets(date_from, date_to)
    return details

# ============================================================
# MIGRASI
# ============================================================

def write_sheet(ws, title, rows, columns):
    ws.title = title
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c, "") for c in columns])
    ws.freeze_panes = "A2"
    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(left=Side(style="thin", color="D3D3D3"),
                  right=Side(style="thin", color="D3D3D3"),
                  top=Side(style="thin", color="D3D3D3"),
                  bottom=Side(style="thin", color="D3D3D3"))
    for ci in range(1, ws.max_column + 1):
        c = ws.cell(row=1, column=ci)
        c.fill, c.font, c.alignment = header_fill, header_font, center
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for c in row:
            c.border = thin
    for col in ws.columns:
        ml = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(ml + 3, 12)


def build_payload(d, abbr, items_master, company_actual=None):
    """Kembalikan (payload, row_audit, problems)."""
    company = company_actual or safe_str((d.get("division") or {}).get("name"))
    pdate = safe_str(d.get("date"))[:10]
    source_wh = WAREHOUSE_MAP.get(safe_str((d.get("supplier") or {}).get("name")))
    exp_acct = f"4210.000 - HPP Pembelian - {abbr}"
    ctr = f"Main - {abbr}"

    rows_payload, row_audit, problems = [], [], []
    for it in d.get("purchase_items") or []:
        inv = it.get("inventory") or {}
        sku = safe_str(inv.get("sku"))
        qty = safe_float(it.get("amount"))
        price = safe_float(it.get("price"))
        unit = safe_str((it.get("unit") or {}).get("name"))
        code = ITEM_MAP.get(sku)
        if not code:
            problems.append(f"SKU {sku} ({safe_str(inv.get('name'))}) belum ada mapping")
            continue
        # ATURAN OLESAN PASTRY (keputusan user 2026-10-12): PO ke PT. JURI
        # bertanggal > 2026-10-08 -> BB26012 'Olesan Pastry (250 g)'.
        if sku in ("BB0174", "BB0352"):
            code = olesan_pastry_item(pdate, default=code)
        master = items_master.get(code)
        if not master:
            problems.append(f"Item {code} tidak ada di master BARU")
            continue
        if master.get("disabled"):
            problems.append(f"Item {code} disabled")
            continue
        if qty <= 0:
            problems.append(f"SKU {sku} qty {qty}")
            continue
        uom, cf, note = resolve_uom(code, unit, master)
        rows_payload.append({
            "item_code": code, "qty": qty, "uom": uom, "conversion_factor": cf,
            "rate": price, "warehouse": f"Stores - {abbr}",
            "expense_account": exp_acct, "cost_center": ctr, "schedule_date": pdate,
        })
        row_audit.append({
            "PO Number": safe_str(d.get("purchase_number")), "Tanggal": pdate,
            "Company (Divisi)": company, "SKU": sku,
            "Nama eResto": safe_str(inv.get("name")), "Qty eResto": qty,
            "Unit eResto": unit, "Harga eResto": price, "Item Code": code,
            "Item Name": master.get("item_name"), "UOM ERPNext": uom, "CF": cf,
            "Stock Qty": qty * cf, "Rate": price, "Amount": round(qty * price, 2),
            "Catatan UOM": note, "Status eResto": safe_str(d.get("status")),
        })

    payload = {
        "doctype": "Purchase Order",
        "company": company,
        "supplier": JURI,
        "transaction_date": pdate,
        "schedule_date": pdate,
        "set_warehouse": f"Stores - {abbr}",
        "buying_price_list": PRICE_LIST,
        "currency": "IDR",
        "title": f"ERESTO {safe_str(d.get('purchase_number'))}",
        "items": rows_payload,
    }
    if source_wh:
        payload["set_from_warehouse"] = source_wh
    return payload, row_audit, problems


def run_migrate(args):
    erp = Client(ERPNEXT_URL, ERPNEXT_KEY, label="BARU", pause=0.15)

    details = load_details(refresh=args.refresh, date_from=args.date_from,
                           date_to=args.date_to)
    if args.only:
        only = set(args.only)
        details = [d for d in details if d.get("purchase_number") in only]

    print("=" * 84)
    print("  MIGRASI PO eResto (internal JURI) -> ERPNext BARU")
    print(f"  Jendela: {args.date_from}..{args.date_to} | Target: {ERPNEXT_URL}")
    print("=" * 84)
    print(f"  {len(details)} PO internal_supplier di jendela.")

    existing = get_existing_po_index(erp)
    abbrs = {c["name"]: c.get("abbr") for c in get_all(erp, "Company", ["name", "abbr"])}
    # Resolusi nama company case-insensitive: eResto 'HJ.' vs ERPNext 'Hj.' dsb.
    _by_lower = {}
    for _n in abbrs:
        _by_lower.setdefault(_n.lower(), _n)

    # Alias eksplisit nama divisi eResto -> nama Company ERPNext (beda ejaan).
    COMPANY_ALIAS = {
        # eResto 'MUHAMMAD YAMIN' vs ERPNext 'MUHAMMADIYAH YAMIN' (terverifikasi
        # Company RSPY dibuat 13 Okt 2026 via onboard_company_rs_12okt.py)
        "RSUD PROF H MUHAMMAD YAMIN SH": "RSUD PROF H MUHAMMADIYAH YAMIN SH",
    }

    def resolve_company(name):
        """(nama_erpnext_aktual, abbr) — cocokkan case-insensitive bila exact gagal."""
        name = COMPANY_ALIAS.get(name, name)
        if name in abbrs:
            return name, abbrs[name]
        actual = _by_lower.get(name.lower())
        return (actual, abbrs.get(actual)) if actual else (name, None)

    allowed = set()
    st, sup = erp.get(f"/api/resource/Supplier/{urllib.parse.quote(JURI)}",
                      params={"fields": json.dumps(["name", "companies"])})
    if st == 200:
        body = sup.get("data", sup)
        allowed = {x.get("company") for x in (body.get("companies") or []) if x.get("company")}
    _allowed_lower = {a.lower() for a in allowed}

    def is_allowed(name):
        return name in allowed or name.lower() in _allowed_lower

    all_codes = set()
    for d in details:
        _pd = safe_str(d.get("date"))[:10]
        for it in d.get("purchase_items") or []:
            _sku = safe_str((it.get("inventory") or {}).get("sku"))
            code = ITEM_MAP.get(_sku)
            if code and _sku in ("BB0174", "BB0352"):
                code = olesan_pastry_item(_pd, default=code)
            if code:
                all_codes.add(code)
    items_master = get_items_master(erp, all_codes)
    print(f"  {len(existing)} PO existing | {len(abbrs)} company | "
          f"{len(allowed)} company izin supplier | {len(items_master)} item master.")

    migrated, skipped, raw_rows = [], [], []

    for d in details:
        pnum = safe_str(d.get("purchase_number"))
        pdate = safe_str(d.get("date"))[:10]
        company = safe_str((d.get("division") or {}).get("name"))
        st_eresto = safe_str(d.get("status"))
        total_cost = safe_float(d.get("total_cost"))
        base = {"PO Number": pnum, "Tanggal": pdate, "Company": company,
                "Total eResto": total_cost, "Status eResto": st_eresto}

        # -- skip status rejected/canceled (keputusan user) --
        if st_eresto in SKIP_STATUS:
            skipped.append({**base, "Kategori": "Status eResto",
                            "Alasan": f"status eResto '{st_eresto}' -> tidak dimigrasi"})
            continue

        # -- idempotensi / self-heal rename --
        if pnum in existing:
            rec = existing[pnum]
            actual = rec.get("name")
            note = ""
            if actual != pnum and not args.dry_run:
                st_ren, body = erp.rename("Purchase Order", actual, pnum)
                if st_ren == 200:
                    print(f"  [REPAIR] {pnum}: rename {actual} -> {pnum}")
                    note = f"repair rename {actual} -> {pnum}"
                    rec = get_po(erp, pnum) or rec
                    actual = pnum
                else:
                    note = f"rename gagal: {str(body)[:120]}"
            print(f"  [SKIP_EXISTS] {pnum} -> {actual} (ds={rec.get('docstatus')})")
            migrated.append({**base, "ERPNext ID": actual, "Supplier": JURI,
                             "Target WH": f"Stores - {abbrs.get(company, '?')}",
                             "Source WH": WAREHOUSE_MAP.get(
                                 safe_str((d.get("supplier") or {}).get("name")), ""),
                             "Jumlah Baris": "", "Grand Total ERPNext": rec.get("grand_total"),
                             "Docstatus": rec.get("docstatus"),
                             "Status ERPNext": rec.get("status"),
                             "Hasil": "SKIP_EXISTS", "Catatan": note or "sudah ada"})
            continue

        # -- guard prasyarat --
        company_actual, abbr = resolve_company(company)
        problems = []
        if not abbr:
            problems.append(f"Company '{company}' belum ada di BARU")
        if not is_allowed(company_actual):
            problems.append("Company belum diizinkan di supplier JURI")
        if abbr:
            for dt, nm in (("Warehouse", f"Stores - {abbr}"),
                           ("Account", f"4210.000 - HPP Pembelian - {abbr}"),
                           ("Cost Center", f"Main - {abbr}")):
                st_c, _ = erp.get(f"/api/resource/{urllib.parse.quote(dt)}/"
                                  f"{urllib.parse.quote(nm)}")
                if st_c != 200:
                    problems.append(f"{dt} '{nm}' tidak ada")
        if problems:
            skipped.append({**base, "Kategori": "Prasyarat", "Alasan": "; ".join(problems)})
            continue

        # -- susun baris --
        payload, row_audit, row_problems = build_payload(d, abbr, items_master,
                                                         company_actual=company_actual)
        if row_problems or not payload["items"]:
            skipped.append({**base, "Kategori": "Item",
                            "Alasan": "; ".join(row_problems) or "tidak ada baris valid"})
            continue
        raw_rows.extend(row_audit)

        expected_ds = 1 if st_eresto in SUBMIT_STATUS else 0

        if args.dry_run:
            print(f"  [DRY-RUN] {pnum} -> {company} ({len(payload['items'])} baris, "
                  f"{st_eresto}, target ds={expected_ds})")
            migrated.append({**base, "ERPNext ID": "(dry-run)", "Supplier": JURI,
                             "Target WH": f"Stores - {abbr}",
                             "Source WH": payload.get("set_from_warehouse", ""),
                             "Jumlah Baris": len(payload["items"]),
                             "Grand Total ERPNext": "", "Docstatus": "",
                             "Status ERPNext": "", "Hasil": "DRY_RUN", "Catatan": ""})
            continue

        # -- POST + rename + submit --
        print(f"  [POST] {pnum} -> {company} ({len(payload['items'])} baris) ...")
        st_p, res = erp.post("Purchase Order", payload)
        if st_p != 200:
            msg = json.dumps(res)[:300]
            print(f"    [GAGAL] {pnum}: {msg}")
            migrated.append({**base, "ERPNext ID": "", "Supplier": JURI,
                             "Target WH": f"Stores - {abbr}",
                             "Source WH": payload.get("set_from_warehouse", ""),
                             "Jumlah Baris": len(payload["items"]),
                             "Grand Total ERPNext": "", "Docstatus": 0,
                             "Status ERPNext": "GAGAL", "Hasil": "FAIL_POST",
                             "Catatan": msg})
            continue

        temp = (res.get("data") or {}).get("name")
        final, note = temp, ""
        st_ren, _ = erp.rename("Purchase Order", temp, pnum)
        if st_ren == 200:
            final = pnum
        else:
            note = f"rename gagal -> tetap '{temp}'"

        ds_final, status_final = 0, "Draft"
        if expected_ds == 1:
            st_s, res_s = erp.submit("Purchase Order", final)
            if st_s == 200:
                ds_final = 1
                status_final = (res_s.get("data") or {}).get("status", "Submitted")
            else:
                note = (note + " | " if note else "") + \
                    f"submit gagal: {json.dumps(res_s)[:120]}"

        rec = get_po(erp, final) or {}
        print(f"    [OK] {final} | ds={ds_final} | {status_final} | "
              f"Rp {rec.get('grand_total', '?')}")
        migrated.append({**base, "ERPNext ID": final, "Supplier": JURI,
                         "Target WH": f"Stores - {abbr}",
                         "Source WH": payload.get("set_from_warehouse", ""),
                         "Jumlah Baris": len(payload["items"]),
                         "Grand Total ERPNext": rec.get("grand_total", ""),
                         "Docstatus": ds_final, "Status ERPNext": status_final,
                         "Hasil": "BERHASIL", "Catatan": note})

    # ---------- verifikasi ----------
    ver_rows, verify_issues = [], []
    if not args.dry_run:
        for m in migrated:
            if m["Hasil"] not in ("BERHASIL", "SKIP_EXISTS"):
                continue
            rec = get_po(erp, m["ERPNext ID"])
            if not rec:
                verify_issues.append(f"{m['PO Number']}: dokumen hilang")
                continue
            exp = m["Total eResto"]
            got = safe_float(rec.get("grand_total"))
            diff = abs(got - exp)
            line_ok = all(
                abs(safe_float(x.get("qty")) * safe_float(x.get("rate"))
                    - safe_float(x.get("amount"))) < 0.01
                and abs(safe_float(x.get("stock_qty"))
                        - safe_float(x.get("qty")) * safe_float(x.get("conversion_factor"))) < 1e-6
                for x in rec.get("items", []))
            ok = diff <= 1.0 and line_ok
            if not ok:
                verify_issues.append(f"{m['PO Number']}: diff Rp{diff:,.2f} line_ok={line_ok}")
            ver_rows.append({"PO": m["PO Number"], "ERPNext ID": m["ERPNext ID"],
                             "Total eResto": exp, "Grand Total ERPNext": got,
                             "Selisih": diff, "Baris": len(rec.get("items", [])),
                             "Line OK": line_ok, "Status": "LOLOS" if ok else "CEK"})

    # ---------- Excel ----------
    wb = Workbook()
    rekap_cols = ["PO Number", "ERPNext ID", "Tanggal", "Company", "Supplier",
                  "Target WH", "Source WH", "Jumlah Baris", "Total eResto",
                  "Grand Total ERPNext", "Docstatus", "Status ERPNext", "Hasil",
                  "Status eResto", "Catatan"]
    raw_cols = ["PO Number", "Tanggal", "Company (Divisi)", "SKU", "Nama eResto",
                "Qty eResto", "Unit eResto", "Harga eResto", "Item Code", "Item Name",
                "UOM ERPNext", "CF", "Stock Qty", "Rate", "Amount", "Catatan UOM",
                "Status eResto"]
    skip_cols = ["PO Number", "Tanggal", "Company", "Total eResto", "Status eResto",
                 "Kategori", "Alasan"]
    ver_cols = ["PO", "ERPNext ID", "Total eResto", "Grand Total ERPNext", "Selisih",
                "Baris", "Line OK", "Status"]

    write_sheet(wb.active, "Rekap_Migrasi", migrated, rekap_cols)
    write_sheet(wb.create_sheet(), "Raw_eResto", raw_rows, raw_cols)
    write_sheet(wb.create_sheet(), "Skipped", skipped, skip_cols)
    write_sheet(wb.create_sheet(), "Verifikasi", ver_rows, ver_cols)

    n_ok = sum(1 for m in migrated if m["Hasil"] == "BERHASIL")
    n_skip = sum(1 for m in migrated if m["Hasil"] == "SKIP_EXISTS")
    n_fail = sum(1 for m in migrated if m["Hasil"] == "FAIL_POST")
    n_dry = sum(1 for m in migrated if m["Hasil"] == "DRY_RUN")
    tot_e = sum(m["Total eResto"] for m in migrated if m["Hasil"] == "BERHASIL")
    tot_g = sum(safe_float(m["Grand Total ERPNext"]) for m in migrated
                if m["Hasil"] == "BERHASIL")
    ds_ok = all(m["Docstatus"] == (1 if m["Status eResto"] in SUBMIT_STATUS else 0)
                for m in migrated if m["Hasil"] == "BERHASIL")
    val_rows = [
        {"Parameter": "PO internal JURI di jendela", "Nilai": len(details), "Status": "OK",
         "Keterangan": f"{args.date_from}..{args.date_to}"},
        {"Parameter": "BERHASIL dimigrasikan", "Nilai": n_ok,
         "Status": "DRY-RUN" if args.dry_run else ("LOLOS" if n_fail == 0 else "CEK"),
         "Keterangan": "ID identik eResto (rename_doc)"},
        {"Parameter": "SKIP_EXISTS (rerun idempoten)", "Nilai": n_skip, "Status": "OK",
         "Keterangan": "0 duplikat"},
        {"Parameter": "DRY_RUN", "Nilai": n_dry, "Status": "OK", "Keterangan": ""},
        {"Parameter": "GAGAL POST", "Nilai": n_fail,
         "Status": "LOLOS" if n_fail == 0 else "CEK", "Keterangan": "harus 0"},
        {"Parameter": "Di-skip (status/prasyarat/item)", "Nilai": len(skipped),
         "Status": "OK", "Keterangan": "lihat sheet Skipped"},
        {"Parameter": "Total nilai eResto vs ERPNext",
         "Nilai": f"{tot_e:,.2f} vs {tot_g:,.2f}",
         "Status": "LOLOS" if abs(tot_e - tot_g) <= 1 else "CEK",
         "Keterangan": "toleransi <= Rp1"},
        {"Parameter": "Docstatus sesuai status eResto", "Nilai": "ya" if ds_ok else "tidak",
         "Status": "LOLOS" if ds_ok else "CEK",
         "Keterangan": "processing/completed -> submit; pending -> draft"},
        {"Parameter": "Verifikasi baris (qty×rate=amount, stock_qty=qty×cf)",
         "Nilai": f"{len(ver_rows)} dok",
         "Status": "LOLOS" if not verify_issues else "CEK",
         "Keterangan": "; ".join(verify_issues) if verify_issues else "semua klop"},
    ]
    write_sheet(wb.create_sheet(), "Validation", val_rows,
                ["Parameter", "Nilai", "Status", "Keterangan"])

    note_rows = [
        {"Topik": "Scope",
         "Catatan": f"PO eResto internal_supplier (gudang JURI) {args.date_from}..{args.date_to}; "
                    f"company = divisi outlet pemesan; supplier = '{JURI}'; "
                    f"set_warehouse 'Stores - <ABBR>'; set_from_warehouse gudang JURI asal."},
        {"Topik": "Status rejected/canceled",
         "Catatan": "Keputusan user 2026-10-12: PO rejected/canceled di eResto TIDAK dimigrasi."},
        {"Topik": "FRC di luar scope",
         "Catatan": "PO divisi 'Jogokariyan (FRC)', 'Pasaman (FRC)', 'Ugm (FRC)' di-skip — "
                    "company belum ada di BARU (butuh onboarding)."},
        {"Topik": "Mapping SKU baru (bukti)",
         "Catatan": "BB0006->PEL260006, BB0042->BB260048, BB0076->PEL260068, BB0078->PEL260019, "
                    "BB0160->PEL260052, BB0185->PEL260010, BB0270->PEL260023 (padanan eResto PO332 "
                    "di site LAMA PO332.220926.001/.002); BB0328->SG260012 (Jilbab); "
                    "BB0353->BB26009 'Gula Cair (250 g)' Botol. BB0043 Daun Teh belum ada padanan "
                    "(hanya PO FRC Pasaman) — dibiarkan unmapped."},
    ]
    write_sheet(wb.create_sheet(), "Catatan", note_rows, ["Topik", "Catatan"])

    wb.save(REPORT_PATH)
    print(f"\n  Laporan: {REPORT_PATH}")
    print("=" * 84)
    print(f"  BERHASIL: {n_ok} | SKIP_EXISTS: {n_skip} | GAGAL: {n_fail} | "
          f"SKIPPED: {len(skipped)} | DRY-RUN: {n_dry}")
    if not args.dry_run:
        print(f"  Nilai eResto : Rp {tot_e:,.2f}")
        print(f"  Nilai ERPNext: Rp {tot_g:,.2f}")
        print(f"  Docstatus OK : {ds_ok} | Verifikasi: "
              f"{'OK' if not verify_issues else verify_issues}")
    print("=" * 84)


# ============================================================
# VERIFY (read-only, kondisi live)
# ============================================================

def run_verify(args):
    erp = Client(ERPNEXT_URL, ERPNEXT_KEY, label="BARU", pause=0.0)
    details = load_details(refresh=args.refresh, date_from=args.date_from,
                           date_to=args.date_to)
    if args.only:
        only = set(args.only)
        details = [d for d in details if d.get("purchase_number") in only]
    existing = get_existing_po_index(erp)

    print("=" * 84)
    print(f"  VERIFIKASI PO eResto internal JURI {args.date_from}..{args.date_to} vs BARU")
    print("=" * 84)
    rows, issues = [], []
    for d in sorted(details, key=lambda x: safe_str(x.get("purchase_number"))):
        pnum = safe_str(d.get("purchase_number"))
        st_eresto = safe_str(d.get("status"))
        exp = safe_float(d.get("total_cost"))
        doc = get_po(erp, pnum)
        if st_eresto in SKIP_STATUS:
            rows.append({"PO": pnum, "Status eResto": st_eresto, "Hasil": "SKIP_STATUS",
                         "Total eResto": exp, "Grand Total ERPNext": "", "Selisih": ""})
            continue
        if not doc:
            rows.append({"PO": pnum, "Status eResto": st_eresto, "Hasil": "TIDAK ADA",
                         "Total eResto": exp, "Grand Total ERPNext": "", "Selisih": ""})
            issues.append(f"{pnum}: tidak ada di BARU")
            continue
        got = safe_float(doc.get("grand_total"))
        diff = abs(got - exp)
        exp_ds = 1 if st_eresto in SUBMIT_STATUS else 0
        ok = diff <= 1 and doc.get("docstatus") == exp_ds
        if not ok:
            issues.append(f"{pnum}: diff={diff:.0f} ds={doc.get('docstatus')}!= {exp_ds}")
        rows.append({"PO": pnum, "Status eResto": st_eresto,
                     "Hasil": "LOLOS" if ok else "CEK", "Total eResto": exp,
                     "Grand Total ERPNext": got, "Selisih": diff,
                     "Docstatus": doc.get("docstatus"),
                     "Company": doc.get("company")})
    print(f"  {len(rows)} PO diperiksa | issues: {issues if issues else 'tidak ada'}")
    for r in rows:
        print(f"   {r['PO']:22} {r['Status eResto']:11} {r['Hasil']:12} "
              f"Rp {safe_float(r['Total eResto']):>14,.0f} vs {r['Grand Total ERPNext']}")


# ============================================================
# CLI
# ============================================================

def run():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--date-from", default=DATE_FROM_DEFAULT)
    ap.add_argument("--date-to", default=DATE_TO_DEFAULT)
    ap.add_argument("--dump", action="store_true", help="tarik ulang eResto -> dump")
    ap.add_argument("--refresh", action="store_true", help="pakai tarikan eResto live")
    ap.add_argument("--dry-run", action="store_true", help="rencana saja, tanpa POST")
    ap.add_argument("--verify", action="store_true", help="mode verifikasi read-only")
    ap.add_argument("--only", nargs="*", default=None, help="batasi nomor PO tertentu")
    args = ap.parse_args()

    if args.dump:
        run_dump(args)
        return
    if args.verify:
        run_verify(args)
        return
    run_migrate(args)


if __name__ == "__main__":
    run()
