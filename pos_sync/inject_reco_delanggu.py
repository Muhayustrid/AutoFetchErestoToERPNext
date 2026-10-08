"""Revisi Stock Reconciliation Opening Stock untuk RS PKU MUHAMMADIYAH DELANGGU.

Sumber Data: Stock Opname Fisik eResto 30 September 2026 (36 baris).
Target: https://ropierpnext.frappe.cloud
Company: RS PKU MUHAMMADIYAH DELANGGU
Warehouse: Stores - RPMD
Posting Date: 2026-09-30 18:00:00
Purpose: Opening Stock
Expense Account: 1171.000 - Pembukaan sementara - RPMD
Cost Center: Main - RPMD

Aturan Sesuai Instruksi User:
1. Ropi Coklat, Ropi Keju, Ropi Butter, Roti Gempi, Croissant adalah bundle dari matang (PJO*).
2. Yang tidak aktif (BB0106 Olesan Pineapple, BB0108 Olesan Strawberry) dilewati (skip).
3. Item dengan After Adj Qty = 0 dilewati (skip) agar tidak mengotori buku besar dengan SLE nol.
4. Nilai After Adj Qty dimasukkan ke custom_qty_after & dikonversi ke qty stok.
5. Nilai Cost Per Unit dimasukkan ke custom_valuation_rate_per_uom & dikonversi ke valuation_rate stok.
6. Mengisi field standar akuntansi dan field kustom UI (warehouse_app dual-field population).
"""

import os
import sys
import time
import requests
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE_URL = "https://ropierpnext.frappe.cloud"
API_KEY = os.environ.get("ERPNEXT_API_KEY", "")

COMPANY = "RS PKU MUHAMMADIYAH DELANGGU"
WAREHOUSE = "Stores - RPMD"
POSTING_DATE = "2026-09-30"
POSTING_TIME = "18:00:00"
EXPENSE_ACCOUNT = "1171.000 - Pembukaan sementara - RPMD"
COST_CENTER = "Main - RPMD"
EXCEL_OUT = "pos_sync/Injeksi_Reco_Delanggu.xlsx"

RAW_DATA = [
    # 1. BB0312 Air Mineral Outlet 600 ml -> PJ260009
    {"sku": "BB0312", "name_eresto": "Air Mineral Outlet 600 ml", "item_code": "PJ260009", "after_qty": 6.0, "unit": "Botol", "cost": 2354.17, "conv": 1.0},
    # 2. BB0171 Box Gempi (Isi 100 pcs) -> PEL260069
    {"sku": "BB0171", "name_eresto": "Box Gempi (Isi 100 pcs)", "item_code": "PEL260069", "after_qty": 71.0, "unit": "Pcs", "cost": 3500.00, "conv": 1.0},
    # 3. BB0311 Bubuk Hepi Drink Chocolate -> Qty 0 -> Skip
    {"sku": "BB0311", "name_eresto": "Bubuk Hepi Drink Chocolate", "item_code": "BB260060", "after_qty": 0.0, "unit": "Sachet", "cost": 1895.90, "conv": 1.0, "skip": True, "reason": "Qty 0"},
    # 4. BB0310 Bubuk Hepi Drink Kopi Karamel -> Qty 0 -> Skip
    {"sku": "BB0310", "name_eresto": "Bubuk Hepi Drink Kopi Karamel (30 Gr)", "item_code": "BB260059", "after_qty": 0.0, "unit": "Sachet", "cost": 5360.00, "conv": 1.0, "skip": True, "reason": "Qty 0"},
    # 5. BB0305 Gula Cair Sachet (Isi 100 Sachet) -> BB260038
    {"sku": "BB0305", "name_eresto": "Gula Cair Sachet (Isi 100 Sachet)", "item_code": "BB260038", "after_qty": 84.0, "unit": "Sachet", "cost": 1000.00, "conv": 1.0},
    # 6. BB0112 Paper Bag Kapasitas 1 Roti -> PEL260076
    {"sku": "BB0112", "name_eresto": "Paper Bag Kapasitas 1 Roti", "item_code": "PEL260076", "after_qty": 510.0, "unit": "Pcs", "cost": 310.00, "conv": 1.0},
    # 7. RR0002 Ropi Coklat -> PJO2601 -> Qty 0 -> Skip
    {"sku": "RR0002", "name_eresto": "Ropi Coklat", "item_code": "PJO2601", "after_qty": 0.0, "unit": "Pcs", "cost": 3750.00, "conv": 1.0, "skip": True, "reason": "Qty 0"},
    # 8. RR0008 Roti Croissant -> PJO2605 (Croissant Matang)
    {"sku": "RR0008", "name_eresto": "Roti Croissant", "item_code": "PJO2605", "after_qty": 2.0, "unit": "Pcs", "cost": 7833.00, "conv": 1.0},
    # 9. BB0071 Gula Stik -> BB260039
    {"sku": "BB0071", "name_eresto": "Gula Stik", "item_code": "BB260039", "after_qty": 389.0, "unit": "Pcs", "cost": 280.00, "conv": 1.0},
    # 10. BB0298 Gelas 14 Oz (Isi 150 pcs) -> PEL260070
    {"sku": "BB0298", "name_eresto": "Gelas 14 Oz (Isi 150 pcs)", "item_code": "PEL260070", "after_qty": 943.0, "unit": "Pcs", "cost": 745.63, "conv": 1.0},
    # 11. RR0001 Ropi Butter -> PJO2603 -> Qty 0 -> Skip
    {"sku": "RR0001", "name_eresto": "Ropi Butter", "item_code": "PJO2603", "after_qty": 0.0, "unit": "Pcs", "cost": 3750.00, "conv": 1.0, "skip": True, "reason": "Qty 0"},
    # 12. RR0009 Roti Gempi -> PJO2604 (Gempi Matang)
    {"sku": "RR0009", "name_eresto": "Roti Gempi", "item_code": "PJO2604", "after_qty": 2.0, "unit": "Pcs", "cost": 5690.68, "conv": 1.0},
    # 13. BB0329 Jasmine Green Tea 750g -> BB260030 (konversi Pack -> g, 1 Pack = 750 g)
    {"sku": "BB0329", "name_eresto": "Jasmine Green Tea 750g", "item_code": "BB260030", "after_qty": 1.46, "unit": "Pack", "cost": 86682.20, "conv": 750.0},
    # 14. BB0085 Krim Kopi -> PJ260016
    {"sku": "BB0085", "name_eresto": "Krim Kopi", "item_code": "PJ260016", "after_qty": 16.0, "unit": "Pcs", "cost": 25999.99, "conv": 1.0},
    # 15. BB0308 Bubuk Hepi Drink Moccacino (30 Gr) -> Qty 0 -> Skip
    {"sku": "BB0308", "name_eresto": "Bubuk Hepi Drink Moccacino (30 Gr)", "item_code": "BB260057", "after_qty": 0.0, "unit": "Sachet", "cost": 5360.00, "conv": 1.0, "skip": True, "reason": "Qty 0"},
    # 16. BB0341 Coffee Classic 120g -> BB260032 (konversi Pack -> g, 1 Pack = 120 g)
    {"sku": "BB0341", "name_eresto": "Coffee Classic 120g", "item_code": "BB260032", "after_qty": 2.1, "unit": "Pack", "cost": 64815.13, "conv": 120.0},
    # 17. BB0340 Coffee Latte 500g -> BB260031 (konversi Pack -> g, 1 Pack = 500 g)
    {"sku": "BB0340", "name_eresto": "Coffee Latte 500g", "item_code": "BB260031", "after_qty": 5.47, "unit": "Pack", "cost": 72360.00, "conv": 500.0},
    # 18. BB0342 Dark Cocoa Drink 1Kg -> BB260033 (konversi Pack -> g, 1 Pack = 1000 g)
    {"sku": "BB0342", "name_eresto": "Dark Cocoa Drink 1Kg", "item_code": "BB260033", "after_qty": 1.35, "unit": "Pack", "cost": 106800.00, "conv": 1000.0},
    # 19. BB0046 Dough Butter -> PJ260012 (konversi Pack -> Pcs, 1 Pack = 12 Pcs)
    {"sku": "BB0046", "name_eresto": "Dough Butter", "item_code": "PJ260012", "after_qty": 17.0, "unit": "Pack", "cost": 32000.00, "conv": 12.0},
    # 20. BB0047 Dough Coklat -> PJ260013 (konversi Pack -> Pcs, 1 Pack = 12 Pcs)
    {"sku": "BB0047", "name_eresto": "Dough Coklat", "item_code": "PJ260013", "after_qty": 7.0, "unit": "Pack", "cost": 32000.00, "conv": 12.0},
    # 21. BB0168 Dough Croissant -> PJ260010 (konversi Pack -> Pcs, 1 Pack = 6 Pcs)
    {"sku": "BB0168", "name_eresto": "Dough Croissant", "item_code": "PJ260010", "after_qty": 27.0, "unit": "Pack", "cost": 45000.00, "conv": 6.0},
    # 22. BB0261 Dough Gempi -> PJ260011 (konversi Pack -> Pcs, 1 Pack = 3 Pcs)
    {"sku": "BB0261", "name_eresto": "Dough Gempi", "item_code": "PJ260011", "after_qty": 18.0, "unit": "Pack", "cost": 17072.05, "conv": 3.0},
    # 23. BB0049 Dough Keju -> PJ260014 (konversi Pack -> Pcs, 1 Pack = 12 Pcs)
    {"sku": "BB0049", "name_eresto": "Dough Keju", "item_code": "PJ260014", "after_qty": 7.0, "unit": "Pack", "cost": 33200.00, "conv": 12.0},
    # 24. BB0257 Gelas Cup 9 Oz (Isi 50 pcs) -> PEL260071
    {"sku": "BB0257", "name_eresto": "Gelas Cup 9 Oz (Isi 50 pcs)", "item_code": "PEL260071", "after_qty": 380.0, "unit": "Pcs", "cost": 1100.00, "conv": 1.0},
    # 25. BB0066 Goodie Bag -> PEL260072
    {"sku": "BB0066", "name_eresto": "Goodie Bag", "item_code": "PEL260072", "after_qty": 73.0, "unit": "Pcs", "cost": 1900.00, "conv": 1.0},
    # 26. BB0080 Kardus Ropi Kapasitas 20 roti (25 pcs/pack) -> PEL260074
    {"sku": "BB0080", "name_eresto": "Kardus Ropi Kapasitas 20 roti (25 pcs/pack)", "item_code": "PEL260074", "after_qty": 28.0, "unit": "Pcs", "cost": 4713.33, "conv": 1.0},
    # 27. BB0322 Kopi Hitam @6gr (Isi 20 Sachet) -> BB260061
    {"sku": "BB0322", "name_eresto": "Kopi Hitam @6gr (Isi 20 Sachet)", "item_code": "BB260061", "after_qty": 40.0, "unit": "Pcs", "cost": 921.52, "conv": 1.0},
    # 28. BB0102 Olesan Choco Crunchy (1 Kg) -> BB260043 (konversi kg -> g, 1 kg = 1000 g)
    {"sku": "BB0102", "name_eresto": "Olesan Choco Crunchy (1 Kg)", "item_code": "BB260043", "after_qty": 1.52, "unit": "kg", "cost": 60000.00, "conv": 1000.0},
    # 29. BB0103 Olesan Dark Choco (1 Kg) -> BB260044 (konversi kg -> g, 1 kg = 1000 g)
    {"sku": "BB0103", "name_eresto": "Olesan Dark Choco (1 Kg)", "item_code": "BB260044", "after_qty": 2.06, "unit": "kg", "cost": 75900.00, "conv": 1000.0},
    # 30. BB0107 Olesan Sarikaya (1 Kg) -> Qty 0 -> Skip
    {"sku": "BB0107", "name_eresto": "Olesan Sarikaya (1 Kg)", "item_code": "BB260049", "after_qty": 0.0, "unit": "kg", "cost": 32000.00, "conv": 1000.0, "skip": True, "reason": "Qty 0"},
    # 31. BB0110 Olesan Tiramisu (1 Kg) -> BB260047 (konversi kg -> g, 1 kg = 1000 g)
    {"sku": "BB0110", "name_eresto": "Olesan Tiramisu (1 Kg)", "item_code": "BB260047", "after_qty": 0.23, "unit": "kg", "cost": 67000.00, "conv": 1000.0},
    # 32. BB0117 Plastik Alas Gempi -> PEL260077
    {"sku": "BB0117", "name_eresto": "Plastik Alas Gempi", "item_code": "PEL260077", "after_qty": 38.0, "unit": "Pcs", "cost": 290.00, "conv": 1.0},
    # 33. BB0145 Susu Bubuk untuk Campuran Minuman (1 Kg) -> Qty 0 -> Skip
    {"sku": "BB0145", "name_eresto": "Susu Bubuk untuk Campuran Minuman (1 Kg)", "item_code": "BB260062", "after_qty": 0.0, "unit": "kg", "cost": 128000.00, "conv": 1000.0, "skip": True, "reason": "Qty 0"},
    # 34. BB0106 TIDAK AKTIF Olesan Pineapple (1 Kg) -> Skip
    {"sku": "BB0106", "name_eresto": "TIDAK AKTIF Olesan Pineapple (1 Kg)", "item_code": "BB260050", "after_qty": 0.0, "unit": "kg", "cost": 58000.00, "conv": 1000.0, "skip": True, "reason": "Tidak Aktif"},
    # 35. BB0108 TIDAK AKTIF Olesan Strawberry (1 Kg) -> Skip
    {"sku": "BB0108", "name_eresto": "TIDAK AKTIF Olesan Strawberry (1 Kg)", "item_code": "BB260051", "after_qty": 0.0, "unit": "kg", "cost": 59000.00, "conv": 1000.0, "skip": True, "reason": "Tidak Aktif"},
    # 36. RR0003 Ropi Keju -> PJO2602 (Ropi Keju Matang)
    {"sku": "RR0003", "name_eresto": "Ropi Keju", "item_code": "PJO2602", "after_qty": 9.0, "unit": "Pcs", "cost": 3850.00, "conv": 1.0},
]


def make_session():
    s = requests.Session()
    s.headers.update({
        "Authorization": f"token {API_KEY}",
        "Accept": "application/json"
    })
    return s


def main():
    print("=" * 75)
    print("REVISI STOCK RECONCILIATION OPENING STOCK - RS PKU MUHAMMADIYAH DELANGGU")
    print("=" * 75)
    s = make_session()

    # Filter item aktif
    active_items = [it for it in RAW_DATA if not it.get("skip")]
    skipped_items = [it for it in RAW_DATA if it.get("skip")]

    print(f"Total baris eResto   : {len(RAW_DATA)}")
    print(f"Item dilewati (skip) : {len(skipped_items)}")
    for sk in skipped_items:
        print(f"  - {sk['sku']} {sk['name_eresto']} ({sk['reason']})")
    print(f"Item aktif diinjeksi : {len(active_items)}\n")

    # Siapkan baris item untuk Stock Reconciliation
    reco_items = []
    total_inventory_value = 0.0

    for it in active_items:
        code = it["item_code"]
        r_it = s.get(f"{BASE_URL}/api/resource/Item/{code}").json().get("data", {})
        stock_uom = r_it.get("stock_uom")
        conv = it.get("conv", 1.0)
        native_qty = round(it["after_qty"] * conv, 4)
        native_rate = round(it["cost"] / conv, 4)
        item_val = round(it["after_qty"] * it["cost"], 2)
        total_inventory_value += item_val

        reco_items.append({
            "item_code": code,
            "warehouse": WAREHOUSE,
            "qty": native_qty,
            "valuation_rate": native_rate,
            "uom": stock_uom,
            "custom_uom": it["unit"],
            "custom_conversion_factor": round(conv, 4),
            "custom_qty_after": round(it["after_qty"], 4),
            "custom_qty_before": 0.0,
            "custom_qty_difference_per_uom": round(it["after_qty"], 4),
            "custom_valuation_rate_per_uom": round(it["cost"], 4),
            "allow_zero_valuation_rate": 1
        })

    print(f"Total Nilai Persediaan Fisik: Rp {total_inventory_value:,.2f}")

    # Buat Dokumen Stock Reconciliation
    payload = {
        "doctype": "Stock Reconciliation",
        "company": COMPANY,
        "purpose": "Opening Stock",
        "set_posting_time": 1,
        "posting_date": POSTING_DATE,
        "posting_time": POSTING_TIME,
        "set_warehouse": WAREHOUSE,
        "expense_account": EXPENSE_ACCOUNT,
        "cost_center": COST_CENTER,
        "items": reco_items,
        "docstatus": 1  # Langsung submit
    }

    print("\nMengirim dokumen Stock Reconciliation ke Frappe Cloud...")
    r_create = s.post(f"{BASE_URL}/api/resource/Stock Reconciliation", json=payload, timeout=60)
    if r_create.status_code not in (200, 201):
        print(f"FATAL Gagal submit Stock Reconciliation: HTTP {r_create.status_code} {r_create.text}")
        sys.exit(1)

    reco_doc = r_create.json().get("data", {})
    reco_name = reco_doc.get("name")
    print(f"[OK] Stock Reconciliation Berhasil Dibuat & Tersubmit: {reco_name} (Status: Submitted)")

    # Simpan Laporan Excel
    os.makedirs("pos_sync", exist_ok=True)
    report_rows = []
    for it in active_items:
        conv = it.get("conv", 1.0)
        report_rows.append({
            "SKU eResto": it["sku"],
            "Nama eResto": it["name_eresto"],
            "Item Code ERPNext": it["item_code"],
            "Qty (UOM Asli)": it["after_qty"],
            "UOM Asli": it["unit"],
            "Cost Per Unit (Rp)": it["cost"],
            "Konversi": conv,
            "Qty ERPNext": round(it["after_qty"] * conv, 4),
            "Valuation Rate (Rp)": round(it["cost"] / conv, 4),
            "Total Nilai (Rp)": round(it["after_qty"] * it["cost"], 2),
            "Status": "SUBMITTED (LOLOS)"
        })

    for sk in skipped_items:
        report_rows.append({
            "SKU eResto": sk["sku"],
            "Nama eResto": sk["name_eresto"],
            "Item Code ERPNext": sk["item_code"],
            "Qty (UOM Asli)": sk["after_qty"],
            "UOM Asli": sk["unit"],
            "Cost Per Unit (Rp)": sk["cost"],
            "Konversi": sk.get("conv", 1.0),
            "Qty ERPNext": 0,
            "Valuation Rate (Rp)": 0,
            "Total Nilai (Rp)": 0,
            "Status": f"DILEWATI ({sk['reason']})"
        })

    df = pd.DataFrame(report_rows)
    with pd.ExcelWriter(EXCEL_OUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Revisi_Reco_Delanggu", index=False)

    wb = openpyxl.load_workbook(EXCEL_OUT)
    ws = wb["Revisi_Reco_Delanggu"]
    ws.freeze_panes = "A2"

    hdr_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    skip_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    hdr_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            if "SUBMITTED" in str(cell.value):
                cell.fill = pass_fill
                cell.font = Font(name="Calibri", size=10, bold=True, color="375623")
            elif "DILEWATI" in str(cell.value):
                cell.fill = skip_fill
                cell.font = Font(name="Calibri", size=10, color="7F6000")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(EXCEL_OUT)
    print(f"Laporan Excel berhasil disimpan: {EXCEL_OUT}")
    print("=" * 75)


if __name__ == "__main__":
    main()
