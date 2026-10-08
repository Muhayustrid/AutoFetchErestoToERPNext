"""Injeksi POS Invoice LOKAL: eResto (cache pos-fetch) -> erp.localhost:8081.

Alur resmi POSNext per (outlet=tanggal):
  1. POS Opening Shift  (REST + submit; period_start_date/posting_date historis)
  2. POS Invoice per order_no lewat lane BACKDATE HO
     (pos_next.api.backdate_invoices.submit_backdate_invoice: posting_date +
     posting_time asli, 1 panggilan = draft + submit)
  3. POS Closing Shift  (rekap server get_closing_shift_data -> REST create +
     submit; posting_date = tanggal shift, closing_amount = expected -> 0)

Batasan POSNext yang menentukan bentuk runner: satu user hanya boleh punya
SATU shift terbuka. Karena itu grup (outlet, tanggal) diproses BERURUTAN:
buka shift -> inject semua invoice grup (paralel di dalam grup) -> tutup.
Shift yang sudah tertutup (mis. pilot TJBK 4 Okt) dan masih punya order
tersisa dibuka ulang lewat jalur resmi `reopen_shift`, lalu ditutup lagi.

Idempotency: marker `remarks = "ERESTO <order_no>"` di POS Invoice;
rerun tidak menduplikasi. Shift dideteksi dari (company, posting_date, user).

Skip + lapor (tanpa menebak): SKU unmapped, total 0 / tanpa item bayar,
total != jumlah baris, jumlah bayar != total, outlet tanpa company/profil.

Flag sementara (dipulihkan di akhir kecuali --keep-flags):
  - POS Next Global Settings.allow_change_posting_date = 1 (gerbang backdate)
  - POS Next Global Settings.allow_negative_stock = 1 (stok outlet belum diisi)

Valuasi: item komponen keluar dari gudang outlet kosong -> ERPNext minta
valuation rate. Isi dulu `set_item_valuation.py` (SLE site lama/lokal).

Usage:
  python3 inject_pos_invoice.py --date-from 2026-10-01 --date-to 2026-10-05
      [--limit-groups N] [--outlet FILTER] [--dry-run] [--workers 8]
      [--only-shifts] [--no-close] [--keep-flags] [--refresh-list]

Env: /tmp/eresto_mig.env (ERPNEXT_URL, ERPNEXT_API_KEY site LOKAL;
     ERESTO_TOKEN/EMAIL/USER_ID hanya bila --refresh-list).
Output: pos_sync/PosInvoice_Lokal_<from>_sd_<to>.xlsx + state JSON.
"""
import argparse
import collections
import csv
import json
import os
import sys
import threading
import time
import urllib.parse
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prod_common import Client, load_env  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
POSFETCH = os.path.join(ROOT, "experiments", "pos-fetch")
DET_DIR = os.path.join(POSFETCH, ".cache_details")
MAP_CSV = os.path.join(ROOT, "experiments", "sku_to_item_code.csv")
STATE_PATH = os.path.join(HERE, "pos_invoice_state.json")
DEFAULT_ENV = "/tmp/eresto_mig.env"

MARK = "ERESTO "          # remarks marker
PAY_FIELDS = [("cash_amount", "Cash"), ("ewallet_amount", "QRIS"),
              ("debit_amount", "Debit Card"), ("credit_amount", "Credit Card"),
              ("transfer_amount", "Wire Transfer")]
MODES = ["Cash", "QRIS", "Debit Card", "Credit Card"]
TOL = 0.9

_print_lock = threading.Lock()


def log(msg):
    with _print_lock:
        print(msg, flush=True)


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def time_of(iso):
    """'2026-10-01T05:03:27.000+07:00' -> '05:03:27' (None bila gagal)."""
    try:
        t = str(iso).split("T", 1)[1]
        return t[:8]
    except (IndexError, TypeError):
        return None


def fetch_discount_code(erp, today=None):
    """Kode konfirmasi diskon HO yang aktif hari ini (dipakai baris rate<plr)."""
    today = today or time.strftime("%Y-%m-%d")
    st, d = erp.get("/api/resource/POS Discount Confirmation Code", {
        "filters": json.dumps([["status", "=", "Active"]]),
        "fields": json.dumps(["name", "code", "valid_from", "valid_upto"]),
        "limit_page_length": 50})
    for r in d.get("data", []) if st == 200 else []:
        vf = str(r.get("valid_from") or "2000-01-01")[:10]
        vu = str(r.get("valid_upto") or "2500-12-31")[:10]
        if vf <= today <= vu:
            return r.get("code") or r.get("name")
    return ""


def style_sheet(ws):
    ws.freeze_panes = "A2"
    try:
        ws.auto_filter.ref = ws.dimensions
    except Exception:
        pass
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill, c.font = fill, Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 45)


# ---------------------------------------------------------------- eResto ----

class Eresto:
    def __init__(self, env):
        self.token = env.get("ERESTO_TOKEN", "")
        self.email = env.get("ERESTO_EMAIL", "")
        self.user_id = env.get("ERESTO_USER_ID", "")
        self.s = requests.Session()
        self.s.headers.update({
            "accept": "application/json, text/plain, */*",
            "authorization": f"Token token={self.token}, email={self.email}",
            "referer": "https://ropi.eresto.id/manager/v2/index.html",
            "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        })

    def get(self, path, params):
        last = None
        for a in range(1, 4):
            try:
                r = self.s.get("https://ropi.eresto.id" + path, params=params,
                               timeout=90)
                if r.status_code in (429, 500, 502, 503, 504):
                    last = f"HTTP {r.status_code}"
                    time.sleep(2 ** a)
                    continue
                r.raise_for_status()
                return r.json(), ""
            except (requests.RequestException, ValueError) as e:
                last = str(e)[:150]
                time.sleep(2 ** a)
        return None, last

    def list_orders(self, dfrom, dto):
        orders, page = [], 1
        while True:
            j, err = self.get("/v3/orders", {
                "current_user_id": self.user_id,
                "date[geq]": f"{dfrom} 00:00:00.000",
                "date[leq]": f"{dto} 23:59:59.999",
                "domain": "ropi", "outlet_id": "",
                "page": page, "per_page": 200,
                "sort[property]": "date", "sort[order]": "asc",
                "token": self.token, "_": int(time.time() * 1000)})
            if err:
                raise SystemExit(f"eResto list page {page} gagal: {err}")
            batch = j.get("data", [])
            orders.extend(batch)
            if len(batch) < 200 or not (j.get("meta") or {}).get("next_page"):
                break
            page += 1
        seen, uniq = set(), []
        for o in orders:
            if o.get("id") and o["id"] not in seen:
                seen.add(o["id"])
                uniq.append(o)
        return uniq

    def detail(self, oid):
        p = os.path.join(DET_DIR, f"{oid}.json")
        if os.path.exists(p):
            try:
                return json.load(open(p)), ""
            except ValueError:
                os.unlink(p)
        j, err = self.get(f"/v3/orders/{oid}", {
            "domain": "ropi", "token": self.token,
            "current_user_id": self.user_id, "_": int(time.time() * 1000)})
        if err:
            return None, err
        root = j.get("data") if isinstance(j.get("data"), dict) else j
        attrs = (root or {}).get("attributes")
        if not isinstance(attrs, dict):
            return None, "payload tanpa data.attributes"
        json.dump(attrs, open(p, "w"))
        return attrs, ""


# ------------------------------------------------------------------ plan ----

class Planner:
    """Rakit daftar order -> payload POS Invoice; tandai yang di-skip."""

    def __init__(self, erp):
        self.erp = erp
        self.mapping = {}
        with open(MAP_CSV) as f:
            for r in csv.DictReader(f):
                self.mapping[r["sku_eresto"]] = r["item_code_erpnext"]
        self.companies = {}   # lower -> canonical name
        for c in erp.list_all("Company", ["name"]):
            self.companies[c["name"].strip().lower()] = c["name"]
        self.profiles = {}
        for p in erp.list_all("POS Profile",
                              ["name", "company", "warehouse", "customer",
                               "selling_price_list", "disabled"],
                              filters=[["name", "like", "SHIFT 1 - %"]]):
            self.profiles[p["company"]] = p
        fields = ["name", "item_name", "stock_uom", "is_stock_item"]
        self.items = {i["name"]: i for i in erp.list_all("Item", fields)}
        # harga jual per price list: dipakai menandai baris rate<price_list_rate
        # sebagai diskon eksplisit + memilih kode konfirmasi (meniru resolusi
        # server `_fetch_uom_prices_map`: valid pada tanggal posting, uom menang)
        self.prices = {}
        rows = erp.list_all(
            "Item Price",
            ["item_code", "price_list", "price_list_rate", "valid_from",
             "valid_upto", "uom"],
            filters=[["selling", "=", 1]])
        for r in rows:
            self.prices.setdefault(r["price_list"], {}).setdefault(
                r["item_code"], []).append(r)

    def server_price(self, price_list, code, date, uom):
        """Harga server untuk (code, price list, tanggal) — valid_from/upto
        dihormati, per-uom menang, sama seperti `_fetch_uom_prices_map`."""
        best = None
        for r in (self.prices.get(price_list, {}) or {}).get(code, []):
            vf = str(r.get("valid_from") or "2000-01-01")[:10]
            vu = str(r.get("valid_upto") or "2500-12-31")[:10]
            if not (vf <= date <= vu):
                continue
            if (r.get("uom") or "") != uom:
                continue
            if best is None or vf > best[0]:
                best = (vf, float(r["price_list_rate"]))
        return best[1] if best else None

    def company_of(self, outlet):
        return self.companies.get((outlet or "").strip().lower())

    def resolve(self, attrs, detail, discount_code=""):
        """-> (payload | None, reason)."""
        outlet = attrs.get("outlet_name") or ""
        date = str(attrs.get("date"))[:10]
        company = self.company_of(outlet)
        if not company:
            return None, "OUTLET_TANPA_COMPANY"
        prof = self.profiles.get(company)
        if not prof:
            return None, "TANPA_PROFIL_SHIFT1"
        if prof.get("disabled"):
            return None, "PROFIL_DISABLED"
        lines = [it for it in detail.get("order_items") or []
                 if not it.get("void") and num(it.get("paid_quantity")) > 0]
        if not lines:
            return None, "TANPA_ITEM_BAYAR"
        p = detail.get("payment") or {}
        total = num(p.get("total"))
        if total <= 0:
            return None, "TOTAL_NOL"
        unmapped = sorted({(it.get("product_sku") or "") for it in lines
                           if (it.get("product_sku") or "") not in self.mapping})
        if unmapped:
            return None, "SKU_UNMAPPED: " + ",".join(unmapped)
        line_sum = sum(num(it.get("paid_amount"))
                       if it.get("paid_amount") not in (None, "")
                       else num(it.get("paid_quantity")) * num(it.get("price"))
                       for it in lines)
        if abs(line_sum - total) >= 1:
            return None, f"TOTAL_BEDA (header {total:g} vs baris {line_sum:g})"
        payments = []
        for field, mode in PAY_FIELDS:
            amt = num(p.get(field))
            if amt > 0:
                payments.append({"mode_of_payment": mode, "amount": amt})
        pay_sum = sum(x["amount"] for x in payments)
        if not payments or abs(pay_sum - total) >= 1:
            return None, f"BAYAR_BEDA (total {total:g} vs bayar {pay_sum:g})"
        agg = collections.OrderedDict()
        for it in lines:
            code = self.mapping[it["product_sku"]]
            rate = num(it.get("price"))
            key = (code, rate)
            agg[key] = agg.get(key, 0.0) + num(it.get("paid_quantity"))
        items = []
        needs_code = False
        for (code, rate), qty in agg.items():
            master = self.items.get(code, {})
            uom = master.get("stock_uom") or "Pcs"
            row = {
                "item_code": code,
                "item_name": master.get("item_name") or code,
                "qty": qty, "rate": rate, "price_list_rate": rate,
                "uom": uom,
                "warehouse": prof["warehouse"], "conversion_factor": 1,
                "discount_percentage": 0, "discount_amount": 0,
            }
            # harga eResto di bawah list ERPNext = diskon nyata. Kirim plr asli
            # + field diskon eksplisit supaya baris tidak jatuh ke lane
            # "manual rate edit" (yang butuh allow_user_to_edit_rate) dan
            # sebaliknya diotorisasi kode konfirmasi diskon.
            plr = self.server_price(prof.get("selling_price_list"), code,
                                    date, uom)
            if plr and rate < plr - 0.005:
                row["price_list_rate"] = plr
                row["discount_amount"] = round(plr - rate, 2)
                row["discount_percentage"] = round((plr - rate) / plr * 100, 2)
                needs_code = True
            items.append(row)
        payload = {
            "customer": prof.get("customer") or "Walk-in Customer",
            "items": items,
            "payments": payments,
            "discount_amount": 0,
            "apply_discount_on": "Grand Total",
            "remarks": MARK + attrs.get("order_no"),
        }
        if needs_code:
            if not discount_code:
                return None, "BUTUH_KODE_DISKON"
            payload["discount_confirmation_code"] = discount_code
        return payload, ""


# --------------------------------------------------------------- injektor ---

class Injector:
    def __init__(self, erp, args):
        self.erp = erp
        self.args = args
        self.state = {"orders": {}, "groups": {}}
        if os.path.exists(STATE_PATH):
            try:
                self.state = json.load(open(STATE_PATH))
            except ValueError:
                pass
        self.state.setdefault("orders", {})
        self.state.setdefault("groups", {})

    def save_state(self):
        tmp = STATE_PATH + ".tmp"
        json.dump(self.state, open(tmp, "w"), indent=1, default=str)
        os.replace(tmp, STATE_PATH)

    # --- flags ---
    def set_single(self, field, value):
        return self.erp.req("POST", "/api/method/frappe.client.set_value", {
            "doctype": "POS Next Global Settings",
            "name": "POS Next Global Settings",
            "fieldname": field, "value": value})

    def read_single(self, field):
        st, d = self.erp.get("/api/resource/POS Next Global Settings/"
                             "POS Next Global Settings",
                             {"fields": json.dumps([field])})
        return d.get("data", {}).get(field)

    # --- master prefetch ---
    def prefetch(self):
        rows = self.erp.list_all(
            "POS Invoice", ["name", "remarks", "posting_date", "grand_total",
                            "docstatus", "status"],
            filters=[["remarks", "like", MARK.strip() + " %"],
                     ["docstatus", "!=", 2]])
        self.existing = {r["remarks"]: r for r in rows if r.get("remarks")}
        log(f"  invoice marker ada: {len(self.existing)}")
        shifts = self.erp.list_all(
            "POS Opening Shift",
            ["name", "company", "posting_date", "user", "status", "docstatus",
             "pos_closing_shift", "period_start_date"],
            filters=[["user", "=", self.args.user]])
        self.shifts = {}
        for s in shifts:
            key = (s["company"], str(s["posting_date"])[:10])
            self.shifts[key] = s
        log(f"  shift (user {self.args.user}): {len(self.shifts)}")

    # --- shift ---
    def ensure_shift(self, company, date, reopen=True):
        """-> (shift, action, err). action: baru/pakai/reopen."""
        key = (company, date)
        s = self.shifts.get(key)
        if s:
            status = s.get("status")
            if status == "Open" and s.get("docstatus") == 1:
                return s, "pakai", ""
            if status == "Closed" and s.get("pos_closing_shift"):
                if not reopen:
                    return s, "tertutup", ""
                # jalur resmi HO: cancel closing -> opening kembali Open
                st, d = self.erp.req(
                    "POST",
                    "/api/method/pos_next.api.backdate_invoices.reopen_shift",
                    {"pos_opening_shift": s["name"]})
                if st != 200:
                    return None, "", f"reopen HTTP {st}: {json.dumps(d)[:300]}"
                s["status"] = "Open"
                s["pos_closing_shift"] = None
                return s, "reopen", ""
            return None, "", f"shift {s['name']} status {status}/" \
                             f"{s.get('docstatus')} tidak bisa dipakai"
        prof = self.profiles[company]
        payload = {
            "doctype": "POS Opening Shift",
            "pos_profile": prof["name"],
            "company": company,
            "user": self.args.user,
            "period_start_date": f"{date} 00:00:00",
            "posting_date": date,
            "balance_details": [{"mode_of_payment": m, "amount": 0}
                                for m in MODES],
        }
        st, d = self.erp.post("POS Opening Shift", payload)
        if st != 200:
            return None, "", f"buat shift HTTP {st}: {json.dumps(d)[:300]}"
        name = d["data"]["name"]
        st, d = self.erp.submit("POS Opening Shift", name)
        if st != 200:
            return None, "", f"submit shift HTTP {st}: {json.dumps(d)[:300]}"
        s = {"name": name, "company": company, "posting_date": date,
             "status": "Open", "docstatus": 1, "pos_closing_shift": None}
        self.shifts[key] = s
        return s, "baru", ""

    # --- invoice ---
    def submit_invoice(self, order_no, shift_name, date, posting_time, payload):
        inv = dict(payload)
        inv["posa_pos_opening_shift"] = shift_name
        inv["posting_date"] = date
        if posting_time:
            inv["posting_time"] = posting_time
        st, d = self.erp.req(
            "POST",
            "/api/method/pos_next.api.backdate_invoices."
            "submit_backdate_invoice",
            {"invoice": inv})
        if st != 200:
            return None, f"HTTP {st}: {json.dumps(d)[:400]}"
        msg = d.get("message") or {}
        if not msg.get("name"):
            return None, f"tanpa name: {json.dumps(msg)[:300]}"
        return msg, ""

    # --- closing ---
    def close_shift(self, company, date):
        key = (company, date)
        s = self.shifts[key]
        if s.get("pos_closing_shift"):
            return s["pos_closing_shift"], "sudah"
        st, d = self.erp.req(
            "POST", "/api/method/pos_next.api.shifts.get_closing_shift_data",
            {"opening_shift": s["name"]})
        if st != 200:
            return None, f"preview HTTP {st}: {json.dumps(d)[:300]}"
        prev = d.get("message") or {}
        recon = []
        for row in prev.get("payment_reconciliation") or []:
            exp = num(row.get("expected_amount"))
            recon.append({"mode_of_payment": row.get("mode_of_payment"),
                          "opening_amount": num(row.get("opening_amount")),
                          "expected_amount": exp,
                          "closing_amount": exp})
        txns = [{k: row.get(k) for k in
                 ("pos_invoice", "sales_invoice", "posting_date", "customer",
                  "grand_total", "transaction_currency", "transaction_amount")}
                for row in prev.get("pos_transactions") or []]
        taxes = [{k: row.get(k) for k in ("account_head", "rate", "amount")}
                 for row in prev.get("taxes") or []]
        payload = {
            "doctype": "POS Closing Shift",
            "pos_opening_shift": s["name"],
            "pos_profile": prev.get("pos_profile"),
            "user": prev.get("user") or self.args.user,
            "company": prev.get("company") or company,
            "period_start_date": prev.get("period_start_date")
            or f"{date} 00:00:00",
            "period_end_date": f"{date} 23:59:59",
            "posting_date": date,
            "payment_reconciliation": recon,
            "pos_transactions": txns,
            "taxes": taxes,
            "grand_total": num(prev.get("grand_total")),
            "net_total": num(prev.get("net_total")),
            "total_quantity": num(prev.get("total_quantity")),
        }
        st, d = self.erp.post("POS Closing Shift", payload)
        if st != 200:
            return None, f"buat closing HTTP {st}: {json.dumps(d)[:300]}"
        name = d["data"]["name"]
        st, d = self.erp.submit("POS Closing Shift", name)
        if st != 200:
            return None, f"submit closing HTTP {st}: {json.dumps(d)[:300]}"
        s["pos_closing_shift"] = name
        s["status"] = "Closed"
        return name, ""


# ------------------------------------------------------------------ main ----

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date-from", required=True)
    ap.add_argument("--date-to", required=True)
    ap.add_argument("--outlet", action="append", default=[],
                    help="filter substring outlet_name (boleh berkali-kali)")
    ap.add_argument("--limit-groups", type=int, default=0)
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--only-shifts", action="store_true")
    ap.add_argument("--no-close", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--keep-flags", action="store_true")
    ap.add_argument("--refresh-list", action="store_true")
    ap.add_argument("--env", default=DEFAULT_ENV)
    args = ap.parse_args()

    env = load_env(args.env)
    erp = Client(env["ERPNEXT_URL"], env["ERPNEXT_API_KEY"],
                 label="local", pause=0.0)
    args.user = erp.get("/api/method/frappe.auth.get_logged_user")[1].get(
        "message", "")
    log(f"== user: {args.user} | site: {env['ERPNEXT_URL']} ==")

    es = Eresto(env)
    tag = f"{args.date_from}_sd_{args.date_to}"
    list_path = os.path.join(POSFETCH, f".cache_details_list_{tag}.json")
    if args.refresh_list or not os.path.exists(list_path):
        log("== tarik list eResto ==")
        orders = es.list_orders(args.date_from, args.date_to)
        json.dump(orders, open(list_path, "w"))
    else:
        orders = json.load(open(list_path))
    orders = [o for o in orders
              if args.date_from <= str(o["attributes"].get("date"))[:10]
              <= args.date_to]
    if args.outlet:
        orders = [o for o in orders if any(
            f.lower() in (o["attributes"].get("outlet_name") or "").lower()
            for f in args.outlet)]
    log(f"== order eResto {tag}: {len(orders)} ==")

    log("== master lokal ==")
    planner = Planner(erp)
    discount_code = fetch_discount_code(erp)
    log(f"  company {len(planner.companies)} | profil SHIFT 1 "
        f"{len(planner.profiles)} | item {len(planner.items)} | "
        f"kode diskon aktif: {discount_code or '(TIDAK ADA)'}")

    # detail (cache + fetch yang bolong)
    missing = [o["id"] for o in orders
               if not os.path.exists(os.path.join(DET_DIR, o["id"] + ".json"))]
    details = {}
    if missing:
        log(f"  detail bolong: {len(missing)} -> tarik")
        with ThreadPoolExecutor(max_workers=10) as pool:
            futs = {pool.submit(es.detail, oid): oid for oid in missing}
            done = 0
            for f in as_completed(futs):
                attrs, err = f.result()
                done += 1
                if err:
                    log(f"  !! detail {futs[f]}: {err}")
                else:
                    details[futs[f]] = attrs
                if done % 500 == 0 or done == len(missing):
                    log(f"  detail {done}/{len(missing)}")
    for o in orders:
        if o["id"] not in details:
            p = os.path.join(DET_DIR, o["id"] + ".json")
            if os.path.exists(p):
                details[o["id"]] = json.load(open(p))

    # rencana
    plan, skipped = [], []
    for o in orders:
        a = o["attributes"]
        det = details.get(o["id"])
        if det is None:
            skipped.append({"order_no": a.get("order_no"),
                            "outlet": a.get("outlet_name"),
                            "tanggal": str(a.get("date"))[:10],
                            "total": num((a.get("payment") or {}).get("total")),
                            "alasan": "DETAIL_GAGAL"})
            continue
        payload, reason = planner.resolve(a, det, discount_code)
        if reason:
            skipped.append({"order_no": a.get("order_no"),
                            "outlet": a.get("outlet_name"),
                            "tanggal": str(a.get("date"))[:10],
                            "total": num((det.get("payment") or {}).get("total")),
                            "alasan": reason})
            continue
        plan.append({"order_no": a.get("order_no"),
                     "outlet": a.get("outlet_name"),
                     "company": planner.company_of(a.get("outlet_name")),
                     "tanggal": str(a.get("date"))[:10],
                     "posting_time": time_of(a.get("date")),
                     "total": num((det.get("payment") or {}).get("total")),
                     "payload": payload})

    groups = collections.defaultdict(list)
    for row in plan:
        groups[(row["company"], row["tanggal"])].append(row)
    glist = sorted(groups.items(), key=lambda kv: (kv[0][1], kv[0][0]))
    if args.limit_groups:
        glist = sorted(glist, key=lambda kv: len(kv[1]))[:args.limit_groups]
        keep = {k for k, _ in glist}
        plan = [r for r in plan if (r["company"], r["tanggal"]) in keep]
        groups = {k: v for k, v in glist}
    log(f"== rencana: {len(plan)} invoice, {len(groups)} outlet-hari, "
        f"skip {len(skipped)} ==")
    for reason, cnt in collections.Counter(
            s["alasan"].split(":")[0] for s in skipped).most_common():
        log(f"   skip {reason}: {cnt}")

    if args.dry_run:
        for (c, d), rows in glist[:5]:
            log(f"   {d} {c}: {len(rows)} invoice, contoh "
                f"{rows[0]['order_no']} Rp{rows[0]['total']:,.0f}")
        log("DRY-RUN selesai")
        return

    inj = Injector(erp, args)
    inj.profiles = planner.profiles
    inj.prefetch()

    # flag sementara
    orig = {}
    if not args.only_shifts:
        for field in ("allow_change_posting_date", "allow_negative_stock"):
            orig[field] = inj.read_single(field)
            st, d = inj.set_single(field, 1)
            log(f"  set {field}=1 (asli {orig[field]}) -> HTTP {st}")
        st, d = erp.get("/api/resource/Stock Settings/Stock Settings",
                        {"fields": '["allow_negative_stock"]'})
        log(f"  Stock Settings.allow_negative_stock = "
            f"{d.get('data', {}).get('allow_negative_stock')}")

    t0 = time.time()
    results = []
    counters = collections.Counter()
    lock = threading.Lock()

    def handle_group(company, date, rows, ex):
        """Satu outlet-hari. Dipanggil berurutan (satu shift terbuka/user)."""
        missing_rows = [r for r in rows
                        if MARK + r["order_no"] not in inj.existing]
        out = []
        for r in rows:
            old = inj.existing.get(MARK + r["order_no"])
            if old:
                out.append((r["order_no"], "ADA", old["name"]))
        if not missing_rows:
            # semua sudah ada; tutup shift bila masih Open dari run gagal
            s = inj.shifts.get((company, date))
            if s and s.get("status") == "Open" and s.get("docstatus") == 1 \
                    and not args.no_close:
                cname, cerr = inj.close_shift(company, date)
                if cerr and cerr != "sudah":
                    out.append((f"CLOSING {date} {company}", "CLOSING_GAGAL",
                                cerr))
            return out
        s, action, err = inj.ensure_shift(company, date, reopen=True)
        if err:
            return out + [(r["order_no"], "SHIFT_GAGAL", err)
                          for r in missing_rows]
        if action in ("baru", "reopen"):
            log(f"  shift {action}: {s['name']} ({len(missing_rows)} invoice "
                f"menyusul)")
        futs = {ex.submit(inj.submit_invoice, r["order_no"], s["name"],
                          r["tanggal"], r["posting_time"], r["payload"]): r
                for r in missing_rows}
        done_ok = 0
        for f in as_completed(futs):
            r = futs[f]
            msg, err = f.result()
            if err:
                out.append((r["order_no"], "GAGAL", err))
                continue
            gt = num(msg.get("grand_total"))
            if abs(gt - r["total"]) >= 1:
                out.append((r["order_no"], "BEDA_TOTAL",
                            f"{msg.get('name')} gt={gt:g} vs {r['total']:g}"))
            else:
                out.append((r["order_no"], "OK", msg.get("name")))
                done_ok += 1
                with lock:
                    inj.existing[MARK + r["order_no"]] = {
                        "name": msg.get("name"), "remarks": MARK + r["order_no"]}
        # retry serial: race tabSeries (naming) sering 500 saat paralel
        for rnd in (1, 2):
            retry = [o for o in out if o[1] == "GAGAL"]
            if not retry:
                break
            log(f"  retry serial r{rnd}: {len(retry)} invoice "
                f"({company} {date})")
            fixed = []
            for ono, _st, _info in retry:
                r = next(x for x in missing_rows if x["order_no"] == ono)
                msg, err = inj.submit_invoice(r["order_no"], s["name"],
                                              r["tanggal"], r["posting_time"],
                                              r["payload"])
                if err:
                    fixed.append((ono, "GAGAL", err))
                    continue
                gt = num(msg.get("grand_total"))
                if abs(gt - r["total"]) >= 1:
                    fixed.append((ono, "BEDA_TOTAL",
                                  f"{msg.get('name')} gt={gt:g} vs "
                                  f"{r['total']:g}"))
                else:
                    fixed.append((ono, "OK", msg.get("name")))
                    done_ok += 1
                    with lock:
                        inj.existing[MARK + r["order_no"]] = {
                            "name": msg.get("name"),
                            "remarks": MARK + r["order_no"]}
            for f in fixed:
                out = [o for o in out
                       if not (o[0] == f[0] and o[1] == "GAGAL")] + [f]
        ok_all = all(o[1] in ("OK", "ADA") for o in out)
        if not args.no_close and ok_all:
            cname, cerr = inj.close_shift(company, date)
            if cerr and cerr != "sudah":
                out.append((f"CLOSING {date} {company}", "CLOSING_GAGAL",
                            cerr))
            else:
                log(f"  closing {cname}: {company} {date} "
                    f"({done_ok} baru)")
        elif not ok_all:
            log(f"  LEWAT closing {company} {date}: ada kegagalan")
        return out

    if args.only_shifts:
        # info saja: shift existing per grup + aksi yang akan diambil.
        # (tidak membuka shift — satu user hanya boleh satu shift terbuka)
        for (c, d), rows in glist[:40]:
            s = inj.shifts.get((c, d))
            log(f"  {d} {c}: {len(rows)} invoice | shift "
                f"{s['name'] + ' ' + s['status'] if s else '(belum ada)'}")
        log(f"  ... total {len(glist)} grup")
        return

    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        for idx, ((c, d), rows) in enumerate(glist, 1):
            res = handle_group(c, d, rows, ex)
            with lock:
                for ono, status, info in res:
                    counters[status] += 1
                    results.append({"order_no": ono, "status": status,
                                    "info": info})
                inj.save_state()
            if idx % 10 == 0 or idx == len(glist):
                el = time.time() - t0
                ok = counters["OK"] + counters["ADA"]
                log(f"  grup {idx}/{len(glist)} | invoice OK {ok} | "
                    f"gagal {counters['GAGAL'] + counters['BEDA_TOTAL']} | "
                    f"closing gagal {counters['CLOSING_GAGAL']} | "
                    f"{el / 60:.1f}mnt")

    if not args.keep_flags:
        for field, val in orig.items():
            st, d = inj.set_single(field, val)
            log(f"  restore {field}={val} -> HTTP {st}")

    log("== hasil ==")
    for k, v in counters.most_common():
        log(f"  {k}: {v}")

    # ---- verifikasi ----
    log("== verifikasi ==")
    plan_map = {r["order_no"]: r for r in plan}
    ok_names = [r["info"] for r in results if r["status"] in ("OK", "ADA")]
    inv_rows = []
    for i in range(0, len(ok_names), 100):
        chunk = ok_names[i:i + 100]
        inv_rows += erp.list_all(
            "POS Invoice", ["name", "posting_date", "grand_total", "status",
                            "remarks", "pos_profile"],
            filters=[["name", "in", chunk]])
    inv_by_name = {r["name"]: r for r in inv_rows}
    checks = []
    bad_date = [r for r in inv_rows
                if r["remarks"] and r["remarks"].startswith(MARK)
                and r["remarks"][len(MARK):] in plan_map
                and str(r["posting_date"])[:10] !=
                plan_map[r["remarks"][len(MARK):]]["tanggal"]]
    checks.append((f"posting_date invoice = tanggal eResto "
                   f"({len(inv_rows)} dicek, {len(bad_date)} beda)",
                   not bad_date))
    gl = sum(num(r["grand_total"]) for r in inv_rows)
    src = sum(plan_map[r["remarks"][len(MARK):]]["total"]
              for r in inv_rows
              if r["remarks"] and r["remarks"].startswith(MARK)
              and r["remarks"][len(MARK):] in plan_map)
    checks.append((f"Rp invoice ({gl:,.0f}) = Rp eResto ({src:,.0f})",
                   abs(gl - src) < 1))
    closings = [s["pos_closing_shift"] for s in inj.shifts.values()
                if s.get("pos_closing_shift")]
    cl_rows = []
    for i in range(0, len(closings), 100):
        cl_rows += erp.list_all(
            "POS Closing Shift", ["name", "posting_date", "grand_total",
                                  "period_end_date", "pos_opening_shift"],
            filters=[["name", "in", closings[i:i + 100]]])
    diff_bad = []
    for i in range(0, len(cl_rows), 50):
        st, d = erp.get("/api/resource/POS Closing Shift",
                        {"filters": json.dumps(
                            [["name", "in", [x["name"] for x in
                                             cl_rows[i:i + 50]]]]),
                         "fields": '["name","payment_reconciliation"]',
                         "limit_page_length": 50})
        for row in d.get("data", []):
            for rec in row.get("payment_reconciliation") or []:
                if abs(num(rec.get("difference"))) > 0.001:
                    diff_bad.append((row["name"], rec.get("mode_of_payment"),
                                     rec.get("difference")))
    checks.append((f"closing selisih 0 ({len(cl_rows)} closing, "
                   f"{len(diff_bad)} beda)", not diff_bad))
    gl_bal_bad = []
    sample = [r["name"] for r in inv_rows[:30]]
    for name in sample:
        st, d = erp.get("/api/resource/GL Entry", {
            "filters": json.dumps([["voucher_no", "=", name]]),
            "fields": '["debit","credit"]', "limit_page_length": 0})
        rows = d.get("data", [])
        bal = sum(num(x["debit"]) - num(x["credit"]) for x in rows)
        if abs(bal) > 0.01:
            gl_bal_bad.append((name, bal))
    checks.append((f"GL balance 0 (sample {len(sample)}, "
                   f"{len(gl_bal_bad)} beda)", not gl_bal_bad))
    for n, ok in checks:
        log(f"  [{'LOLOS' if ok else 'GAGAL'}] {n}")

    # rekap
    df_plan = pd.DataFrame([{k: r[k] for k in
                             ("order_no", "outlet", "company", "tanggal",
                              "total")} for r in plan])
    if len(df_plan):
        by_day = df_plan.groupby("tanggal").agg(
            invoice=("order_no", "count"), outlet=("company", "nunique"),
            omzet=("total", "sum")).reset_index()
        by_outlet = df_plan.groupby(["tanggal", "company"], as_index=False).agg(
            invoice=("order_no", "count"), omzet=("total", "sum"))
    else:
        by_day = pd.DataFrame(columns=["tanggal"])
        by_outlet = pd.DataFrame(columns=["tanggal"])
    shift_rows = []
    for (c, d), s in inj.shifts.items():
        if (c, d) in groups:
            shift_rows.append({"tanggal": d, "company": c,
                               "opening": s["name"],
                               "closing": s.get("pos_closing_shift") or "",
                               "shift_posting": str(s.get("posting_date"))[:10]})
    out = os.path.join(HERE, f"PosInvoice_Lokal_{tag}.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        by_day.to_excel(w, sheet_name="Rekap_Harian", index=False)
        by_outlet.to_excel(w, sheet_name="Rekap_Outlet", index=False)
        pd.DataFrame(shift_rows).to_excel(w, sheet_name="Shift", index=False)
        pd.DataFrame(results).to_excel(w, sheet_name="Raw_Invoice", index=False)
        pd.DataFrame(skipped).to_excel(w, sheet_name="Skipped", index=False)
        pd.DataFrame([{"cek": n, "hasil": "LOLOS" if ok else "GAGAL"}
                      for n, ok in checks]).to_excel(
            w, sheet_name="Validation", index=False)
        for ws in w.sheets.values():
            style_sheet(ws)
    log(f"Excel: {out}")
    log(f"total {time.time() - t0:.0f}s | OK {counters['OK']} | "
        f"ADA {counters['ADA']} | GAGAL {counters['GAGAL']} | "
        f"BEDA {counters['BEDA_TOTAL']} | skip {len(skipped)}")
    if not all(ok for _, ok in checks):
        sys.exit(1)


if __name__ == "__main__":
    main()
