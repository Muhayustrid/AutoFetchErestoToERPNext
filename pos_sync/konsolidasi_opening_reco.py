#!/usr/bin/env python3
"""pos_sync/konsolidasi_opening_reco.py — konsolidasi Stock Reconciliation JURI.

Instruksi user 2026-10-09:
  - Akun RECO salah: dokumen "Stock Reconciliation" (penyesuaian) ada yang
    posting ke 1171 Pembukaan sementara; harusnya 5110.020 Penyesuaian Stock.
  - Opening harus JADI SATU per gudang @2026-09-30 23:00:00 (sekarang
    tersebar di 31 dokumen, jam 11:00-23:59 + 3 dok bertanggal 2 Okt).
  - Dokumen cancelled boleh dihapus permanen (aturan repo).

Fase:
  --plan (default, read-only)  -> reco_consolidate_plan.json + ringkasan
  --execute [--warehouses "A,B"] [--enable-negative-window]
      cancel dok lama -> submit opening gabungan per gudang -> verifikasi
      trail penuh -> poll repost.
      GATE: kalau subset ybs diprediksi kena NegativeStockError (validasi
      native v16, allow_negative_stock=0) dan --enable-negative-window
      TIDAK diberikan -> script REFUSE (STOP + lapor). Flag itu HANYA
      dinyalakan dengan izin user, durasi eksekusi saja, restore 0 di
      finally.
  --cleanup [--execute-cleanup]  hapus permanen dok RECO cancelled JURI
      (prosedur resmi: backup -> purge Print Log -> delete_linked_ledger
      -> delete -> restore -> cek ledger; target = CANCEL_A + SR7).
  --report     tulis Laporan_Konsolidasi_Opening_RECO.xlsx

Env: /tmp/eresto_prod.env (site BARU). State: reco_consolidate_state.json.
"""
import argparse
import json
import os
import sys
import time
import urllib.parse
from collections import defaultdict
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prod_common import Client, load_env  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
PLAN_JSON = os.path.join(HERE, "reco_consolidate_plan.json")
STATE_JSON = os.path.join(HERE, "reco_consolidate_state.json")
REPORT = os.path.join(HERE, "Laporan_Konsolidasi_Opening_RECO.xlsx")
BACKUP_DIR = os.path.join(HERE, "..", "audit_logs")

JURI = "PT. JUARA ROTI INDONESIA"
POST_DATE = "2026-09-30"
POST_TIME = "23:00:00"
ACC_OPEN = "1171.000 - Pembukaan sementara - JURI"
ACC_ADJ = "5110.020 - Penyesuaian Stock - JURI"
COST_CENTER = "Main - JURI"
MARK = "KONSOLIDASI OPENING 30SEP"
MARK_SR7 = "KONSOLIDASI ADJ 5110"
EPS = 0.004

# 31 dokumen opening lama -> dibatalkan lalu DIHAPUS (scope: JURI)
CANCEL_A = [
    "RECO-JURI-261006001", "RECO-JURI-261006003", "RECO-JURI-261006005",
    "RECO-JURI-261006006", "RECO-JURI-261006007", "RECO-JURI-261006008",
    "RECO-JURI-261007001", "RECO-JURI-261007002", "RECO-JURI-261007003",
    "RECO-JURI-261007007", "RECO-JURI-261008001", "RECO-JURI-261009001",
    "RECO-JURI-261009002", "RECO-JURI-261009003", "RECO-JURI-261009004",
    "RECO-JURI-261009005", "RECO-JURI-261009006", "RECO-JURI-261009008",
    "RECO-JURI-261009009", "RECO-JURI-261009010", "RECO-JURI-261009011",
    "RECO-JURI-261009012", "RECO-JURI-261009013", "RECO-JURI-261009017",
    "RECO-JURI-261009018", "RECO-JURI-261009020", "RECO-JURI-261009022",
    "RECO-JURI-261009035", "RECO-JURI-261009036", "RECO-JURI-261009037",
    "RECO-JURI-261009039",
]
# 4 dokumen tidak disentuh (dough GBJ dulu + 2 koreksi 5110 yang sudah benar)
KEEP_A = ["RECO-JURI-261006002", "RECO-JURI-261009014",
          "RECO-JURI-261009040", "RECO-JURI-261009041"]
# 7 dokumen penyesuaian salah akun -> dibuat ulang dengan 5110.020
SR7 = ["RECO-JURI-261009026", "RECO-JURI-261009029", "RECO-JURI-261009030",
       "RECO-JURI-261009031", "RECO-JURI-261009032", "RECO-JURI-261009033",
       "RECO-JURI-261009034"]
ALL_DOCS = CANCEL_A + KEEP_A + SR7

SLE_F = ["name", "posting_date", "posting_time", "creation", "voucher_type",
         "voucher_no", "item_code", "warehouse", "actual_qty",
         "qty_after_transaction", "stock_value", "stock_value_difference"]
ROW_F = ["item_code", "warehouse", "qty", "valuation_rate", "uom",
         "batch_no", "allow_zero_valuation_rate"]


def log(m):
    print(m, flush=True)


def tosec(t):
    s = str(t or "0").split(".")
    p = s[0].split(":")
    while len(p) < 3:
        p.append("0")
    try:
        return int(float(p[0])) * 3600 + int(float(p[1])) * 60 + float(p[2])
    except Exception:
        return -1


def key(r):
    return (str(r["posting_date"])[:10], tosec(r["posting_time"]),
            str(r.get("creation")), r["name"])


T0 = (POST_DATE, tosec(POST_TIME))
EOD = (POST_DATE, 24 * 3600 - 1)


def f2(t):
    return "%.2f" % float(t or 0)


def fetch(c, dt, fields, filters, order, page=500):
    out, start = [], 0
    while True:
        st, d = c.get("/api/resource/" + urllib.parse.quote(dt), {
            "fields": json.dumps(fields), "filters": json.dumps(filters),
            "order_by": order, "limit_page_length": page,
            "limit_start": start})
        rows = d.get("data", []) if st == 200 else []
        if st != 200:
            raise SystemExit("GET %s gagal HTTP %s %s" % (dt, st, str(d)[:300]))
        out += rows
        start += len(rows)
        if len(rows) < page:
            break
    return out


def fetch_sle(c):
    return fetch(c, "Stock Ledger Entry", SLE_F,
                 [["company", "=", JURI], ["is_cancelled", "=", 0]],
                 "posting_date asc, posting_time asc, creation asc, name asc")


def build_idx(sle):
    idx = defaultdict(list)
    for r in sle:
        idx[(r["warehouse"], r["item_code"])].append(r)
    for k in idx:
        idx[k].sort(key=key)
    return idx


def value_at(rows, cut):
    best = None
    for r in rows:
        if key(r)[:2] <= cut:
            best = r
    return float(best["qty_after_transaction"] or 0) if best else 0.0


def get_doc(c, name):
    st, d = c.get("/api/resource/Stock Reconciliation/"
                  + urllib.parse.quote(name))
    if st != 200:
        raise SystemExit("GET doc %s gagal HTTP %s" % (name, st))
    return d["data"]


def load_state():
    try:
        return json.load(open(STATE_JSON))
    except (OSError, ValueError):
        return {"cancelled": [], "created": {}, "sr7": {}, "snapshot": None}


def save_state(st):
    json.dump(st, open(STATE_JSON, "w"), indent=1, default=str)


def snapshot_full(c):
    sle = fetch_sle(c)
    idx = build_idx(sle)
    rows_out = {}
    for (wh, it), rows in idx.items():
        rows_out["%s|%s" % (wh, it)] = [
            [str(r["posting_date"])[:10], str(r["posting_time"]),
             r["voucher_no"], str(r["actual_qty"]), str(r["qty_after_transaction"])]
            for r in rows]
    vals = {}
    for (wh, it), rows in idx.items():
        vals["%s|%s" % (wh, it)] = [value_at(rows, T0), value_at(rows, EOD),
                                    float(rows[-1]["qty_after_transaction"] or 0)]
    bins = {(r["item_code"], r["warehouse"]):
            [round(float(r["actual_qty"] or 0), 3),
             round(float(r.get("stock_value") or 0), 2)]
            for r in fetch(c, "Bin",
                           ["item_code", "warehouse", "actual_qty",
                            "stock_value", "valuation_rate"],
                           [["company", "=", JURI], ["actual_qty", "!=", 0]],
                           "name asc")}
    neg = sum(1 for r in sle
              if str(r["posting_date"])[:10] >= "2026-10-01"
              and float(r["qty_after_transaction"] or 0) < -EPS)
    return {"values": vals,
            "rows": rows_out,
            "bins": {"%s|%s" % k: v for k, v in bins.items()},
            "negrows": neg}


def poll_repost(c, tries=40, label=""):
    for i in range(1, tries + 1):
        st, d = c.get("/api/resource/Repost Item Valuation", {
            "filters": json.dumps([["company", "=", JURI],
                                   ["status", "in", ["Queued", "In Progress"]]]),
            "fields": json.dumps(["name"]), "limit_page_length": 1})
        n = len(d.get("data", [])) if st == 200 else -1
        if n == 0:
            log("  [repost] selesai %s" % label)
            return
        if i == 8:  # safety: force worker
            c.post("/api/method/erpnext.stock.doctype.repost_item_valuation"
                   ".repost_item_valuation.execute_repost_item_valuation", {})
        log("  [repost] ... masih %s antre (i=%d) %s" % (n, i, label))
        time.sleep(10)


# ------------------------------------------------------------------ plan
def cancels_blocked(idx, doc):
    """prediksi baris yang ditolak validate_negative_qty_in_future_sle v16."""
    T = (str(doc["posting_date"])[:10], tosec(doc["posting_time"]))
    n = 0
    for row in doc.get("items", []):
        wh, it = row["warehouse"], row["item_code"]
        q = float(row.get("qty") or 0)
        if not q:
            continue
        R = [r for r in idx.get((wh, it), [])
             if r["voucher_no"] != doc["name"]]
        B = None
        for r in R:
            if r["voucher_type"] == "Stock Reconciliation" and key(r)[:2] > T:
                if B is None or key(r) < B:
                    B = key(r)
        for r in R:
            k2 = key(r)
            if k2[:2] < T:
                continue
            v2 = float(r["qty_after_transaction"] or 0)
            shifted = v2 - q if (k2[:2] > T and (B is None or k2 < B)) else v2
            if shifted < -EPS:
                n += 1
                break
    return n


def build_plan(c):
    log("== fetch SLE + dokumen ==")
    sle = fetch_sle(c)
    idx = build_idx(sle)
    log("  SLE aktif: %d | posisi: %d" % (len(sle), len(idx)))
    docs = {}
    for n in ALL_DOCS:
        docs[n] = get_doc(c, n)
    log("  dokumen: %d (docstatus=1: %d)"
        % (len(docs), sum(1 for d in docs.values() if d["docstatus"] == 1)))

    child_map = {}   # (voucher, wh, item) -> child row
    for n, d in docs.items():
        for row in d.get("items", []):
            child_map[(n, row["warehouse"], row["item_code"])] = row

    cancel_set = set(CANCEL_A)
    # posisi yang disentuh dok CANCEL_A
    touched = set()
    for n in CANCEL_A:
        for row in docs[n].get("items", []):
            if float(row.get("qty") or 0) != 0:
                touched.add((row["warehouse"], row["item_code"]))

    cons = defaultdict(list)
    cons_meta, skipped = {}, []
    for (wh, it) in sorted(touched):
        rows = idx.get((wh, it), [])
        anchors = [r for r in rows if r["voucher_no"] in cancel_set]
        if not anchors:
            skipped.append((wh, it, "-", "tak ada anchor CANCEL_A terlihat"))
            continue
        G = anchors[-1]  # anchor CANCEL_A terakhir (paling baru)
        # ada reco lain (KEEP/SR7) SETELAH anchor? -> posisi sudah dipegang
        # dokumen itu, jangan diutak-atik (biar tidak dobel).
        later = [r for r in rows
                 if r["voucher_type"] == "Stock Reconciliation"
                 and key(r)[:2] > key(G)[:2]
                 and r["voucher_no"] not in cancel_set]
        if later:
            skipped.append((wh, it, G["voucher_no"],
                            "dipegang reco lain: %s"
                            % later[0]["voucher_no"]))
            continue
        ch = child_map.get((G["voucher_no"], wh, it))
        if ch is None:
            skipped.append((wh, it, G["voucher_no"], "child row hilang"))
            continue
        v2300 = value_at(rows, T0)
        vEOD = value_at(rows, EOD)
        gk = key(G)[:2]
        backdate = gk > EOD
        if not backdate:
            # anchor <= 30 Sep EOD: nilai = keadaan akhir 30 Sep (kenyataan
            # "malam itu"); sama dgn 23:00 kecuali kasus BB26008 (007007).
            qty_new = vEOD
            if abs(qty_new) < 1e-9:
                skipped.append((wh, it, G["voucher_no"],
                                "nilai EOD nol (anchor lama)"))
                continue
        else:
            # anchor SETELAH 30 Sep (mis. dok 2 Okt) -> backdate ke 23:00.
            # koreksi: kurangi pergerakan non-CANCEL antara 23:00..anchor.
            m = sum(float(r["actual_qty"] or 0)
                    for r in rows
                    if T0 < key(r)[:2] < gk
                    and r["voucher_no"] not in cancel_set)
            qty_new = float(ch.get("qty") or 0) - m
            if abs(qty_new) < 1e-9:
                skipped.append((wh, it, G["voucher_no"],
                                "nilai backdate nol"))
                continue
        rate = float(ch.get("valuation_rate") or 0)
        row = {"item_code": it, "warehouse": wh,
               "qty": round(qty_new, 6)}
        if rate > 0:
            row["valuation_rate"] = rate
        elif ch.get("allow_zero_valuation_rate"):
            row["allow_zero_valuation_rate"] = 1
        if ch.get("uom"):
            row["uom"] = ch["uom"]
        if ch.get("batch_no"):
            row["batch_no"] = ch["batch_no"]
            row["use_serial_batch_fields"] = 1
        cons[wh].append(row)
        cons_meta["%s|%s" % (wh, it)] = {
            "src": G["voucher_no"], "qty": qty_new,
            "rate": rate, "v2300": v2300,
            "chg2300": abs(v2300 - qty_new) > 1e-6}

    # dokumen yang harus dibatalkan per gudang (dok satu gudang)
    wh_docs = defaultdict(list)
    for n in CANCEL_A:
        d = docs[n]
        if d["docstatus"] != 1:
            continue
        wh = d.get("set_warehouse") or (
            d["items"][0]["warehouse"] if d.get("items") else None)
        if not wh:
            skipped.append(("-", "-", n, "tanpa gudang (dok manual)"))
            continue
        wh_docs[wh].append(n)

    # urutan cancel WAJIB ascending (creation, name): ERPNext v16 membatasi
    # shift cancel di "RECO berikutnya" yang masih visible; meng-cancel
    # dokumen termuda dulu membuat sisa dokumen tua men-shift penuh.
    order = {n: (str(docs[n].get("creation")), n) for n in CANCEL_A}
    for w in wh_docs:
        wh_docs[w] = sorted(wh_docs[w], key=lambda n: order[n])
    global_order = sorted([n for n in CANCEL_A if docs[n]["docstatus"] == 1],
                          key=lambda n: order[n])

    pred_cancel = {n: (cancels_blocked(idx, docs[n])
                       if docs[n]["docstatus"] == 1 else 0)
                   for n in CANCEL_A}
    pred_sr7 = {n: (cancels_blocked(idx, docs[n])
                    if docs[n]["docstatus"] == 1 else 0) for n in SR7}
    pred_new = {}
    for wh, rws in cons.items():
        n = 0
        for row in rws:
            R = idx.get((row["warehouse"], row["item_code"]), [])
            if any(key(r)[:2] >= T0
                   and float(r["qty_after_transaction"] or 0) < -EPS
                   for r in R):
                n += 1
        pred_new[wh] = n

    plan = {
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "cancel_docs": CANCEL_A, "keep_docs": KEEP_A, "sr7_docs": SR7,
        "consolidated": {w: sorted(r, key=lambda x: x["item_code"])
                         for w, r in cons.items()},
        "wh_docs": {w: sorted(v) for w, v in wh_docs.items()},
        "cons_meta": cons_meta,
        "pred_cancel_blocks": pred_cancel,
        "pred_sr7_blocks": pred_sr7,
        "pred_new_blocks": pred_new,
        "skipped": skipped,
        "sr7_def": {n: {"posting_date": str(docs[n]["posting_date"]),
                        "posting_time": str(docs[n]["posting_time"]),
                        "purpose": docs[n]["purpose"],
                        "set_warehouse": docs[n].get("set_warehouse"),
                        "cost_center": docs[n].get("cost_center")}
                    for n in SR7},
        "sr7_rows": {n: [{k2: r.get(k2) for k2 in ROW_F}
                         for r in docs[n].get("items", [])] for n in SR7},
    }
    json.dump(plan, open(PLAN_JSON, "w"), indent=1, default=str)
    return plan


def print_plan(plan):
    log("")
    log("=" * 72)
    log("=== RENCANA KONSOLIDASI RECO JURI ===")
    log("  cancel  : %d dokumen opening lama" % len(plan["cancel_docs"]))
    log("  keep    : %s" % ", ".join(plan["keep_docs"]))
    log("  recreate: %d dok 5110: %s" % (len(plan["sr7_docs"]),
                                        ", ".join(plan["sr7_docs"])))
    log("")
    log("  %-38s %5s %6s %8s  %s" % ("OPENING BARU / gudang", "rows",
                                     "chg23", "blokNEW", "dok lama dibatalkan"))
    tot = 0
    for w in sorted(plan["consolidated"]):
        rws = plan["consolidated"][w]
        tot += len(rws)
        chg = sum(1 for r in rws if plan["cons_meta"]
                  ["%s|%s" % (w, r["item_code"])]["chg2300"])
        log("  %-38s %5d %6d %8d  %s" % (
            w, len(rws), chg, plan["pred_new_blocks"].get(w, 0),
            ",".join(x[-6:] for x in plan["wh_docs"].get(w, []))))
    log("  TOTAL baris opening baru: %d" % tot)
    log("")
    log("  == prediksi blokir CANCEL (validasi native v16, flag=0) ==")
    bad = {k: v for k, v in plan["pred_cancel_blocks"].items() if v}
    for k in sorted(bad):
        log("    %-22s %d baris" % (k, bad[k]))
    log("  dok CANCEL_A tanpa risiko: %d/%d"
        % (len(plan["pred_cancel_blocks"]) - len(bad),
           len(plan["pred_cancel_blocks"])))
    sb = {k: v for k, v in plan["pred_sr7_blocks"].items() if v}
    for k in sorted(sb):
        log("    SR7 %-18s %d baris" % (k, sb[k]))
    nb = {k: v for k, v in plan["pred_new_blocks"].items() if v}
    log("  == prediksi blokir CREATE (baris minus historis) ==")
    for k in sorted(nb):
        log("    %-38s %d baris" % (k, nb[k]))
    if plan["skipped"]:
        log("  == dilewati (%d) ==" % len(plan["skipped"]))
        for s in plan["skipped"][:25]:
            log("    %s" % str(s))
    chg = [(k, m["v2300"], m["qty"], m["src"])
           for k, m in sorted(plan["cons_meta"].items()) if m["chg2300"]]
    log("  == posisi dgn nilai 23:00 berubah (backdate) ==")
    for k, v, q, src in chg:
        log("    %-52s %10.2f -> %10.2f  (%s)" % (k, v, q, src))
    log("OK — plan: %s" % PLAN_JSON)
    return bool(bad) or bool(nb) or bool(sb)


# ------------------------------------------------------------------ execute
def subst_material_receipt(c, sr7_name, posting_date, posting_time, rows):
    """Pulihkan efek zero-out SR7 yg tak bisa di-recreate tanpa flag:
    SE Material Receipt dgn |saldo negatif sekarang| & rate baris plan
    (saldo -> 0). Return nama SE atau None kalau tak ada yang negatif."""
    items = []
    for r in rows:
        wh, it = r["warehouse"], r["item_code"]
        sl = fetch(c, "Stock Ledger Entry",
                   ["name", "posting_date", "posting_time", "creation",
                    "qty_after_transaction"],
                   [["item_code", "=", it], ["warehouse", "=", wh],
                    ["is_cancelled", "=", 0]],
                   "posting_date asc, posting_time asc, creation asc, name asc")
        cur = float(sl[-1]["qty_after_transaction"] or 0) if sl else 0.0
        if cur < -EPS:
            rate = float(r.get("valuation_rate") or 0)
            if rate <= 0:
                # ambil rate Bin sekarang (biar nilai tetap sama)
                st2, d2 = c.get("/api/resource/Bin", {
                    "filters": json.dumps([["item_code", "=", it],
                                           ["warehouse", "=", wh]]),
                    "fields": json.dumps(["valuation_rate"])})
                rate = float((d2.get("data") or [{}])[0].get("valuation_rate") or 0)
            items.append({"item_code": it, "t_warehouse": wh,
                          "qty": abs(cur), "basic_rate": rate,
                          "expense_account": ACC_ADJ,
                          "cost_center": COST_CENTER})
    if not items:
        return None
    stn, dn = c.post("Stock Entry", {
        "doctype": "Stock Entry", "purpose": "Material Receipt",
        "stock_entry_type": "Material Receipt",
        "company": JURI, "set_posting_time": 1,
        "posting_date": posting_date, "posting_time": posting_time,
        "remarks": "%s SUBST %s" % (MARK_SR7, sr7_name),
        "items": items, "docstatus": 1})
    if stn not in (200, 201):
        log("     [WARN] substitusi %s gagal: HTTP %s %s"
            % (sr7_name, stn, str(dn)[:200]))
        return None
    return dn["data"]["name"]


def check_gate(plan, args):
    """Gate izin flag:
    - blokir pada dokumen GUDANG -> wajib --enable-negative-window (raise).
    - blokir pada SR7 -> per-dokumen SKIP (tidak raise); dilaporkan.
    - mode --sr7-only: cek SR7 saja.
    """
    sb = {k: v for k, v in plan["pred_sr7_blocks"].items() if v}
    if args.sr7_only:
        if sb and not args.enable_negative_window:
            log("  [WARN] SR7 butuh flag (akan di-SKIP, tidak di-cancel):")
            for k in sorted(sb):
                log("     %s (%d baris) -> skip" % (k, sb[k]))
        return {}, sb
    whs = plan["consolidated"]
    if args.warehouses:
        want = set(w.strip() for w in args.warehouses.split(",") if w.strip())
        whs = {w: v for w, v in whs.items() if w in want}
    bad_whs = {}
    for w in whs:
        for n in plan["wh_docs"].get(w, []):
            k = plan["pred_cancel_blocks"].get(n, 0)
            if k:
                bad_whs.setdefault(w, []).append((n, k))
        k = plan["pred_new_blocks"].get(w, 0)
        if k:
            bad_whs.setdefault(w, []).append(("(create)", k))
    if bad_whs and not args.enable_negative_window:
        log("")
        log("STOP: gudang berikut pasti ditolak dgn allow_negative_stock=0:")
        for w in sorted(bad_whs):
            for n, k in bad_whs[w]:
                log("   %-38s %-24s %d baris" % (w, n, k))
        log("  -> butuh --enable-negative-window (HANYA dengan izin user)")
        raise SystemExit(2)
    if sb and not args.enable_negative_window:
        log("  [WARN] SR7 butuh flag (akan di-SKIP, tidak di-cancel):")
        for k in sorted(sb):
            log("     %s (%d baris) -> skip" % (k, sb[k]))
    return bad_whs, sb


def execute(args):
    if not os.path.exists(PLAN_JSON):
        raise SystemExit("plan belum ada — jalankan --plan dulu")
    plan = json.load(open(PLAN_JSON))
    _, sr7_skip = check_gate(plan, args)

    prod = load_env("/tmp/eresto_prod.env")
    c = Client(prod["ERPNEXT_URL"], prod["ERPNEXT_API_KEY"], "BARU")
    st = load_state()

    if args.sr7_only:
        whs = []
        log("  mode: SR7-only (5 dok yang aman; 032/034 di-skip)")
    elif args.warehouses:
        want = [w.strip() for w in args.warehouses.split(",") if w.strip()]
        whs = [w for w in want if w in plan["consolidated"]]
        log("  mode terbatas: %s" % whs)
    else:
        whs = sorted(plan["consolidated"])

    log("")
    log("== SNAPSHOT pra-eksekusi ==")
    if not st.get("snapshot"):
        st["snapshot"] = snapshot_full(c)
        save_state(st)
    snap = st["snapshot"]
    pre_rows = snap["rows"]
    pre_vals = {k: tuple(v) for k, v in snap["values"].items()}
    pre_bins = snap["bins"]
    log("  posisi=%d baris-trail=%d Bin=%d baris minus(>=1 Okt)=%s"
        % (len(pre_vals), sum(len(v) for v in pre_rows.values()),
           len(pre_bins), snap.get("negrows")))

    log("")
    log("== kuras antrean repost dulu ==")
    poll_repost(c, tries=12, label="pre")

    flag_restore = False
    try:
        if args.enable_negative_window:
            log("")
            log("== FLAG WINDOW: allow_negative_stock 0 -> 1 (sementara) ==")
            st2, _ = c.put("Stock Settings", "Stock Settings",
                           {"allow_negative_stock": 1})
            log("  set 1: HTTP %s" % st2)
            flag_restore = True

        for wh in whs:
            log("")
            log("=== [%s] ===" % wh)
            for n in plan["wh_docs"].get(wh, []):
                if n in st["cancelled"]:
                    continue
                stc, d = c.put("Stock Reconciliation", n, {"docstatus": 2})
                if stc != 200:
                    log("  [FAIL] cancel %s: HTTP %s %s"
                        % (n, stc, str(d)[:250]))
                    raise SystemExit("STOP: cancel gagal (cek flag/negatif)")
                st["cancelled"].append(n)
                save_state(st)
                log("  [OK] cancelled %s" % n)

            rows = plan["consolidated"].get(wh, [])
            name = st["created"].get(wh)
            if name:
                log("  [SKIP] opening sudah ada: %s" % name)
            elif rows:
                payload = {
                    "doctype": "Stock Reconciliation", "company": JURI,
                    "purpose": "Opening Stock", "set_posting_time": 1,
                    "posting_date": POST_DATE, "posting_time": POST_TIME,
                    "set_warehouse": wh, "expense_account": ACC_OPEN,
                    "cost_center": COST_CENTER,
                    "remarks": "%s %s" % (MARK, wh),
                    "items": rows, "docstatus": 1,
                }
                stn, dn = c.post("Stock Reconciliation", payload)
                if stn not in (200, 201):
                    raise SystemExit("GAGAL submit opening %s: HTTP %s %s"
                                     % (wh, stn, str(dn)[:400]))
                name = dn["data"]["name"]
                st["created"][wh] = name
                save_state(st)
                log("  [OK] opening baru: %s (%d baris)" % (name, len(rows)))
            else:
                log("  [INFO] tidak ada baris utk gudang ini")

            poll_repost(c, tries=20, label=wh)
            newnames = [v for k2, v in st["created"].items() if k2 == wh]
            bad = verify_slice(c, plan, pre_rows, pre_vals, {wh}, newnames)
            if bad:
                log("  [FAIL] verifikasi %s: %d masalah" % (wh, len(bad)))
                for b in bad[:15]:
                    log("     %s" % b)
                raise SystemExit("STOP: verifikasi gudang %s tidak lolos" % wh)
            log("  [OK] verifikasi %s lolos" % wh)

        if (not args.warehouses) or args.sr7_only:
            log("")
            log("=== recreate dokumen penyesuaian (akun 5110) ===")
            for n in SR7:
                if n in st["sr7"]:
                    continue
                if n in sr7_skip and not args.enable_negative_window:
                    log("  [SKIP] %s (butuh flag; %d baris)"
                        % (n, sr7_skip[n]))
                    continue
                stc, d = c.get("/api/resource/Stock Reconciliation/"
                               + urllib.parse.quote(n))
                if stc == 200 and d["data"]["docstatus"] == 1:
                    stp, dd = c.put("Stock Reconciliation", n, {"docstatus": 2})
                    if stp != 200:
                        raise SystemExit("STOP: cancel %s HTTP %s" % (n, stp))
                    if n not in st["cancelled"]:
                        st["cancelled"].append(n)
                    log("  [OK] cancelled %s" % n)
                defd = plan["sr7_def"][n]
                rows = [dict(r) for r in plan["sr7_rows"][n]]
                for r in rows:
                    r.pop("batch_no", None)
                    if float(r.get("valuation_rate") or 0) <= 0:
                        r["allow_zero_valuation_rate"] = 1
                payload = {
                    "doctype": "Stock Reconciliation", "company": JURI,
                    "purpose": defd["purpose"], "set_posting_time": 1,
                    "posting_date": defd["posting_date"],
                    "posting_time": defd["posting_time"],
                    "expense_account": ACC_ADJ,
                    "cost_center": defd.get("cost_center") or COST_CENTER,
                    "remarks": "%s %s" % (MARK_SR7, n),
                    "items": rows, "docstatus": 1,
                }
                if defd.get("set_warehouse"):
                    payload["set_warehouse"] = defd["set_warehouse"]
                stn, dn = c.post("Stock Reconciliation", payload)
                if stn not in (200, 201):
                    txt = str(dn)
                    if "EmptyStockReconciliationItemsError" in txt:
                        log("  [NOOP] %s: tak ada perubahan qty/value — "
                            "dok kosong, dibiarkan cancelled (dihapus saat "
                            "cleanup)" % n)
                        st.setdefault("sr7_noop", {})[n] = True
                        save_state(st)
                        continue
                    if "NegativeStockError" in txt:
                        # cancel dok lama meninggalkan saldo negatif di anchor
                        # (zero-out) -> recreate MUSTAHIL dgn flag=0 (validasi
                        # anchor v16). Pulihkan via SE Material Receipt dgn
                        # qty = |saldo sekarang| & rate yang sama (bin -> 0,
                        # GL setara; akun tetap 5110).
                        log("  [SUBST] %s: zero-out tak bisa di-recreate tanpa "
                            "flag — pakai substitusi SE Material Receipt" % n)
                        stn2 = subst_material_receipt(
                            c, n, defd["posting_date"], defd["posting_time"],
                            rows)
                        if stn2:
                            st["sr7"][n] = stn2
                            save_state(st)
                            log("  [OK] substitusi %s -> %s" % (n, stn2))
                        continue
                    raise SystemExit("GAGAL recreate %s: HTTP %s %s"
                                     % (n, stn, txt[:300]))
            poll_repost(c, tries=30, label="SR7")

        log("")
        if args.sr7_only:
            end_scope = set()
            pos_filter = set()
            for n, newname in st["sr7"].items():
                for r in plan["sr7_rows"][n]:
                    end_scope.add(r["warehouse"])
                    pos_filter.add("%s|%s" % (r["warehouse"], r["item_code"]))
            log("== VERIFIKASI AKHIR (scope SR7: %s) ==" % sorted(end_scope))
            bad = verify_slice(c, plan, pre_rows, pre_vals, end_scope,
                               list(st["sr7"].values()), pos_filter)
        else:
            log("== VERIFIKASI AKHIR (scope %s) =="
                % ("semua" if not args.warehouses else str(whs)))
            end_scope = set(whs) if args.warehouses else None
            bad = verify_slice(c, plan, pre_rows, pre_vals, end_scope,
                               list(st["created"].values()))
        if bad:
            log("  [FAIL] %d masalah" % len(bad))
            for b in bad[:30]:
                log("     %s" % b)
            raise SystemExit("STOP: verifikasi akhir tidak lolos")
        log("  [LOLOS] trail > 30 Sep identik; checkpoints sesuai rencana")
        post = snapshot_full(c)
        post_bins = post["bins"]
        diff = {k: (pre_bins.get(k), post_bins.get(k))
                for k in set(pre_bins) | set(post_bins)
                if pre_bins.get(k) != post_bins.get(k)}
        log("  [%s] Bin banding: %d selisih"
            % ("LOLOS" if not diff else "FAIL", len(diff)))
        for k, v in list(diff.items())[:10]:
            log("     %s: %s" % (k, v))
        log("  [INFO] baris minus (>=1 Okt): pra=%s pasca=%s"
            % (snap.get("negrows"), post["negrows"]))
    finally:
        if flag_restore:
            c.put("Stock Settings", "Stock Settings",
                  {"allow_negative_stock": 0})
            stx, d = c.get("/api/method/frappe.client.get_value", {
                "doctype": "Stock Settings",
                "fieldname": "allow_negative_stock"})
            cur = (d.get("message") or {}).get("allow_negative_stock")
            log("")
            log("== FLAG WINDOW ditutup: allow_negative_stock=%s ==" % cur)
        save_state(st)
    log("")
    log("OK — lanjutkan: --cleanup lalu --report")


def verify_slice(c, plan, pre_rows, pre_vals, scope_whs, created_names=None,
                 pos_filter=None):
    """Verifikasi tepat:
    - posisi di luar scope/pos_filter tidak diaudit.
    - baris trail setelah timestamp dokumen tersentuh terakhir (L) harus
      IDENTIK (voucher, jam, actual, qty_after).
    - checkpoint 23:00 posisi konsolidasi = qty plan; EOD & NOW tidak berubah.
    """
    sle = fetch_sle(c)
    idx = build_idx(sle)
    vals_new = {}
    for (wh, it), rows in idx.items():
        vals_new["%s|%s" % (wh, it)] = [
            value_at(rows, T0), value_at(rows, EOD),
            float(rows[-1]["qty_after_transaction"] or 0)]
    cons_qty = {}
    for w, rws in plan["consolidated"].items():
        if scope_whs and w not in scope_whs:
            continue
        for r in rws:
            cons_qty["%s|%s" % (w, r["item_code"])] = r["qty"]

    # dokumen tersentuh dlm scope (cancel + recreate + nama baru)
    scope_docs = set(created_names or [])
    scope_docs.update(plan["sr7_docs"])          # mungkin sudah di-cancel
    scope_docs.update(plan.get("_sr7new", {}).values())
    for w in (scope_whs if scope_whs else plan["consolidated"]):
        scope_docs.update(plan["wh_docs"].get(w, []))
    L = list(T0)
    for k, rows in pre_rows.items():
        for r in rows:
            if r[2] in scope_docs:
                kk = (r[0], tosec(r[1]))
                if kk > tuple(L):
                    L = list(kk)
    L = tuple(L)

    problems = []
    for k in set(list(pre_vals) + list(vals_new)):
        if pos_filter is not None and k not in pos_filter:
            continue
        wh = k.split("|")[0]
        if scope_whs and wh not in scope_whs:
            continue
        pre = pre_vals.get(k)
        new = vals_new.get(k)
        if pre is None:
            problems.append("%s: posisi baru muncul %s" % (k, new))
            continue
        if new is None:
            problems.append("%s: posisi hilang" % k)
            continue
        v23o, eodo, nowo = pre
        v23n, eodn, nown = new
        meta = plan["cons_meta"].get(k) or {}
        # posisi backdate (sumber dok setelah 30 Sep) SAH pindah EOD;
        # posisi lain EOD wajib tetap.
        if abs(eodn - eodo) > EPS and not meta.get("chg2300"):
            problems.append("%s: EOD %.3f -> %.3f" % (k, eodo, eodn))
        if abs(nown - nowo) > EPS:
            problems.append("%s: NOW %.3f -> %.3f" % (k, nowo, nown))
        if k in cons_qty:
            if abs(v23n - cons_qty[k]) > EPS:
                problems.append("%s: 23:00 %.3f != rencana %.3f"
                                % (k, v23n, cons_qty[k]))
        else:
            if abs(v23n - v23o) > EPS:
                problems.append("%s: 23:00 berubah %.3f -> %.3f"
                                % (k, v23o, v23n))
        # trail kunci setelah L wajib identik (voucher, jam, actual, after)
        old_after = [r for r in pre_rows.get(k, [])
                     if (r[0], tosec(r[1])) > L]
        new_after = [[str(r["posting_date"])[:10], str(r["posting_time"]),
                      r["voucher_no"], str(r["actual_qty"]),
                      str(r["qty_after_transaction"])]
                     for r in idx.get(tuple(k.split("|")), [])
                     if (str(r["posting_date"])[:10], tosec(r["posting_time"])) > L]
        if len(old_after) != len(new_after):
            problems.append("%s: jumlah baris >L beda %d->%d"
                            % (k, len(old_after), len(new_after)))
            continue
        for a, b in zip(old_after, new_after):
            if (a[2] != b[2]
                    or abs(float(a[3]) - float(b[3])) > 1e-6
                    or abs(float(a[4]) - float(b[4])) > EPS):
                problems.append("%s: baris >L beda %s vs %s" % (k, a, b))
                break
    return problems


# ------------------------------------------------------------------ cleanup
def cleanup(args):
    st = load_state()
    prod = load_env("/tmp/eresto_prod.env")
    c = Client(prod["ERPNEXT_URL"], prod["ERPNEXT_API_KEY"], "BARU")
    expect = set(CANCEL_A) | set(SR7)

    rows = fetch(c, "Stock Reconciliation",
                 ["name", "docstatus", "company"],
                 [["docstatus", "=", 2], ["company", "=", JURI]], "name asc")
    names = sorted(r["name"] for r in rows)
    log("== dokumen RECO cancelled JURI: %d ==" % len(names))
    for n in names:
        log("  %s %s" % ("[target]" if n in expect else "[EXTRA]", n))
    if not names:
        log("[SKIP] tidak ada dokumen cancelled.")
        return
    todo = [n for n in names if n in expect or n in st.get("cancelled", [])]
    if len(todo) != len(names):
        raise SystemExit("STOP: ada dok cancelled di luar daftar target — "
                         "periksa manual dulu.")
    if not args.execute_cleanup:
        log("")
        log("[DRY-RUN] jalankan --cleanup --execute-cleanup utk menghapus.")
        return

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(BACKUP_DIR, "backup_reco_consolidated_%s.json" % stamp)
    os.makedirs(BACKUP_DIR, exist_ok=True)
    backup = {n: get_doc(c, n) for n in todo}
    json.dump(backup, open(path, "w"), indent=1, default=str)
    log("[OK] backup %d dok -> %s" % (len(backup), os.path.abspath(path)))

    purged = 0
    for n in todo:
        stt, d = c.get("/api/resource/POS%20Print%20Log", {
            "filters": json.dumps([["reference_name", "=", n]]),
            "fields": json.dumps(["name"]), "limit_page_length": 0})
        for r in (d.get("data") or []):
            c.req("DELETE", "/api/resource/POS%20Print%20Log/"
                  + urllib.parse.quote(r["name"]))
            purged += 1
    log("[OK] POS Print Log dipurge: %d" % purged)

    c.put("Accounts Settings", "Accounts Settings",
          {"delete_linked_ledger_entries": 1})
    deleted = []
    try:
        for n in todo:
            stt, dd = c.req("DELETE", "/api/resource/Stock%20Reconciliation/"
                            + urllib.parse.quote(n))
            ok = stt in (200, 202)
            log("  [%s] delete %s: %s" % ("OK" if ok else "FAIL", n, stt))
            deleted.append((n, ok))
            time.sleep(0.05)
    finally:
        c.put("Accounts Settings", "Accounts Settings",
              {"delete_linked_ledger_entries": 0})
        log("[OK] restore delete_linked_ledger_entries=0")

    log("")
    log("== CEK LEDGER per dokumen (harus 0) ==")
    okall = True
    for n, ok in deleted:
        if not ok:
            okall = False
            continue
        for ldt in ("GL Entry", "Stock Ledger Entry", "Payment Ledger Entry"):
            stt, dct = c.req("GET", "/api/method/frappe.client.get_count",
                             params={"doctype": ldt, "filters": json.dumps(
                                 [["voucher_no", "=", n]])})
            v = dct.get("message", 0) if stt == 200 else "?"
            if isinstance(v, int) and v > 0:
                okall = False
                log("  [WARN] %s sisa %s=%s" % (n, ldt, v))
    log("  [%s] ledger bersih" % ("LOLOS" if okall else "PERIKSA"))

    post = snapshot_full(c)
    pre_bins = (st.get("snapshot") or {}).get("bins") or {}
    post_bins = post["bins"]
    diff = {k: (pre_bins.get(k), post_bins.get(k))
            for k in set(pre_bins) | set(post_bins)
            if pre_bins.get(k) != post_bins.get(k)}
    log("  [%s] Bin vs snapshot pra: %d selisih"
        % ("LOLOS" if not diff else "PERIKSA", len(diff)))
    st["cleanup"] = {"deleted": [n for n, ok in deleted if ok],
                     "backup": os.path.abspath(path)}
    save_state(st)
    log("OK")
    log("Lanjut: python pos_sync/konsolidasi_opening_reco.py --report")


# ------------------------------------------------------------------ report
def report(args):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    plan = json.load(open(PLAN_JSON))
    st = load_state()
    prod = load_env("/tmp/eresto_prod.env")
    c = Client(prod["ERPNEXT_URL"], prod["ERPNEXT_API_KEY"], "BARU")
    active = fetch(c, "Stock Reconciliation",
                   ["name", "purpose", "posting_date", "posting_time",
                    "expense_account", "docstatus"],
                   [["company", "=", JURI]], "name asc")

    wb = openpyxl.Workbook()
    hf = PatternFill("solid", fgColor="D9EAF7")
    bold = Font(bold=True)

    def sheet(name, cols, rows, widths):
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for i in range(1, len(cols) + 1):
            cc = ws.cell(row=1, column=i)
            cc.fill = hf
            cc.font = bold
            cc.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["Laporan", "Konsolidasi Opening RECO JURI 30 Sep 2026 23:00"])
    ws.append(["Instruksi", "user 2026-10-09 — opening satu per gudang, "
                            "akun 5110 utk penyesuaian, dok cancelled dihapus"])
    ws.append(["Dok dibatalkan", len(st.get("cancelled", []))])
    ws.append(["Opening baru dibuat", len(st.get("created", {}))])
    ws.append(["SR7 di-recreate (5110)", len(st.get("sr7", {}))])
    ws.append(["Dihapus permanen", len(st.get("cleanup", {}).get("deleted", []))])
    ws.append(["RECO aktif JURI sekarang",
               len([a for a in active if a["docstatus"] == 1])])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70

    sheet("Opening_Per_Gudang",
          ["Gudang", "Baris", "Doc Baru", "Dok Lama dibatalkan"],
          [[w, len(rws), st.get("created", {}).get(w, ""),
            ", ".join(plan["wh_docs"].get(w, []))]
           for w, rws in sorted(plan["consolidated"].items())],
          [38, 7, 24, 90])

    rows_all = []
    for w, rws in sorted(plan["consolidated"].items()):
        for r in rws:
            meta = plan["cons_meta"]["%s|%s" % (w, r["item_code"])]
            rows_all.append([r["item_code"], w, r["qty"],
                             r.get("valuation_rate", ""), meta["src"],
                             st.get("created", {}).get(w, ""),
                             "CHG" if meta["chg2300"] else ""])
    sheet("Rows", ["Item", "Gudang", "Qty", "Rate", "Sumber doc",
                   "Doc baru", "Cat"], rows_all,
          [16, 34, 12, 16, 24, 22, 6])

    sheet("SR7", ["Doc lama", "Doc baru (5110)", "Akun lama", "Akun baru"],
          [[n, st.get("sr7", {}).get(n, ""),
            "1171.000 - Pembukaan sementara", ACC_ADJ]
           for n in plan["sr7_docs"]], [24, 24, 34, 34])

    sheet("Dihapus", ["Dokumen"],
          [[n] for n in st.get("cleanup", {}).get("deleted", [])], [26])

    sheet("Sisa_Reco",
          ["Dokumen", "Purpose", "Tanggal", "Jam", "Akun"],
          [[a["name"], a["purpose"], str(a["posting_date"]),
            str(a["posting_time"]),
            "5110" if "5110" in (a["expense_account"] or "") else
            ("1171" if "1171" in (a["expense_account"] or "") else "")]
           for a in active if a["docstatus"] == 1],
          [26, 20, 12, 16, 10])

    wb.save(REPORT)
    log("Laporan ditulis: %s" % REPORT)
    log("OK")


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--plan", action="store_true")
    ap.add_argument("--execute", action="store_true")
    ap.add_argument("--warehouses", default=None)
    ap.add_argument("--sr7-only", action="store_true",
                    help="hanya recreate dokumen penyesuaian (5110)")
    ap.add_argument("--enable-negative-window", action="store_true")
    ap.add_argument("--cleanup", action="store_true")
    ap.add_argument("--execute-cleanup", action="store_true")
    ap.add_argument("--report", action="store_true")
    args = ap.parse_args()
    prod = load_env("/tmp/eresto_prod.env")
    c = Client(prod["ERPNEXT_URL"], prod["ERPNEXT_API_KEY"], "BARU")

    if args.execute:
        execute(args)
    elif args.cleanup:
        cleanup(args)
    elif args.report:
        report(args)
    else:
        log("=" * 72)
        log("FASE PLAN (dry-run; tidak menulis ke situs)")
        plan = build_plan(c)
        print_plan(plan)
        return


if __name__ == "__main__":
    main()
