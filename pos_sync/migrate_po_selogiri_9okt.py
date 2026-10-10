#!/usr/bin/env python3
"""PO eResto PO222.091026.001/.002 (RS MUHAMMADIYAH SELOGIRI) -> ERPNext Baru.

Konteks (lihat memori `po-eresto-internal-gudang-mapping` + PROJECT_STATE):
- Dua PO outlet baru di eResto tanggal 9 Okt 2026:
    * PO222.091026.001  supplier 'Gudang Outlet - PH Klaten'   Rp 6.078.849
    * PO222.091026.002  supplier 'Gudang Peralatan Klaten'     Rp 1.958.000
  Keduanya divisi pemesan RS MUHAMMADIYAH SELOGIRI, status eResto 'processing'.
- Company 'RS MUHAMMADIYAH SELOGIRI' belum ada di BARU (TUNDA di Kandidat
  karena 0 trx POS 30 hari) -> di-onboard dulu persis pola RSUD KRAMAT JATI:
  Company (anak OUTLET) + Fiscal Year 2026 + Allowed To Transact With di
  supplier JURI + customer internal.
- Pemetaan dokumen internal gudang (sama dengan `migrate_po_gudang_jakarta_9okt.py`):
    * Company       = divisi outlet pemesan
    * Supplier      = 'PT. JUARA ROTI INDONESIA'
    * set_warehouse = 'Stores - <ABBR>'
    * set_from_warehouse = gudang asal JURI ('Gudang Outlet - JURI' /
      'Gudang Peralatan - JURI')
    * Buying Price List = 'Harga Outlet'; beban '4210.000 - HPP Pembelian -
      <ABBR>'; cost center 'Main - <ABBR>'.
- ID dokumen = nomor PO eResto (rename_doc). status processing/completed ->
  submit (1); lainnya -> draft (0).
- Idempoten: dokumen yang sudah ada (by name ATAU title 'ERESTO <PO>') di-skip;
  rerun melakukan self-heal rename bila dokumen tertinggal dengan nama fallback.

Bukti pemetaan SKU peralatan (bukan tebakan): PO lama `PO332.100926.002`
(RS PKU AISYIYAH BOYOLALI, Gudang Outlet - PH Klaten) berisi set item yang
IDENTIK 1:1 dengan PO222.091026.001 (nama/qty/harga klop), dan di ERPNext
site lama tersimpan sebagai PEL260056/57/58/59/60. BB0291->PER260004 sudah
ada di ITEM_MAP `migrate_po_eresto_8_9_okt.py` (PO223/PO332 seri sama).

CLI:
  python pos_sync/migrate_po_selogiri_9okt.py [--dry-run] [--verify]
      [--only PO222.091026.001 ...]

Env: /tmp/eresto_mig.env (ERESTO_*) + /tmp/eresto_prod.env (ERPNEXT_URL,
ERPNEXT_API_KEY). Token tidak pernah ditulis ke output.
"""
import argparse
import json
import os
import sys
import time
import urllib.parse
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prod_common import Client, JURI  # noqa: E402

# ============================================================
# ENV & KONSTANTA
# ============================================================

def load_env_file(filepath):
    env = {}
    if os.path.exists(filepath):
        with open(filepath) as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env[k.strip()] = v.strip().strip('"').strip("'")
    return env

ENV_MIG = load_env_file("/tmp/eresto_mig.env")
ENV_PROD = load_env_file("/tmp/eresto_prod.env")

ERESTO_BASE_URL = "https://ropi.eresto.id"
ERESTO_DOMAIN = "ropi"
ERESTO_TOKEN = os.environ.get("ERESTO_TOKEN") or ENV_MIG.get("ERESTO_TOKEN", "")
ERESTO_EMAIL = os.environ.get("ERESTO_EMAIL") or ENV_MIG.get("ERESTO_EMAIL", "")
ERESTO_USER_ID = os.environ.get("ERESTO_USER_ID") or ENV_MIG.get("ERESTO_USER_ID", "")

ERPNEXT_URL = os.environ.get("ERPNEXT_URL") or ENV_PROD.get(
    "ERPNEXT_URL", "https://ropierpnext.frappe.cloud")
ERPNEXT_KEY = os.environ.get("ERPNEXT_API_KEY") or ENV_PROD.get("ERPNEXT_API_KEY", "")

PO_LIST = ["PO222.091026.001", "PO222.091026.002"]
PO_DATE = "2026-10-09"          # jangkar tarikan eResto (date + created_at window)
COMPANY = "RS MUHAMMADIYAH SELOGIRI"
ABBR = "RSMS"                   # unik di site (dicek vs seluruh 175 company)
PRICE_LIST = "Harga Outlet"
REPORT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "Laporan_Migrasi_PO_Selogiri_9Okt2026.xlsx")

# Supplier eResto (nama gudang) -> source warehouse JURI.
WAREHOUSE_MAP = {
    "Gudang Outlet - PH Klaten": "Gudang Outlet - JURI",
    "Gudang Peralatan Klaten": "Gudang Peralatan - JURI",
    "Gudang Barang Jadi - PH Klaten": "Gudang Barang Jadi - JURI",
    "Gudang Bahan Baku - PH Klaten": "Gudang Bahan Baku - JURI",
}

# SKU eResto -> Item Code ERPNext Baru (bukti: PO332.100926.002 site lama).
ITEM_MAP = {
    "BB0183": "PEL260058",  # 'Memory Card'        -> MMC CCTV (128 GB)
    "BB0199": "PEL260060",  # 'USB Adapter'
    "BB0217": "PEL260056",  # 'CCTV Outlet'
    "BB0238": "PEL260057",  # 'Alat Kasir'         -> Mesin Kasir POS
    "BB0262": "PEL260059",  # 'Orbit Outlet'       -> Modem WiFi Outlet
    "BB0291": "PER260004",  # 'Panel 3 Phase'
}

# (item_code, uom_eresto) -> (uom_erpnext, cf, catatan); kosong = semua unit
# ada langsung di master Item (terverifikasi).
UOM_SPECIAL = {}

# ============================================================
# KLIEN eRESTO (read-only)
# ============================================================

def safe_float(v):
    try:
        if v is None or v == "":
            return 0.0
        if isinstance(v, str):
            return float(v.strip().replace(",", "."))
        return float(v)
    except Exception:
        return 0.0


def safe_str(v):
    return "" if v is None else str(v).strip()


class ErestoClient:
    def __init__(self, base_url, token, email, user_id):
        self.base_url = base_url.rstrip("/")
        self.token = token
        self.email = email
        self.user_id = user_id
        self.session = requests.Session()
        self.session.headers.update({
            "accept": "application/json, text/plain, */*",
            "authorization": f"Token token={self.token}, email={self.email}",
            "referer": f"{self.base_url}/manager/v2/index.html",
        })

    def _fetch_paged(self, window):
        out = {}
        page = 1
        while True:
            params = dict(window)
            params.update({"current_user_id": self.user_id, "domain": ERESTO_DOMAIN,
                           "token": self.token, "per_page": 200, "page": page})
            res = self.session.get(f"{self.base_url}/v2/purchases", params=params,
                                   timeout=60)
            if res.status_code != 200:
                raise SystemExit(f"[eResto] GET /v2/purchases HTTP "
                                 f"{res.status_code}: {res.text[:200]}")
            rows = res.json().get("purchases", [])
            for p in rows:
                out[p["id"]] = p
            if len(rows) < 200:
                break
            page += 1
        return out

    def fetch_purchases(self, date_from, date_to):
        results = self._fetch_paged({"date[geq]": date_from, "date[leq]": date_to})
        results.update(self._fetch_paged({
            "created_at[geq]": f"{date_from}T00:00:00",
            "created_at[leq]": f"{date_to}T23:59:59"}))
        return list(results.values())

    def fetch_purchase_detail(self, purchase_id):
        params = {"domain": ERESTO_DOMAIN, "token": self.token}
        res = self.session.get(f"{self.base_url}/v2/purchases/{purchase_id}/detail",
                               params=params, timeout=60)
        if res.status_code == 200:
            return res.json().get("purchase_detail", {})
        return {}


_q = urllib.parse.quote


def resolve_uom(code, unit, master):
    """(uom, cf, catatan). Tidak pernah menebak diam-diam: fallback dicatat."""
    su = master["stock_uom"]
    uoms = master["uoms"]
    unit = {"Kg": "kg", "KG": "kg"}.get(unit, unit)
    if unit == su:
        return unit, uoms.get(su, 1.0), ""
    if unit in uoms:
        return unit, uoms[unit], ""
    special = UOM_SPECIAL.get((code, unit))
    if special:
        return special[0], special[1], special[2]
    return su, uoms.get(su, 1.0), f"UOM '{unit}' tidak ada di master -> fallback {su} cf=1"


# ============================================================
# ONBOARDING COMPANY (idempoten)
# ============================================================

def ensure_company(erp, dry_run, log):
    st, d = erp.get("/api/resource/Company/" + _q(COMPANY))
    if st == 200:
        log.append({"Step": "Company", "Detail": f"{COMPANY} (abbr {ABBR})",
                    "Status": "SUDAH ADA"})
    elif dry_run:
        log.append({"Step": "Company", "Detail": f"akan dibuat: {COMPANY} (abbr {ABBR})",
                    "Status": "DRY-RUN"})
        return
    else:
        st, d = erp.req("POST", "/api/resource/Company", {
            "company_name": COMPANY, "abbr": ABBR, "default_currency": "IDR",
            "country": "Indonesia", "parent_company": "OUTLET",
            "create_chart_of_accounts_based_on": "Existing Company",
            "existing_company": "OUTLET"}, timeout=300)
        if st not in (200, 201):
            st2, _ = erp.get("/api/resource/Company/" + _q(COMPANY))
            if st2 != 200:
                raise SystemExit(f"[prep] GAGAL buat Company {COMPANY}: HTTP {st} "
                                 f"{json.dumps(d)[:300]}")
            log.append({"Step": "Company", "Detail": f"{COMPANY} (abbr {ABBR})",
                        "Status": "DIBUAT (recover setelah timeout)"})
        else:
            log.append({"Step": "Company", "Detail": f"{COMPANY} (abbr {ABBR})",
                        "Status": "DIBUAT"})
            print(f"  [prep] company dibuat: {COMPANY} ({ABBR})")

    # master turunan wajib ada (dibuat otomatis oleh Company insert)
    missing = []
    for dt, nm in (("Warehouse", f"Stores - {ABBR}"),
                   ("Account", f"4210.000 - HPP Pembelian - {ABBR}"),
                   ("Cost Center", f"Main - {ABBR}")):
        st, _ = erp.get(f"/api/resource/{_q(dt)}/{_q(nm)}")
        if st != 200:
            missing.append(f"{dt} '{nm}'")
    if missing:
        raise SystemExit("[prep] master turunan company kurang: " + "; ".join(missing))
    log.append({"Step": "Master turunan", "Detail": f"Stores/4210/Cost Center - {ABBR}",
                "Status": "OK"})

    st, d = erp.get("/api/resource/Company/" + _q(COMPANY))
    if st == 200 and not d["data"].get("enable_perpetual_inventory"):
        st, r = erp.put("Company", COMPANY, {"enable_perpetual_inventory": 1})
        if st != 200:
            raise SystemExit(f"[prep] GAGAL set perpetual inventory: "
                             f"{json.dumps(r)[:200]}")
        log.append({"Step": "Perpetual inventory", "Detail": "0 -> 1 (parity outlet lain)",
                    "Status": "DIPERBAIKI"})


def ensure_fiscal_year(erp, dry_run, log):
    st, d = erp.get("/api/resource/Fiscal Year/2026")
    if st != 200:
        raise SystemExit("[prep] Fiscal Year 2026 tidak ada")
    comps = [x.get("company") for x in d["data"].get("companies", [])]
    if COMPANY in comps:
        log.append({"Step": "Fiscal Year", "Detail": "2026",
                    "Status": "SUDAH ADA"})
        return
    if dry_run:
        log.append({"Step": "Fiscal Year", "Detail": "akan didaftarkan ke 2026",
                    "Status": "DRY-RUN"})
        return
    st, r = erp.put("Fiscal Year", "2026",
                    {"companies": [{"company": c} for c in comps + [COMPANY]]})
    if st != 200:
        raise SystemExit(f"[prep] GAGAL update Fiscal Year: {json.dumps(r)[:250]}")
    log.append({"Step": "Fiscal Year", "Detail": "2026",
                "Status": "DITAMBAHKAN"})
    print(f"  [prep] Fiscal Year 2026: +{COMPANY}")


def ensure_supplier_allowed(erp, dry_run, log):
    st, d = erp.get("/api/resource/Supplier/" + _q(JURI))
    if st != 200:
        raise SystemExit("[prep] Supplier JURI tidak ditemukan")
    comps = [x["company"] for x in d["data"].get("companies", [])]
    if COMPANY in comps:
        log.append({"Step": "Allowed To Transact With", "Detail": JURI,
                    "Status": "SUDAH ADA"})
        return
    if dry_run:
        log.append({"Step": "Allowed To Transact With",
                    "Detail": f"akan ditambah ke {JURI}", "Status": "DRY-RUN"})
        return
    st, r = erp.put("Supplier", JURI,
                    {"companies": [{"company": c} for c in comps + [COMPANY]]})
    if st != 200:
        raise SystemExit(f"[prep] GAGAL update supplier: {json.dumps(r)[:250]}")
    log.append({"Step": "Allowed To Transact With", "Detail": JURI,
                "Status": f"DITAMBAHKAN (total {len(comps) + 1})"})
    print(f"  [prep] supplier JURI: +{COMPANY} diizinkan")


def ensure_customer_internal(erp, dry_run, log):
    st, d = erp.get("/api/resource/Customer", {
        "filters": json.dumps([["represents_company", "=", COMPANY],
                               ["is_internal_customer", "=", 1]]),
        "fields": json.dumps(["name", "companies"]), "limit_page_length": 1})
    rows = d.get("data", []) if st == 200 else []
    if rows:
        code = rows[0]["name"]
        comps = [x.get("company") for x in rows[0].get("companies", [])]
        if JURI not in comps and not dry_run:
            st, r = erp.put("Customer", code,
                            {"companies": [{"company": c} for c in comps + [JURI]]})
            if st != 200:
                raise SystemExit(f"[prep] GAGAL set companies {code}: "
                                 f"{json.dumps(r)[:200]}")
            log.append({"Step": "Customer internal", "Detail": f"{code} +allowed {JURI}",
                        "Status": "DIPERBAIKI"})
        else:
            log.append({"Step": "Customer internal", "Detail": code,
                        "Status": "SUDAH ADA"})
        return
    if dry_run:
        log.append({"Step": "Customer internal", "Detail": "akan dibuat (internal)",
                    "Status": "DRY-RUN"})
        return
    st, d = erp.post("Customer", {
        "customer_name": COMPANY, "customer_type": "Company",
        "customer_group": "Outlet Reguler", "territory": "Indonesia",
        "is_internal_customer": 1, "represents_company": COMPANY,
        "companies": [{"company": JURI}]})
    if st != 200:
        raise SystemExit(f"[prep] GAGAL buat customer {COMPANY}: "
                         f"{json.dumps(d)[:250]}")
    code = d["data"]["name"]
    log.append({"Step": "Customer internal", "Detail": f"{code} (internal)",
                "Status": "DIBUAT"})
    print(f"  [prep] customer internal dibuat: {COMPANY} -> {code}")


# ============================================================
# ALUR UTAMA
# ============================================================

def get_existing_index(erp):
    """{name: row} + alias title 'ERESTO <PO>' -> row."""
    index = {}
    for row in erp.list_all("Purchase Order",
                            ["name", "title", "docstatus", "grand_total", "status"]):
        index[row["name"]] = row
        title = row.get("title") or ""
        if title.startswith("ERESTO "):
            index[title.replace("ERESTO ", "").strip()] = row
    return index


def fetch_targets(eresto, only=None):
    raw = eresto.fetch_purchases(PO_DATE, PO_DATE)
    targets = [p for p in raw if p.get("purchase_number") in set(PO_LIST)]
    if only:
        targets = [p for p in targets if p.get("purchase_number") in set(only)]
    details = []
    with ThreadPoolExecutor(max_workers=5) as ex:
        futs = {ex.submit(eresto.fetch_purchase_detail, p["id"]): p for p in targets}
        for fu in as_completed(futs):
            d = fu.result()
            if d:
                details.append(d)
    details.sort(key=lambda x: safe_str(x.get("purchase_number")))
    return details


def run(dry_run=False, verify_only=False, only=None):
    eresto = ErestoClient(ERESTO_BASE_URL, ERESTO_TOKEN, ERESTO_EMAIL, ERESTO_USER_ID)
    erp = Client(ERPNEXT_URL, ERPNEXT_KEY, label="BARU", pause=0.1)

    print("=" * 84)
    print("  MIGRASI PO eResto PO222.091026.001/.002 (RS MUHAMMADIYAH "
          "SELOGIRI) -> ERPNext BARU")
    print(f"  Target: {ERPNEXT_URL}" + ("  [DRY-RUN]" if dry_run else "") +
          ("  [VERIFY]" if verify_only else ""))
    print("=" * 84)

    print(f"\n[1/5] Tarik PO dari eResto (window {PO_DATE}) ...")
    details = fetch_targets(eresto, only)
    print(f"  {len(details)} detail PO222 ditemukan: "
          f"{', '.join(safe_str(d.get('purchase_number')) for d in details)}")
    if not details:
        raise SystemExit("[eResto] tidak ada PO target di window tanggal")

    # ---------------- onboarding (dilewati saat --verify) ----------------
    onboard = []
    if not verify_only:
        print("[2/5] Onboarding company (idempoten) ...")
        ensure_company(erp, dry_run, onboard)
        ensure_fiscal_year(erp, dry_run, onboard)
        ensure_supplier_allowed(erp, dry_run, onboard)
        ensure_customer_internal(erp, dry_run, onboard)

    # ---------------- master item ----------------
    print("[3/5] Ambil master item ERPNext ...")
    items_master = {}
    for code in sorted(set(ITEM_MAP.values())):
        st, d = erp.get("/api/resource/Item/" + _q(code),
                        {"fields": json.dumps(["name", "item_name", "stock_uom",
                                               "disabled", "uoms"])})
        if st == 200:
            items_master[code] = {
                "name": d["data"].get("name"),
                "item_name": d["data"].get("item_name"),
                "stock_uom": d["data"].get("stock_uom"),
                "disabled": bool(d["data"].get("disabled")),
                "uoms": {u["uom"]: float(u["conversion_factor"])
                         for u in d["data"].get("uoms", [])},
            }
    print(f"  {len(items_master)} item master termuat "
          f"(dari {len(set(ITEM_MAP.values()))} target).")

    # ---------------- migrasi ----------------
    print("[4/5] Migrasi ...")
    existing = get_existing_index(erp)
    migrated, skipped, raw_rows = [], [], []
    verify_issues = []

    for d in details:
        pnum = safe_str(d.get("purchase_number"))
        pdate = safe_str(d.get("date"))[:10]
        div = safe_str((d.get("division") or {}).get("name"))
        supp = safe_str((d.get("supplier") or {}).get("name"))
        st_eresto = safe_str(d.get("status"))
        total_cost = safe_float(d.get("total_cost"))
        items = d.get("purchase_items") or []
        charges = d.get("purchase_charges") or []
        discounts = d.get("purchase_discounts") or []

        base_rec = {"PO Number": pnum, "Tanggal": pdate, "Company": div,
                    "Total eResto": total_cost, "Status eResto": st_eresto}

        # guard: charges/diskon tak terduga
        if charges or discounts:
            skipped.append({**base_rec, "Supplier": supp, "Kategori": "Charges/Diskon",
                            "Alasan": f"{len(charges)} charges & {len(discounts)} "
                                      "discounts; butuh pemetaan manual"})
            continue

        # guard: divisi & supplier sesuai target
        if div != COMPANY:
            skipped.append({**base_rec, "Supplier": supp, "Kategori": "Divisi",
                            "Alasan": f"divisi '{div}' bukan '{COMPANY}'"})
            continue
        if supp not in WAREHOUSE_MAP:
            skipped.append({**base_rec, "Supplier": supp, "Kategori": "Supplier",
                            "Alasan": f"supplier '{supp}' belum ada di WAREHOUSE_MAP"})
            continue
        source_wh = WAREHOUSE_MAP[supp]

        # mode --verify: read-only, laporkan dokumen yang belum ada
        if verify_only and pnum not in existing:
            verify_issues.append(f"{pnum}: TIDAK ADA di ERPNext BARU")
            migrated.append({**base_rec, "ERPNext ID": "-", "Supplier": JURI,
                             "Target WH": f"Stores - {ABBR}", "Source WH": source_wh,
                             "Jumlah Baris": len(items),
                             "Grand Total ERPNext": "", "Docstatus": "",
                             "Status ERPNext": "", "Hasil": "TIDAK ADA",
                             "Catatan": ""})
            continue

        # idempotensi (+ self-heal rename)
        if pnum in existing:
            rec = existing[pnum]
            actual_name = rec.get("name")
            rename_note = ""
            if actual_name != pnum and not dry_run and not verify_only:
                st_ren, body = erp.rename("Purchase Order", actual_name, pnum)
                if st_ren == 200:
                    print(f"  [REPAIR] {pnum}: rename {actual_name} -> {pnum}")
                    rename_note = f"repair rename {actual_name} -> {pnum}"
                    actual_name = pnum
                    rec = (erp.get("/api/resource/Purchase Order/" + _q(pnum))[1]
                           .get("data") or rec)
                else:
                    rename_note = f"rename masih gagal: {str(body)[:120]}"
            print(f"  [SKIP_EXISTS] {pnum} -> {actual_name} "
                  f"(ds={rec.get('docstatus')})")
            migrated.append({**base_rec, "ERPNext ID": actual_name,
                             "Supplier": JURI, "Target WH": f"Stores - {ABBR}",
                             "Source WH": source_wh, "Jumlah Baris": "",
                             "Grand Total ERPNext": rec.get("grand_total"),
                             "Docstatus": rec.get("docstatus"),
                             "Status ERPNext": rec.get("status"),
                             "Hasil": "SKIP_EXISTS",
                             "Catatan": rename_note or "sudah ada (rerun aman)"})
            continue

        # susun baris item
        rows_payload, row_audit, row_problems = [], [], []
        for it in items:
            inv = it.get("inventory") or {}
            sku = safe_str(inv.get("sku"))
            qty = safe_float(it.get("amount"))
            price = safe_float(it.get("price"))
            unit = safe_str((it.get("unit") or {}).get("name"))

            code = ITEM_MAP.get(sku)
            if not code:
                row_problems.append(f"SKU {sku} belum ada mapping")
                continue
            master = items_master.get(code)
            if not master:
                row_problems.append(f"Item {code} tidak ada di master")
                continue
            if master.get("disabled"):
                row_problems.append(f"Item {code} disabled")
                continue
            if qty <= 0:
                row_problems.append(f"SKU {sku} qty {qty} (skip baris)")
                continue

            uom, cf, note = resolve_uom(code, unit, master)
            rows_payload.append({
                "item_code": code, "qty": qty, "uom": uom,
                "conversion_factor": cf, "rate": price,
                "warehouse": f"Stores - {ABBR}",
                "expense_account": f"4210.000 - HPP Pembelian - {ABBR}",
                "cost_center": f"Main - {ABBR}", "schedule_date": pdate,
            })
            row_audit.append({
                "PO Number": pnum, "Tanggal": pdate, "Company (Divisi)": div,
                "SKU": sku, "Nama eResto": safe_str(inv.get("name")),
                "Qty eResto": qty, "Unit eResto": unit, "Harga eResto": price,
                "Item Code": code, "Item Name": master.get("item_name"),
                "UOM ERPNext": uom, "CF": cf, "Stock Qty": qty * cf,
                "Rate": price, "Amount": round(qty * price, 2),
                "Catatan UOM": note, "Status eResto": st_eresto})

        if row_problems or not rows_payload:
            skipped.append({**base_rec, "Supplier": supp, "Kategori": "Item",
                            "Alasan": "; ".join(row_problems) or "tidak ada baris valid"})
            continue
        raw_rows.extend(row_audit)

        expected_ds = 1 if st_eresto in ("processing", "completed") else 0
        payload = {
            "doctype": "Purchase Order", "company": COMPANY, "supplier": JURI,
            "transaction_date": pdate, "schedule_date": pdate,
            "set_warehouse": f"Stores - {ABBR}",
            "set_from_warehouse": source_wh,
            "buying_price_list": PRICE_LIST, "currency": "IDR",
            "title": f"ERESTO {pnum}", "items": rows_payload,
        }

        if dry_run:
            hitung = sum(r["qty"] * r["rate"] for r in rows_payload)
            print(f"  [DRY-RUN] {pnum} <- {supp} ({len(rows_payload)} baris, "
                  f"status {st_eresto}, target ds={expected_ds}, "
                  f"baris Rp {hitung:,.0f} vs eResto Rp {total_cost:,.0f})")
            migrated.append({**base_rec, "ERPNext ID": "(dry-run)", "Supplier": JURI,
                             "Target WH": f"Stores - {ABBR}", "Source WH": source_wh,
                             "Jumlah Baris": len(rows_payload),
                             "Grand Total ERPNext": "", "Docstatus": "",
                             "Status ERPNext": "", "Hasil": "DRY_RUN", "Catatan": ""})
            continue

        # POST + rename + submit
        print(f"  [POST] {pnum} <- {supp} ({len(rows_payload)} baris) ...")
        st, res = erp.post("Purchase Order", payload)
        if st != 200:
            msg = json.dumps(res)[:300]
            print(f"    [GAGAL] {pnum}: {msg}")
            migrated.append({**base_rec, "ERPNext ID": "", "Supplier": JURI,
                             "Target WH": f"Stores - {ABBR}", "Source WH": source_wh,
                             "Jumlah Baris": len(rows_payload),
                             "Grand Total ERPNext": "", "Docstatus": 0,
                             "Status ERPNext": "GAGAL", "Hasil": "FAIL_POST",
                             "Catatan": msg})
            continue

        temp_name = res.get("data", {}).get("name")
        final_name, rename_note = temp_name, ""
        st_ren, _ = erp.rename("Purchase Order", temp_name, pnum)
        if st_ren == 200:
            final_name = pnum
        else:
            rename_note = f"rename gagal -> tetap '{temp_name}'"

        docstatus_final, status_final = 0, "Draft"
        if expected_ds == 1:
            st_sub, res_sub = erp.submit("Purchase Order", final_name)
            if st_sub == 200:
                docstatus_final = 1
                status_final = "Submitted"
            else:
                rename_note = (rename_note + " | " if rename_note else "") + \
                              f"submit gagal: {json.dumps(res_sub)[:150]}"

        rec = (erp.get("/api/resource/Purchase Order/" + _q(final_name))[1]
               .get("data") or {})
        print(f"    [OK] {final_name} | ds={docstatus_final} | {status_final} | "
              f"Rp {rec.get('grand_total', '?')}")
        migrated.append({**base_rec, "ERPNext ID": final_name, "Supplier": JURI,
                         "Target WH": f"Stores - {ABBR}", "Source WH": source_wh,
                         "Jumlah Baris": len(rows_payload),
                         "Grand Total ERPNext": rec.get("grand_total", ""),
                         "Docstatus": docstatus_final, "Status ERPNext": status_final,
                         "Hasil": "BERHASIL", "Catatan": rename_note})
        time.sleep(0.2)

    # ---------------- verifikasi ----------------
    print("[5/5] Verifikasi ...")
    ver_rows = []
    if not dry_run:
        for m in migrated:
            if m["Hasil"] not in ("BERHASIL", "SKIP_EXISTS"):
                continue
            st, d = erp.get("/api/resource/Purchase Order/" + _q(m["ERPNext ID"]))
            rec = d.get("data") if st == 200 else None
            if not rec:
                verify_issues.append(f"{m['PO Number']}: dokumen hilang")
                continue
            exp_total = m["Total eResto"]
            got_total = safe_float(rec.get("grand_total"))
            diff = abs(got_total - exp_total)
            n_items = len(rec.get("items", []))
            line_ok = all(
                abs(safe_float(x.get("qty")) * safe_float(x.get("rate"))
                    - safe_float(x.get("amount"))) < 0.01
                and abs(safe_float(x.get("stock_qty"))
                        - safe_float(x.get("qty")) * safe_float(x.get("conversion_factor"))) < 1e-6
                for x in rec.get("items", []))
            st_eresto = next((x["Status eResto"] for x in migrated
                              if x["PO Number"] == m["PO Number"]), "")
            exp_ds = 1 if st_eresto in ("processing", "completed") else 0
            ds_ok = rec.get("docstatus") == exp_ds
            ok = diff <= 1 and line_ok and ds_ok
            if not ok:
                verify_issues.append(f"{m['PO Number']}: diff={diff} "
                                     f"line_ok={line_ok} ds={rec.get('docstatus')} "
                                     f"(harus {exp_ds})")
            ver_rows.append({"PO": m["PO Number"], "ERPNext ID": m["ERPNext ID"],
                             "Grand Total eResto": exp_total,
                             "Grand Total ERPNext": got_total, "Selisih": diff,
                             "Baris": n_items, "Docstatus": rec.get("docstatus"),
                             "Line OK": line_ok, "Status": "LOLOS" if ok else "CEK"})
    else:
        # dry-run: cek hitung baris vs total eResto
        for d in details:
            pnum = safe_str(d.get("purchase_number"))
            hitung = sum(safe_float(it.get("amount")) * safe_float(it.get("price"))
                         for it in (d.get("purchase_items") or []))
            tot = safe_float(d.get("total_cost"))
            ok = abs(hitung - tot) <= 1
            if not ok:
                verify_issues.append(f"{pnum}: jumlah baris {hitung} != total {tot}")
            ver_rows.append({"PO": pnum, "ERPNext ID": "(dry-run)",
                             "Grand Total eResto": tot, "Grand Total ERPNext": hitung,
                             "Selisih": abs(hitung - tot), "Baris": len(d.get("purchase_items") or []),
                             "Docstatus": "", "Line OK": ok,
                             "Status": "LOLOS" if ok else "CEK"})

    # ---------------- laporan Excel ----------------
    wb = Workbook()
    header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    thin = "D3D3D3"
    thin_border = Border(left=Side(style="thin", color=thin),
                         right=Side(style="thin", color=thin),
                         top=Side(style="thin", color=thin),
                         bottom=Side(style="thin", color=thin))

    def write_sheet(ws, title, rows, columns):
        ws.title = title
        ws.append(columns)
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
        ws.freeze_panes = "A2"
        for ci in range(1, ws.max_column + 1):
            c = ws.cell(row=1, column=ci)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                min_col=1, max_col=ws.max_column):
            for c in row:
                c.border = thin_border
        for col in ws.columns:
            ml = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(ml + 3, 12)

    rekap_cols = ["PO Number", "ERPNext ID", "Tanggal", "Company", "Supplier",
                  "Target WH", "Source WH", "Jumlah Baris", "Total eResto",
                  "Grand Total ERPNext", "Docstatus", "Status ERPNext", "Hasil",
                  "Status eResto", "Catatan"]
    raw_cols = ["PO Number", "Tanggal", "Company (Divisi)", "SKU", "Nama eResto",
                "Qty eResto", "Unit eResto", "Harga eResto", "Item Code", "Item Name",
                "UOM ERPNext", "CF", "Stock Qty", "Rate", "Amount", "Catatan UOM",
                "Status eResto"]
    skip_cols = ["PO Number", "Tanggal", "Company", "Supplier", "Total eResto",
                 "Status eResto", "Kategori", "Alasan"]

    write_sheet(wb.active, "Onboarding", onboard, ["Step", "Detail", "Status"])
    write_sheet(wb.create_sheet(), "Rekap_Migrasi", migrated, rekap_cols)
    write_sheet(wb.create_sheet(), "Raw_eResto", raw_rows, raw_cols)
    write_sheet(wb.create_sheet(), "Skipped", skipped, skip_cols)

    n_ok = sum(1 for m in migrated if m["Hasil"] == "BERHASIL")
    n_skip = sum(1 for m in migrated if m["Hasil"] == "SKIP_EXISTS")
    n_fail = sum(1 for m in migrated if m["Hasil"] == "FAIL_POST")
    total_eresto = sum(m["Total eResto"] for m in migrated
                       if m["Hasil"] in ("BERHASIL", "SKIP_EXISTS"))
    total_erp = sum(safe_float(m["Grand Total ERPNext"]) for m in migrated
                    if m["Hasil"] in ("BERHASIL", "SKIP_EXISTS"))
    val_rows = [
        {"Parameter": "PO eResto PO222 ditemukan", "Nilai": len(details),
         "Status": "OK", "Keterangan": f"window {PO_DATE}; divisi {COMPANY}"},
        {"Parameter": "PO BERHASIL dimigrasikan", "Nilai": n_ok,
         "Status": ("DRY-RUN" if dry_run else "LOLOS"),
         "Keterangan": "ID identik eResto (rename_doc)"},
        {"Parameter": "PO SKIP_EXISTS (rerun)", "Nilai": n_skip, "Status": "OK",
         "Keterangan": "idempoten, 0 duplikat"},
        {"Parameter": "PO GAGAL", "Nilai": n_fail,
         "Status": "LOLOS" if n_fail == 0 else "CEK", "Keterangan": "harus 0"},
        {"Parameter": "Dokumen di-skip (guard)", "Nilai": len(skipped),
         "Status": "LOLOS" if not skipped else "CEK",
         "Keterangan": "tidak ditebak; lihat sheet Skipped"},
        {"Parameter": "Total nilai eResto vs ERPNext",
         "Nilai": f"{total_eresto:,.2f} vs {total_erp:,.2f}",
         "Status": "LOLOS" if (not dry_run and abs(total_eresto - total_erp) <= 1)
                   else ("DRY-RUN" if dry_run else "CEK"),
         "Keterangan": "toleransi ≤ Rp1"},
        {"Parameter": "Verifikasi live (total, baris, docstatus, line math)",
         "Nilai": f"{len(ver_rows)} dok",
         "Status": "LOLOS" if not verify_issues else "CEK",
         "Keterangan": "; ".join(verify_issues) if verify_issues else "semua klop"},
    ]
    write_sheet(wb.create_sheet(), "Validation", val_rows,
                ["Parameter", "Nilai", "Status", "Keterangan"])

    note_rows = [
        {"Topik": "Scope",
         "Catatan": f"{', '.join(PO_LIST)} divisi {COMPANY}; supplier internal gudang -> "
                    f"Supplier {JURI}; target 'Stores - {ABBR}'; source WH per supplier; "
                    "status processing -> submit. Rerun aman (idempoten)."},
        {"Topik": "Company baru",
         "Catatan": f"{COMPANY} (abbr {ABBR}) dibuat di bawah OUTLET + Fiscal Year 2026 + "
                    f"Allowed To Transact With supplier JURI + customer internal "
                    "(pola sama dengan RSUD KRAMAT JATI 2026-10-09)."},
        {"Topik": "Bukti pemetaan SKU peralatan",
         "Catatan": "PO lama PO332.100926.002 (RS PKU AISYIYAH BOYOLALI, Gudang Outlet - "
                    "PH Klaten) identik 1:1 dgn PO222.091026.001 (nama/qty/harga): "
                    "BB0262->PEL260059, BB0217->PEL260056, BB0183->PEL260058, "
                    "BB0199->PEL260060, BB0238->PEL260057. BB0291->PER260004 sudah di "
                    "ITEM_MAP migrate_po_eresto_8_9_okt.py (PO223/PO332 seri sama)."},
    ]
    write_sheet(wb.create_sheet(), "Catatan", note_rows, ["Topik", "Catatan"])

    wb.save(REPORT_PATH)
    print(f"\n  Laporan: {REPORT_PATH}")
    print("=" * 84)
    print(f"  BERHASIL : {n_ok} | SKIP_EXISTS : {n_skip} | GAGAL : {n_fail} | "
          f"SKIPPED : {len(skipped)}")
    if not dry_run:
        print(f"  Nilai eResto : Rp {total_eresto:,.2f}")
        print(f"  Nilai ERPNext: Rp {total_erp:,.2f}")
        print(f"  Verifikasi   : {'OK' if not verify_issues else verify_issues}")
    print("=" * 84)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dry-run", action="store_true", help="tanpa POST/PUT")
    ap.add_argument("--verify", action="store_true",
                    help="hanya verifikasi kondisi live (read-only)")
    ap.add_argument("--only", nargs="*", default=None, help="batasi ke nomor PO")
    args = ap.parse_args()
    run(dry_run=args.dry_run, verify_only=args.verify, only=args.only)


if __name__ == "__main__":
    main()
