"""Transfer Item Images from Old ERPNext Site to New ERPNext Cloud.

Old Site: https://rotiropi.j.frappe.cloud
New Site: https://ropierpnext.frappe.cloud

Langkah:
1. Ambil semua Item dengan field `image` terisi dari site lama.
2. Download binary gambar (baik public /files/ maupun private /private/files/).
3. Upload binary gambar ke site baru via /api/method/upload_file dengan is_private=0 (public).
4. Update field `image` pada Item Master di site baru via frappe.client.set_value.
5. Verifikasi bahwa gambar terpasang dan URL gambar dapat diakses secara publik (HTTP 200).
6. Export laporan verifikasi ke Excel pos_sync/Transfer_Item_Images_Report.xlsx.
"""

import os
import sys
import time
import urllib.parse
import requests
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

OLD_BASE = "https://rotiropi.j.frappe.cloud"
OLD_KEY = os.environ.get("ERPNEXT_OLD_CLOUD_API_KEY", "")

NEW_BASE = "https://ropierpnext.frappe.cloud"
NEW_KEY = os.environ.get("ERPNEXT_API_KEY", "")

EXCEL_OUT = "pos_sync/Transfer_Item_Images_Report.xlsx"

# Mapping khusus bila ada perbedaan kode item antara site lama vs site baru
CODE_MAPPING = {
    "GM002": "GM005",  # Gempi Choco Crunchy di master baru memakai kode GM005
}


def make_session(key):
    s = requests.Session()
    s.headers.update({
        "Authorization": f"token {key}",
        "Accept": "application/json"
    })
    return s


def main():
    print("=" * 70)
    print("MIGRASI GAMBAR ITEM DARI SITE LAMA KE SITE BARU")
    print("=" * 70)
    print(f"Site Lama : {OLD_BASE}")
    print(f"Site Baru : {NEW_BASE}")
    print("-" * 70)

    s_old = make_session(OLD_KEY)
    s_new = make_session(NEW_KEY)

    # 1. Verifikasi koneksi kedua site
    r_user_old = s_old.get(f"{OLD_BASE}/api/method/frappe.auth.get_logged_user", timeout=15)
    if r_user_old.status_code != 200:
        print(f"FATAL: Gagal login ke site lama: {r_user_old.status_code} {r_user_old.text}")
        sys.exit(1)
    print(f"Koneksi Site Lama OK: {r_user_old.json().get('message')}")

    r_user_new = s_new.get(f"{NEW_BASE}/api/method/frappe.auth.get_logged_user", timeout=15)
    if r_user_new.status_code != 200:
        print(f"FATAL: Gagal login ke site baru: {r_user_new.status_code} {r_user_new.text}")
        sys.exit(1)
    print(f"Koneksi Site Baru OK: {r_user_new.json().get('message')}")

    # 2. Ambil semua item dari site lama yang memiliki gambar
    r_items = s_old.get(f"{OLD_BASE}/api/resource/Item", params={
        "fields": '["name", "item_code", "item_name", "image"]',
        "filters": '[["image", "is", "set"]]',
        "limit_page_length": 1000
    }, timeout=30)
    if r_items.status_code != 200:
        print(f"FATAL: Gagal mengambil daftar item dari site lama: {r_items.status_code}")
        sys.exit(1)

    old_items = r_items.json().get("data", [])
    print(f"Ditemukan {len(old_items)} item dengan gambar di site lama.\n")

    # Cache file bytes & uploaded URLs untuk efisiensi
    file_cache = {}      # filename -> bytes
    uploaded_cache = {}  # filename -> file_url di site baru

    results = []

    for idx, it in enumerate(old_items, 1):
        old_code = it["item_code"]
        old_name = it["item_name"]
        old_image_path = it["image"]

        target_code = CODE_MAPPING.get(old_code, old_code)

        print(f"[{idx}/{len(old_items)}] Memproses {old_code} -> {target_code} ({old_name})...")

        # Cek apakah target item ada di site baru
        r_target = s_new.get(f"{NEW_BASE}/api/resource/Item/{target_code}", timeout=15)
        if r_target.status_code != 200:
            print(f"  [GAGAL] Target item {target_code} tidak ditemukan di site baru!")
            results.append({
                "No": idx,
                "Old Code": old_code,
                "Old Name": old_name,
                "Target Code": target_code,
                "Old Image": old_image_path,
                "New Image": "-",
                "File Size (KB)": 0,
                "Status HTTP": "-",
                "Result": "GAGAL (Item Not Found)"
            })
            continue

        target_item_data = r_target.json().get("data", {})
        target_name = target_item_data.get("item_name")
        current_new_image = target_item_data.get("image")

        # Ekstrak nama file asli dari URL
        filename = os.path.basename(old_image_path)
        # Tentukan content-type
        lower_fn = filename.lower()
        if lower_fn.endswith(".png"):
            content_type = "image/png"
        elif lower_fn.endswith((".jpg", ".jpeg")):
            content_type = "image/jpeg"
        elif lower_fn.endswith(".webp"):
            content_type = "image/webp"
        else:
            content_type = "application/octet-stream"

        # Download gambar jika belum di-cache
        if filename not in file_cache:
            # Handle URL encoding untuk fetch dari site lama
            encoded_old_path = urllib.parse.quote(old_image_path, safe="/:")
            dl_url = f"{OLD_BASE}{encoded_old_path}"
            r_dl = s_old.get(dl_url, timeout=30)
            if r_dl.status_code != 200:
                print(f"  [GAGAL] Gagal mendownload gambar dari {dl_url}: HTTP {r_dl.status_code}")
                results.append({
                    "No": idx,
                    "Old Code": old_code,
                    "Old Name": old_name,
                    "Target Code": target_code,
                    "Old Image": old_image_path,
                    "New Image": "-",
                    "File Size (KB)": 0,
                    "Status HTTP": f"Download HTTP {r_dl.status_code}",
                    "Result": "GAGAL (Download Failed)"
                })
                continue
            file_cache[filename] = r_dl.content
            print(f"  Downloaded {filename}: {len(r_dl.content):,} bytes")

        file_bytes = file_cache[filename]
        file_size_kb = round(len(file_bytes) / 1024, 2)

        # Upload ke site baru jika file belum pernah diupload
        new_file_url = uploaded_cache.get(filename)
        if not new_file_url:
            files_payload = {
                "file": (filename, file_bytes, content_type)
            }
            data_payload = {
                "is_private": 0,
                "doctype": "Item",
                "docname": target_code,
                "fieldname": "image"
            }
            r_up = s_new.post(f"{NEW_BASE}/api/method/upload_file", files=files_payload, data=data_payload, timeout=45)
            if r_up.status_code != 200:
                print(f"  [GAGAL] Upload file gagal: HTTP {r_up.status_code} {r_up.text[:200]}")
                results.append({
                    "No": idx,
                    "Old Code": old_code,
                    "Old Name": old_name,
                    "Target Code": target_code,
                    "Old Image": old_image_path,
                    "New Image": "-",
                    "File Size (KB)": file_size_kb,
                    "Status HTTP": f"Upload HTTP {r_up.status_code}",
                    "Result": "GAGAL (Upload Failed)"
                })
                continue

            resp_json = r_up.json().get("message", {})
            new_file_url = resp_json.get("file_url")
            if not new_file_url:
                print(f"  [GAGAL] Respon upload tidak memiliki file_url: {r_up.text[:200]}")
                results.append({
                    "No": idx,
                    "Old Code": old_code,
                    "Old Name": old_name,
                    "Target Code": target_code,
                    "Old Image": old_image_path,
                    "New Image": "-",
                    "File Size (KB)": file_size_kb,
                    "Status HTTP": "Upload response missing file_url",
                    "Result": "GAGAL"
                })
                continue
            uploaded_cache[filename] = new_file_url
            print(f"  Uploaded to new site: {new_file_url}")
        else:
            print(f"  Reusing uploaded file: {new_file_url}")

        # Update field `image` pada Item di site baru
        r_set = s_new.post(f"{NEW_BASE}/api/method/frappe.client.set_value", data={
            "doctype": "Item",
            "name": target_code,
            "fieldname": "image",
            "value": new_file_url
        }, timeout=20)

        if r_set.status_code != 200:
            print(f"  [GAGAL] set_value image gagal: HTTP {r_set.status_code} {r_set.text[:200]}")
            results.append({
                "No": idx,
                "Old Code": old_code,
                "Old Name": old_name,
                "Target Code": target_code,
                "Old Image": old_image_path,
                "New Image": new_file_url,
                "File Size (KB)": file_size_kb,
                "Status HTTP": f"set_value HTTP {r_set.status_code}",
                "Result": "GAGAL (set_value)"
            })
            continue

        # Verifikasi item dan fetch gambar dari URL publik
        encoded_new_path = urllib.parse.quote(new_file_url, safe="/:")
        verify_fetch_url = f"{NEW_BASE}{encoded_new_path}"
        r_fetch = requests.get(verify_fetch_url, timeout=15)
        if r_fetch.status_code == 200:
            print(f"  [OK] Berhasil! Image tersimpan & terverifikasi publik (HTTP 200, {len(r_fetch.content):,} bytes)")
            results.append({
                "No": idx,
                "Old Code": old_code,
                "Old Name": old_name,
                "Target Code": target_code,
                "Old Image": old_image_path,
                "New Image": new_file_url,
                "File Size (KB)": file_size_kb,
                "Status HTTP": "200 OK",
                "Result": "LOLOS"
            })
        else:
            print(f"  [WARNING] Set value berhasil tapi verify fetch HTTP {r_fetch.status_code}")
            results.append({
                "No": idx,
                "Old Code": old_code,
                "Old Name": old_name,
                "Target Code": target_code,
                "Old Image": old_image_path,
                "New Image": new_file_url,
                "File Size (KB)": file_size_kb,
                "Status HTTP": f"Verify HTTP {r_fetch.status_code}",
                "Result": "VERIFY_WARN"
            })

        # Jeda kecil antar operasi agar server santai
        time.sleep(0.05)

    # 3. Buat Excel Report
    os.makedirs("pos_sync", exist_ok=True)
    df = pd.DataFrame(results)
    with pd.ExcelWriter(EXCEL_OUT, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Hasil_Migrasi_Gambar", index=False)

    # Styling Excel
    wb = openpyxl.load_workbook(EXCEL_OUT)
    ws = wb["Hasil_Migrasi_Gambar"]
    ws.freeze_panes = "A2"

    hdr_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    hdr_font = Font(name="Calibri", size=11, bold=True, color="000000")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        result_cell = row[8]  # Kolom 'Result'
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            if result_cell.value == "LOLOS":
                result_cell.fill = pass_fill
                result_cell.font = Font(name="Calibri", size=10, bold=True, color="375623")
            elif result_cell.value and "GAGAL" in str(result_cell.value):
                result_cell.fill = fail_fill
                result_cell.font = Font(name="Calibri", size=10, bold=True, color="C65911")

    # Auto fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(EXCEL_OUT)
    print("\n" + "=" * 70)
    print(f"Laporan Excel berhasil disimpan: {EXCEL_OUT}")
    print(f"Total Diproses : {len(results)}")
    print(f"Total Lolos    : {sum(1 for r in results if r['Result'] == 'LOLOS')}")
    print(f"Total Gagal    : {sum(1 for r in results if r['Result'] != 'LOLOS')}")
    print("=" * 70)


if __name__ == "__main__":
    main()
