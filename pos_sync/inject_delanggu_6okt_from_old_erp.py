"""Migrasi & Injeksi Transaksi Delanggu 6 Oktober 2026 dari ERPNext Lama ke ERPNext Baru.

Sumber:
- Old ERPNext: https://rotiropi.j.frappe.cloud (API Key [API-KEY-FROM-ENV])
- Filter: Company = "RS PKU MUHAMMADIYAH DELANGGU", posting_date = "2026-10-06"
- Dokumen: 8 Sales Invoice (is_pos = 1) dengan total omzet Rp 171.000,00.

Target:
- New ERPNext: https://ropierpnext.frappe.cloud
- Company = "RS PKU MUHAMMADIYAH DELANGGU", Warehouse = "Stores - RPMD"
- POS Profile = "SHIFT 1 - RS PKU MUHAMMADIYAH DELANGGU"
- Shift = "POPEN-RPMD-261006006" (periode 2026-10-06)
- POS Package: Paket Hemat 3 New (SINV-RPMD26589) otomatis dikuotasikan via quote_package
  dengan alokasi proporsional (Parent Rp 0, CP001 Rp 14.400, RP003 Rp 9.600).
- Idempotensi: remarks = "OLD-ERP <name>", jika sudah ada dilewati (anti double).
"""

import json
import os
import sys
import time
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import requests

# Load env for new site
ENV_FILE = Path("/tmp/eresto_prod.env")
if not ENV_FILE.exists():
    ENV_FILE = Path("/tmp/eresto_mig.env")

new_env = {}
for line in ENV_FILE.read_text().splitlines():
    if "=" in line and not line.startswith("#"):
        k, v = line.split("=", 1)
        new_env[k.strip()] = v.strip()

NEW_BASE = new_env["ERPNEXT_URL"].rstrip("/")
NEW_HEADERS = {"Authorization": f"token {new_env['ERPNEXT_API_KEY']}"}

OLD_BASE = "https://rotiropi.j.frappe.cloud"
OLD_HEADERS = {"Authorization": f"token {os.environ.get('ERPNEXT_OLD_CLOUD_API_KEY', '')}"}

COMPANY = "RS PKU MUHAMMADIYAH DELANGGU"
WAREHOUSE = "Stores - RPMD"
PROFILE = "SHIFT 1 - RS PKU MUHAMMADIYAH DELANGGU"
OPENING_SHIFT = "POPEN-RPMD-261006006"
CUSTOMER = "Walk-in Customer"
TARGET_DATE = "2026-10-06"

# Mapping pilihan roti untuk Hemat 3 New
HEMAT_3_OPTIONS = {
    "RP001": "6upe790ost",
    "RP002": "bv2elf0eu6",
    "RP003": "bv2el9jp36",
}
HEMAT_3_GROUP = "dzrekvzd"


def fetch_old_invoices():
    r = requests.get(
        f"{OLD_BASE}/api/resource/Sales Invoice"
        f'?filters=[["company","=","{COMPANY}"],["posting_date","=","{TARGET_DATE}"]]'
        f'&fields=["name","posting_date","posting_time","grand_total","docstatus","is_pos"]'
        f"&limit_page_length=200",
        headers=OLD_HEADERS,
    )
    if r.status_code != 200:
        raise RuntimeError(f"Gagal ambil invoice dari site lama: {r.text}")
    invoices = r.json().get("data", [])
    # Urutkan berdasarkan waktu
    invoices.sort(key=lambda x: str(x.get("posting_time")))
    return invoices


def get_existing_new_invoices():
    r = requests.get(
        f"{NEW_BASE}/api/resource/POS Invoice"
        f'?filters=[["company","=","{COMPANY}"],["posting_date","=","{TARGET_DATE}"],["docstatus","!=",2]]'
        f'&fields=["name","remarks","grand_total"]'
        f"&limit_page_length=200",
        headers=NEW_HEADERS,
    )
    res = {}
    for inv in r.json().get("data", []):
        rem = (inv.get("remarks") or "").strip()
        if rem:
            res[rem] = inv
    return res


def quote_hemat_3(roti_code):
    opt_id = HEMAT_3_OPTIONS[roti_code]
    choices = [{"group_key": HEMAT_3_GROUP, "options": [{"option_id": opt_id, "qty": 1}]}]
    r = requests.post(
        f"{NEW_BASE}/api/method/pos_next.api.packages.quote_package",
        json={
            "package": "HEMAT 3 NEW",
            "choices": json.dumps(choices),
            "pos_profile": PROFILE,
        },
        headers=NEW_HEADERS,
    )
    if r.status_code != 200:
        raise RuntimeError(f"Gagal quote HEMAT 3 NEW: {r.text}")
    return r.json().get("message", {})


def close_and_submit_shift(opening_shift_name, date_str):
    print(f"\nClosing shift {opening_shift_name} for date {date_str}...")
    r_prev = requests.post(
        f"{NEW_BASE}/api/method/pos_next.api.shifts.get_closing_shift_data",
        json={"opening_shift": opening_shift_name},
        headers=NEW_HEADERS,
    )
    if r_prev.status_code != 200:
        raise RuntimeError(f"Gagal get_closing_shift_data: {r_prev.text}")

    prev = r_prev.json().get("message", {})
    recon = []
    for row in prev.get("payment_reconciliation") or []:
        exp = float(row.get("expected_amount") or 0)
        recon.append({
            "mode_of_payment": row.get("mode_of_payment"),
            "opening_amount": float(row.get("opening_amount") or 0),
            "expected_amount": exp,
            "closing_amount": exp,
        })

    txns = [
        {
            k: row.get(k)
            for k in (
                "pos_invoice",
                "sales_invoice",
                "posting_date",
                "customer",
                "grand_total",
                "transaction_currency",
                "transaction_amount",
            )
        }
        for row in prev.get("pos_transactions") or []
    ]
    taxes = [
        {k: row.get(k) for k in ("account_head", "rate", "amount")}
        for row in prev.get("taxes") or []
    ]

    closing_payload = {
        "doctype": "POS Closing Shift",
        "pos_opening_shift": opening_shift_name,
        "pos_profile": prev.get("pos_profile"),
        "user": prev.get("user"),
        "company": prev.get("company"),
        "period_start_date": prev.get("period_start_date"),
        "period_end_date": prev.get("period_end_date"),
        "posting_date": date_str,
        "payment_reconciliation": recon,
        "pos_transactions": txns,
        "taxes": taxes,
        "grand_total": float(prev.get("grand_total") or 0),
        "net_total": float(prev.get("net_total") or 0),
        "total_quantity": float(prev.get("total_quantity") or 0),
    }

    r_close = requests.post(
        f"{NEW_BASE}/api/resource/POS Closing Shift",
        json=closing_payload,
        headers=NEW_HEADERS,
    )
    if r_close.status_code != 200:
        raise RuntimeError(f"Gagal create closing shift: {r_close.text}")

    clo_name = r_close.json().get("data", {}).get("name")
    r_sub = requests.put(
        f"{NEW_BASE}/api/resource/POS Closing Shift/{clo_name}",
        json={"docstatus": 1},
        headers=NEW_HEADERS,
    )
    if r_sub.status_code != 200:
        raise RuntimeError(f"Gagal submit closing shift {clo_name}: {r_sub.text}")

    print(f"Successfully closed & submitted {clo_name} (grand total: Rp {prev.get('grand_total'):,.2f})")
    return clo_name


def main():
    print(f"Mengambil transaksi Delanggu per {TARGET_DATE} dari site lama...")
    old_invoices = fetch_old_invoices()
    print(f"Ditemukan {len(old_invoices)} invoice di site lama.")

    existing_in_new = get_existing_new_invoices()
    print(f"Invoice yang sudah ada di site baru dengan tanggal {TARGET_DATE}: {len(existing_in_new)}")

    report_rows = []

    for old_inv_info in old_invoices:
        old_name = old_inv_info["name"]
        rem_marker = f"OLD-ERP {old_name}"

        # Periksa duplikasi
        if rem_marker in existing_in_new:
            print(f"Skipping {old_name}: sudah ada di site baru ({existing_in_new[rem_marker]['name']})")
            report_rows.append({
                "old_invoice": old_name,
                "time": old_inv_info.get("posting_time"),
                "total": old_inv_info.get("grand_total"),
                "new_invoice": existing_in_new[rem_marker]["name"],
                "items_desc": "Sudah Ada",
                "status": "SKIP (SUDAH ADA)",
            })
            continue

        # Ambil full detail doc dari site lama
        r_detail = requests.get(f"{OLD_BASE}/api/resource/Sales Invoice/{old_name}", headers=OLD_HEADERS)
        doc = r_detail.json().get("data", {})
        grand_total = float(doc.get("grand_total") or 0)
        clean_time = str(doc.get("posting_time")).split(".")[0]
        if len(clean_time) == 7:  # misal 9:31:20 -> 09:31:20
            clean_time = "0" + clean_time

        # Persiapkan payments
        pay_rows = []
        for p in doc.get("payments", []):
            mode = p.get("mode_of_payment")
            amt = float(p.get("amount") or 0)
            if amt > 0:
                # Batasi akumulasi tidak melebihi grand_total (karena di kasir ada kolom tendered > net)
                actual_amt = min(amt, grand_total - sum(x["amount"] for x in pay_rows))
                if actual_amt > 0:
                    pay_rows.append({"mode_of_payment": mode, "amount": actual_amt})

        if not pay_rows or sum(x["amount"] for x in pay_rows) != grand_total:
            # Fallback jika payments di doc lama kosong atau selisih
            default_mode = "QRIS" if any(p.get("mode_of_payment") == "QRIS" for p in doc.get("payments", [])) else "Cash"
            pay_rows = [{"mode_of_payment": default_mode, "amount": grand_total}]

        # Persiapkan items
        items = []
        desc_items = []
        has_package = any(it.get("pos_package") or it.get("item_code") == "PAKET26-003" for it in doc.get("items", []))

        if has_package:
            # Cari pilihan roti untuk paket
            roti_code = "RP003"  # default
            for it in doc.get("items", []):
                if it.get("item_code") in ("RP001", "RP002", "RP003") and it.get("pos_package_role") == "Package Item":
                    roti_code = it.get("item_code")

            qres = quote_hemat_3(roti_code)
            inst_id = "pkg-1"

            # Baris parent
            parent_line = qres["lines"][0]
            items.append({
                "item_code": parent_line["item_code"],
                "item_name": parent_line["item_name"],
                "qty": parent_line["qty"],
                "rate": parent_line["rate"],
                "price_list_rate": parent_line["rate"],
                "amount": parent_line["rate"] * parent_line["qty"],
                "pos_package": qres["package"],
                "pos_package_role": parent_line["role"],
                "pos_package_instance": inst_id,
                "pos_package_snapshot": json.dumps(qres["snapshot"]),
                "warehouse": WAREHOUSE,
            })

            # Baris komponen
            for comp in qres["lines"][1:]:
                items.append({
                    "item_code": comp["item_code"],
                    "item_name": comp["item_name"],
                    "qty": comp["qty"],
                    "rate": comp["rate"],
                    "price_list_rate": comp["rate"],
                    "amount": comp["rate"] * comp["qty"],
                    "pos_package": qres["package"],
                    "pos_package_role": comp["role"],
                    "pos_package_instance": inst_id,
                    "warehouse": WAREHOUSE,
                })
            desc_items.append(f"HEMAT 3 NEW (Parent Rp0, {roti_code} Rp9.600, CP001 Rp14.400)")

        else:
            # Non-paket regular items
            for it in doc.get("items", []):
                code = it.get("item_code")
                qty = float(it.get("qty") or 0)
                rate = float(it.get("rate") or 0)
                items.append({
                    "item_code": code,
                    "item_name": it.get("item_name") or code,
                    "qty": qty,
                    "rate": rate,
                    "price_list_rate": rate,
                    "amount": rate * qty,
                    "warehouse": WAREHOUSE,
                })
                desc_items.append(f"{code} ({qty}x @ Rp{rate:,.0f})")

        payload = {
            "customer": CUSTOMER,
            "pos_profile": PROFILE,
            "posa_pos_opening_shift": OPENING_SHIFT,
            "posting_date": TARGET_DATE,
            "posting_time": clean_time,
            "company": COMPANY,
            "items": items,
            "payments": pay_rows,
            "discount_amount": 0,
            "apply_discount_on": "Grand Total",
            "remarks": rem_marker,
        }

        print(f"Injecting {old_name} ({clean_time}) - Rp {grand_total:,.2f}...")
        r_sub = requests.post(
            f"{NEW_BASE}/api/method/pos_next.api.backdate_invoices.submit_backdate_invoice",
            json={"invoice": payload},
            headers=NEW_HEADERS,
        )

        if r_sub.status_code != 200:
            print(f"ERROR: {r_sub.text[:400]}")
            report_rows.append({
                "old_invoice": old_name,
                "time": clean_time,
                "total": grand_total,
                "new_invoice": "ERROR",
                "items_desc": "; ".join(desc_items),
                "status": f"GAGAL: {r_sub.text[:150]}",
            })
        else:
            new_inv_name = r_sub.json().get("message", {}).get("name")
            print(f"-> SUCCESS: {new_inv_name}")
            report_rows.append({
                "old_invoice": old_name,
                "time": clean_time,
                "total": grand_total,
                "new_invoice": new_inv_name,
                "items_desc": "; ".join(desc_items),
                "status": "LOLOS (SUBMITTED)",
            })

    # Close and submit shift for 6 Oktober 2026
    clo_shift = close_and_submit_shift(OPENING_SHIFT, TARGET_DATE)

    # Export report to Excel
    report_file = "pos_sync/Injeksi_Delanggu_6_Okt_Old_ERP.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Injeksi_6_Okt"

    headers = [
        "No",
        "Faktur Site Lama",
        "Jam Transaksi",
        "Total (Rp)",
        "Faktur Site Baru (Resmi)",
        "Rincian Item / Alokasi Paket",
        "Status Verifikasi",
    ]
    ws.append(headers)

    for idx, r in enumerate(report_rows, 1):
        ws.append([
            idx,
            r["old_invoice"],
            r["time"],
            r["total"],
            r["new_invoice"],
            r["items_desc"],
            r["status"],
        ])

    fill_hdr = PatternFill("solid", fgColor="D9EAF7")
    font_bold = Font(bold=True)
    for c in ws[1]:
        c.fill = fill_hdr
        c.font = font_bold
        c.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(w + 3, 12), 65)

    ws.freeze_panes = "A2"
    wb.save(report_file)
    print(f"\nLaporan Excel berhasil disimpan di: {report_file}")


if __name__ == "__main__":
    main()
